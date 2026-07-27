import argparse
import re
import sys
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path

import fitz
import pandas as pd
import tabula
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# =========================
# PDF CONFIG
# =========================
PDF_AREA = [200, 15, 710, 830]
PDF_COLUMNS = [90, 330, 390, 500, 830]

PDF_AREA_SUMMARY = [290, 15, 710, 830]
PDF_COLUMNS_SUMMARY = [180, 300, 370, 500, 830]

BANK_JAKARTA_DIR = Path("pdf_file") / "Bank Jakarta"
EXCEL_DIR = Path("excel_file") / "Bank Jakarta"

MONTH_ALIASES = {
    "JAN": ("Jan", 1),
    "JANUARI": ("Jan", 1),
    "FEB": ("Feb", 2),
    "FEBRUARI": ("Feb", 2),
    "MAR": ("Mar", 3),
    "MARET": ("Mar", 3),
    "APR": ("Apr", 4),
    "APRIL": ("Apr", 4),
    "MEI": ("May", 5),
    "MAY": ("May", 5),
    "JUN": ("Jun", 6),
    "JUNI": ("Jun", 6),
    "JUL": ("Jul", 7),
    "JULI": ("Jul", 7),
    "AGU": ("Aug", 8),
    "AGS": ("Aug", 8),
    "AUG": ("Aug", 8),
    "AGUSTUS": ("Aug", 8),
    "SEP": ("Sep", 9),
    "SEPT": ("Sep", 9),
    "SEPTEMBER": ("Sep", 9),
    "OKT": ("Oct", 10),
    "OCT": ("Oct", 10),
    "OKTOBER": ("Oct", 10),
    "NOV": ("Nov", 11),
    "NOVEMBER": ("Nov", 11),
    "DES": ("Dec", 12),
    "DEC": ("Dec", 12),
    "DESEMBER": ("Dec", 12),
    "JANUARY": ("Jan", 1),
    "FEBRUARY": ("Feb", 2),
    "MARCH": ("Mar", 3),
    "APRIL": ("Apr", 4),
    "JUNE": ("Jun", 6),
    "JULY": ("Jul", 7),
    "AUGUST": ("Aug", 8),
    "OCTOBER": ("Oct", 10),
    "DECEMBER": ("Dec", 12),
}

DATE_MONTHS = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "MEI": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "AGU": 8,
    "SEP": 9,
    "OCT": 10,
    "OKT": 10,
    "NOV": 11,
    "DEC": 12,
    "DES": 12,
}


# =========================
# HELPERS
# =========================
def clean_text(value):
    if pd.isna(value):
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def append_text(base, extra):
    base = clean_text(base)
    extra = clean_text(extra)

    if not extra:
        return base
    if not base:
        return extra
    return f"{base} {extra}"


def parse_number(value):
    text = clean_text(value).upper()
    if text in ("", "NAN"):
        return pd.NA

    text = text.replace(",", "").replace("+", "").strip()

    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return pd.NA


def excel_number(value):
    if pd.isna(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def dataframe_for_excel(df):
    result = df.copy()
    for col in result.columns:
        result[col] = result[col].apply(excel_number)
    return result


def is_money_line(text):
    return bool(re.fullmatch(r"[+-]?\d[\d,]*(?:\.\d{1,2})?", clean_text(text)))


def is_datetime_start(text):
    return bool(re.match(r"^\d{2}\s+[A-Za-z]{3}\s+\d{2}\s+\d{2}:\d{2}$", clean_text(text)))


def date_to_day_month(text):
    match = re.match(r"^(\d{2})\s+([A-Za-z]{3})\s+\d{2}$", clean_text(text))
    if not match:
        return clean_text(text)

    month = DATE_MONTHS.get(match.group(2).upper())
    if not month:
        return clean_text(text)

    return f"{match.group(1)}/{month:02d}"


def parse_mutation_type(jenis, nominal):
    jenis_text = clean_text(jenis).upper()
    number = parse_number(nominal)

    if pd.isna(number):
        return pd.Series([pd.NA, pd.NA], index=["DB", "CR"])
    if "DEBIT" in jenis_text:
        return pd.Series([abs(number), pd.NA], index=["DB", "CR"])
    if "KREDIT" in jenis_text:
        return pd.Series([pd.NA, abs(number)], index=["DB", "CR"])

    return pd.Series([pd.NA, pd.NA], index=["DB", "CR"])


def safe_sheet_name(name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean_text(name)).strip()
    return cleaned[:31] or "Sheet"


def unique_sheet_name(base_name, used_names):
    base_name = safe_sheet_name(base_name)
    candidate = base_name
    counter = 2

    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1

    used_names.add(candidate.lower())
    return candidate


@lru_cache(maxsize=None)
def read_pdf_text_for_metadata(pdf_file):
    pdf_file = Path(pdf_file)
    text_parts = []

    try:
        with fitz.open(pdf_file) as doc:
            for page in doc:
                text_parts.append(page.get_text("text"))
    except Exception:
        return ""

    return "\n".join(text_parts)


def pdf_text_lines(pdf_file):
    text = read_pdf_text_for_metadata(Path(pdf_file))
    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def normalize_metadata_text(text):
    return re.sub(r"\s+", " ", clean_text(text)).upper()


def extract_account_from_text(text):
    normalized = normalize_metadata_text(text)
    match = re.search(r"\bDKI\s*-\s*(\d{6,})\b", normalized)
    if match:
        return match.group(1)

    match = re.search(r"\bREKENING\b[^\d]{0,30}(\d{6,})", normalized)
    if match:
        return match.group(1)

    matches = re.findall(r"\b\d{8,}\b", normalized)
    return matches[0] if matches else None


def extract_period_from_text(text):
    normalized = normalize_metadata_text(text)
    month_pattern = "|".join(sorted(MONTH_ALIASES, key=len, reverse=True))

    match = re.search(rf"\bPERIODE\b\s+({month_pattern})\s+(20\d{{2}})\b", normalized)
    if not match:
        match = re.search(rf"\b({month_pattern})\s+(20\d{{2}})\b", normalized)

    if match:
        month_key, year = match.groups()
        month_name, month_order = MONTH_ALIASES[month_key]
        return f"{month_name}-{year[-2:]}", month_order, year

    return None


def extract_period_from_filename(pdf_file):
    stem = pdf_file.stem.upper()
    year_match = re.search(r"(20\d{2})", stem)
    year = year_match.group(1) if year_match else None

    month_label = None
    month_order = 99
    for token in re.split(r"[^A-Z]+", stem):
        if token in MONTH_ALIASES:
            month_name, month_order = MONTH_ALIASES[token]
            month_label = f"{month_name}-{year[-2:]}" if year else month_name
            break

    if year is None and re.fullmatch(r"20\d{2}", pdf_file.parent.name):
        year = pdf_file.parent.name
        if month_label and "-" not in month_label:
            month_label = f"{month_label}-{year[-2:]}"

    return month_label or pdf_file.stem[:20], month_order, year or "Unknown"


def extract_pdf_metadata(pdf_file):
    pdf_file = Path(pdf_file)
    text = read_pdf_text_for_metadata(pdf_file)
    month_label, month_order, year = extract_period_from_text(text) or extract_period_from_filename(pdf_file)

    account = extract_account_from_text(text)
    if account is None:
        match = re.search(r"(\d{8,})", pdf_file.stem)
        account = match.group(1) if match else pdf_file.stem.split("-")[0]

    return {
        "account": account,
        "month_label": month_label,
        "month_order": month_order,
        "year": str(year),
    }


def extract_month_year(pdf_file):
    metadata = extract_pdf_metadata(pdf_file)
    return metadata["month_label"], metadata["month_order"], metadata["year"]


def extract_account_number_from_text(text):
    return extract_account_from_text(text) or ""


# =========================
# PDF READING
# =========================
def read_pdf_table(pdf_file):
    dfs = tabula.read_pdf(
        str(pdf_file),
        pages="all",
        stream=True,
        guess=False,
        area=PDF_AREA,
        columns=PDF_COLUMNS,
        pandas_options={"header": None},
        multiple_tables=False,
        force_subprocess=True,
        encoding="cp1252",
    )

    if not dfs:
        raise ValueError("No transaction table found. Try adjusting PDF_AREA / PDF_COLUMNS.")

    df = pd.concat(dfs, ignore_index=True)
    df = df.iloc[:, :5].copy()
    df.columns = ["TanggalJam", "Keterangan", "DebitKredit", "Nominal", "SaldoBerjalan"]

    for col in df.columns:
        df[col] = df[col].apply(clean_text)

    return df


def read_summary_page(pdf_file):
    dfs = tabula.read_pdf(
        str(pdf_file),
        pages=1,
        stream=True,
        guess=False,
        area=PDF_AREA_SUMMARY,
        columns=PDF_COLUMNS_SUMMARY,
        pandas_options={"header": None},
        multiple_tables=True,
        force_subprocess=True,
        encoding="cp1252",
    )

    if not dfs:
        return pd.DataFrame(columns=["Rekening", "Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"])

    df = pd.concat(dfs, ignore_index=True)
    df = df.iloc[:, :5].copy()
    df.columns = ["Rekening", "Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"]

    for col in df.columns:
        df[col] = df[col].apply(clean_text)

    return df


def build_summary_from_text(pdf_file):
    lines = pdf_text_lines(pdf_file)
    first_page_end = next(
        (idx for idx, line in enumerate(lines) if re.search(r"E-Statement\s+2\s+dari", line, flags=re.IGNORECASE)),
        len(lines),
    )
    header_lines = lines[:first_page_end]

    account = ""
    for line in header_lines:
        match = re.search(r"\bDKI\s*-\s*(\d{6,})\b", line, flags=re.IGNORECASE)
        if match:
            account = line
            break

    money_values = [line for line in header_lines if is_money_line(line)]
    labels = ["Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"]
    amounts = {label: pd.NA for label in labels}

    if len(money_values) >= 4:
        for label, value in zip(labels, money_values[:4]):
            amounts[label] = parse_number(value)

    return pd.DataFrame([{
        "Rekening": account,
        "Saldo_Awal": amounts["Saldo_Awal"],
        "Transaksi_Masuk": amounts["Transaksi_Masuk"],
        "Transaksi_Keluar": amounts["Transaksi_Keluar"],
        "Saldo_Akhir": amounts["Saldo_Akhir"],
    }], columns=["Rekening", "Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"])


def parse_transactions_from_text(pdf_file):
    lines = pdf_text_lines(pdf_file)
    records = []
    current = None
    in_table = False

    footer_patterns = (
        r"^DISCLAIMER$",
        r"^PT\.?\s*Bank DKI\b",
        r"^www\.bankdki\.co\.id$",
    )
    skip_exact = {
        "TANGGAL & JAM",
        "RINCIAN",
        "SALDO BERJALAN",
        "NOMINAL",
        "DEBIT/KREDIT",
        "E-STATEMENT",
        "RINGKASAN",
    }

    def finish_current():
        nonlocal current
        if not current:
            return

        direction_idx = next(
            (idx for idx in range(len(current["lines"]) - 1, -1, -1)
             if current["lines"][idx].upper() in {"DEBIT", "KREDIT"}),
            None,
        )
        if direction_idx is None:
            current = None
            return

        before_direction = current["lines"][:direction_idx]
        money_indexes = [idx for idx, line in enumerate(before_direction) if is_money_line(line)]
        if len(money_indexes) < 2:
            current = None
            return

        saldo_idx, nominal_idx = money_indexes[-2], money_indexes[-1]
        description_parts = [
            line for idx, line in enumerate(before_direction)
            if idx not in {saldo_idx, nominal_idx}
        ]
        description = " ".join(description_parts).strip()

        records.append({
            "TanggalJam": current["TanggalJam"],
            "Keterangan": description,
            "Jenis": current["lines"][direction_idx],
            "Nominal": before_direction[nominal_idx],
            "Saldo": before_direction[saldo_idx],
        })
        current = None

    for line in lines:
        upper = line.upper()

        if upper == "TANGGAL & JAM":
            in_table = True
            finish_current()
            continue

        if not in_table and is_datetime_start(line):
            in_table = True
            current = {"TanggalJam": line, "lines": []}
            continue

        if not in_table:
            continue

        if any(re.search(pattern, line, flags=re.IGNORECASE) for pattern in footer_patterns):
            finish_current()
            in_table = False
            continue

        if upper in skip_exact:
            continue
        if re.fullmatch(r"E-Statement\s+\d+\s+dari\s+\d+", line, flags=re.IGNORECASE):
            continue
        if re.fullmatch(r"\d+", line):
            continue
        if upper.startswith(("PEMILIK REKENING", "ALAMAT", "MATA UANG", "PERIODE")):
            continue
        if upper in {"RUPIAH"}:
            continue
        if "BANK DKI ADALAH PELAKU" in upper:
            continue

        if is_datetime_start(line):
            finish_current()
            current = {"TanggalJam": line, "lines": []}
            continue

        if current is None:
            continue

        current["lines"].append(line)
        if upper in {"DEBIT", "KREDIT"}:
            finish_current()

    finish_current()
    return pd.DataFrame(records, columns=["TanggalJam", "Keterangan", "Jenis", "Nominal", "Saldo"])


# =========================
# CLEANING
# =========================
def merge_summary_rows(df_summary):
    if df_summary.empty:
        return pd.DataFrame(columns=["Rekening", "Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"])

    records = []
    current = None

    for _, row in df_summary.iterrows():
        rekening = clean_text(row["Rekening"])
        amounts = {
            "Saldo_Awal": parse_number(row["Saldo_Awal"]),
            "Transaksi_Masuk": parse_number(row["Transaksi_Masuk"]),
            "Transaksi_Keluar": parse_number(row["Transaksi_Keluar"]),
            "Saldo_Akhir": parse_number(row["Saldo_Akhir"]),
        }
        has_amount = any(not pd.isna(value) for value in amounts.values())

        if has_amount:
            if current is not None:
                records.append(current)
            current = {"Rekening": rekening, **amounts}
        elif current is not None and rekening:
            current["Rekening"] = append_text(current["Rekening"], rekening)

    if current is not None:
        records.append(current)

    return pd.DataFrame(records, columns=["Rekening", "Saldo_Awal", "Transaksi_Masuk", "Transaksi_Keluar", "Saldo_Akhir"])


def remove_non_transaction_rows(df):
    garbage_regex = (
        r"E-Statement|Ringkasan|Rekening|Saldo Awal|Transaksi Masuk|"
        r"Transaksi Keluar|Saldo Akhir|Tanggal & Jam|Rincian|Debit/Kredit|"
        r"Saldo Berjalan|Bank DKI|Jakarta"
    )

    def keep_row(row):
        joined = clean_text(" ".join(row.astype(str).tolist()))
        if not joined:
            return False
        if re.search(garbage_regex, joined, flags=re.IGNORECASE):
            return False
        return is_datetime_start(row["TanggalJam"]) or not clean_text(row["TanggalJam"])

    return df[df.apply(keep_row, axis=1)].reset_index(drop=True)


def merge_continuation_rows(df):
    merged_rows = []
    current = None

    for _, row in df.iterrows():
        tanggal_jam = clean_text(row["TanggalJam"])
        rincian = clean_text(row["Keterangan"])
        jenis = clean_text(row["DebitKredit"])
        nominal = clean_text(row["Nominal"])
        saldo = clean_text(row["SaldoBerjalan"])

        if is_datetime_start(tanggal_jam):
            if current is not None:
                merged_rows.append(current)

            current = {
                "TanggalJam": tanggal_jam,
                "Keterangan": rincian,
                "Jenis": jenis,
                "Nominal": nominal,
                "Saldo": saldo,
            }
            continue

        if current is None:
            continue

        if tanggal_jam:
            current["Keterangan"] = append_text(current["Keterangan"], tanggal_jam)
        if rincian:
            current["Keterangan"] = append_text(current["Keterangan"], rincian)
        if jenis:
            if current["Jenis"] == "":
                current["Jenis"] = jenis
            else:
                current["Keterangan"] = append_text(current["Keterangan"], jenis)
        if nominal:
            if current["Nominal"] == "":
                current["Nominal"] = nominal
            else:
                current["Keterangan"] = append_text(current["Keterangan"], nominal)
        if saldo:
            if current["Saldo"] == "":
                current["Saldo"] = saldo
            else:
                current["Keterangan"] = append_text(current["Keterangan"], saldo)

    if current is not None:
        merged_rows.append(current)

    return pd.DataFrame(merged_rows, columns=["TanggalJam", "Keterangan", "Jenis", "Nominal", "Saldo"])


def finalize_transactions(df):
    df = df.copy()
    df[["DB", "CR"]] = df.apply(lambda row: parse_mutation_type(row["Jenis"], row["Nominal"]), axis=1)
    df["Saldo"] = df["Saldo"].apply(parse_number)
    df["Tanggal"] = df["TanggalJam"].str.extract(r"^(\d{2}\s+[A-Za-z]{3}\s+\d{2})", expand=False).apply(date_to_day_month)

    df = df[["Tanggal", "Keterangan", "DB", "CR", "Saldo"]].copy()
    return df


# =========================
# SUMMARY / EXPORT HELPERS
# =========================
def summary_value(df_summary, label):
    if df_summary.empty or label not in df_summary.columns:
        return pd.NA
    value = df_summary[label].iloc[0]
    return value if not pd.isna(value) else pd.NA


def sum_amount(df, column):
    values = df[column].dropna()
    if values.empty:
        return pd.NA
    return sum(values, Decimal("0.00"))


def count_amount(df, column):
    return int(df[column].dropna().count())


def sum_by_description(df, amount_column, pattern):
    if df.empty:
        return pd.NA

    mask = df["Keterangan"].str.upper().str.contains(pattern, regex=True, na=False)
    values = df.loc[mask, amount_column].dropna()
    if values.empty:
        return pd.NA
    return sum(values, Decimal("0.00"))


def extract_account_number(df_summary, pdf_file):
    if not df_summary.empty:
        account = extract_account_number_from_text(df_summary["Rekening"].iloc[0])
        if account:
            return account

    fallback = extract_account_number_from_text(pdf_file.stem)
    return fallback or pdf_file.stem.split("-")[0]


def build_month_summary(clean_df, df_summary, pdf_file, metadata=None):
    metadata = metadata or extract_pdf_metadata(pdf_file)

    return {
        "Source File": pdf_file.name,
        "Account": metadata["account"],
        "Year": metadata["year"],
        "MonthOrder": metadata["month_order"],
        "Bulan": metadata["month_label"],
        "Mutasi Debet Nominal (Rp)": summary_value(df_summary, "Transaksi_Keluar"),
        "Mutasi Debet Frek": count_amount(clean_df, "DB"),
        "Mutasi Kredit Nominal (Rp)": summary_value(df_summary, "Transaksi_Masuk"),
        "Mutasi Kredit Frek": count_amount(clean_df, "CR"),
        "Saldo (Rp)": summary_value(df_summary, "Saldo_Akhir"),
        "Saldo Awal (Rp)": summary_value(df_summary, "Saldo_Awal"),
        "Adm": sum_by_description(clean_df, "DB", r"\bADM\b|ADMIN|BIAYA\s+ADMIN|FEE"),
        "Pajak": sum_by_description(clean_df, "DB", r"\bPPH\b|PAJAK|TAX"),
        "Bunga": sum_by_description(clean_df, "CR", r"BUNGA|INTEREST"),
        "Saldo Min": sum_by_description(clean_df, "DB", r"SALDO\s+MIN"),
        "JaGir": sum_by_description(clean_df, "CR", r"JASA\s+GIRO|JAGIR"),
    }


def first_day_label(pdf_file, clean_df):
    _, month_order, _ = extract_month_year(pdf_file)
    if month_order != 99:
        return f"01/{month_order:02d}"

    if not clean_df.empty:
        match = re.match(r"^\d{1,2}/\d{1,2}", clean_text(clean_df["Tanggal"].iloc[0]))
        if match:
            return match.group(0)

    return ""


def build_monthly_display_dataframe(clean_df, df_summary, pdf_file):
    opening_row = pd.DataFrame([{
        "Tanggal": first_day_label(pdf_file, clean_df),
        "Keterangan": "SALDO AWAL",
        "DB": pd.NA,
        "CR": pd.NA,
        "Saldo": summary_value(df_summary, "Saldo_Awal"),
    }])

    return pd.concat([opening_row, clean_df], ignore_index=True)


def reconcile_transactions_with_balance(clean_df, df_summary):
    if clean_df.empty:
        return clean_df

    result = clean_df.copy()
    last_balance = summary_value(df_summary, "Saldo_Awal")

    for idx, row in result.iterrows():
        current_balance = row["Saldo"]
        if pd.isna(current_balance) or pd.isna(last_balance):
            last_balance = current_balance
            continue

        delta = current_balance - last_balance
        if delta < 0:
            result.at[idx, "DB"] = abs(delta)
            result.at[idx, "CR"] = pd.NA
        elif delta > 0:
            result.at[idx, "DB"] = pd.NA
            result.at[idx, "CR"] = delta

        last_balance = current_balance

    return result


def report_extraction_balance(clean_df, df_summary, pdf_file):
    extracted_db = sum(clean_df["DB"].dropna(), Decimal("0.00"))
    extracted_cr = sum(clean_df["CR"].dropna(), Decimal("0.00"))
    summary_db = summary_value(df_summary, "Transaksi_Keluar")
    summary_cr = summary_value(df_summary, "Transaksi_Masuk")

    mismatches = []
    if not pd.isna(summary_db) and extracted_db != summary_db:
        mismatches.append(f"DB extracted {extracted_db} but PDF summary {summary_db}")
    if not pd.isna(summary_cr) and extracted_cr != summary_cr:
        mismatches.append(f"CR extracted {extracted_cr} but PDF summary {summary_cr}")

    if mismatches:
        print(f"PERINGATAN: {Path(pdf_file).name}: " + "; ".join(mismatches))


def auto_fit_columns(sheet):
    for col_idx, col_cells in enumerate(sheet.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        sheet.column_dimensions[col_letter].width = max_length + 2


def add_monthly_sheet_totals(sheet):
    first_data_row = 2
    last_data_row = sheet.max_row
    total_row = last_data_row + 3
    freq_row = total_row + 1
    average_row = total_row + 2

    for row_num, label in {
        total_row: "Total",
        freq_row: "Frekuensi",
        average_row: "Rata2",
    }.items():
        sheet.cell(row=row_num, column=2, value=label)

    for col_num in (3, 4):
        col_letter = get_column_letter(col_num)
        sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=freq_row, column=col_num, value=f"=COUNT({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=average_row, column=col_num, value=f"=IFERROR(AVERAGE({col_letter}{first_data_row}:{col_letter}{last_data_row}),0)")


def style_monthly_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary_start = sheet.max_row - 2

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    for col_letter, width in {"A": 10, "B": 95, "C": 16, "D": 16, "E": 16}.items():
        sheet.column_dimensions[col_letter].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column in (3, 4, 5):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row_num in range(summary_start, sheet.max_row + 1):
        label_cell = sheet.cell(row=row_num, column=2)
        label_cell.fill = header_fill
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_num in (3, 4):
            value_cell = sheet.cell(row=row_num, column=col_num)
            value_cell.number_format = "#,##0.00" if row_num != summary_start + 1 else "0"
            value_cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_num in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 18


def write_monthly_transaction_sheet(writer, clean_df, df_summary, pdf_file, sheet_name):
    display_df = build_monthly_display_dataframe(clean_df, df_summary, pdf_file)
    display_df = dataframe_for_excel(display_df)
    display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    sheet = writer.sheets[sheet_name]
    add_monthly_sheet_totals(sheet)
    style_monthly_sheet(sheet)


def write_summary_sheet(writer, summary_rows):
    workbook = writer.book
    sheet = workbook.create_sheet("Summary", 0)
    writer.sheets["Summary"] = sheet

    headers_top = [
        "No", "Bulan", "Mutasi Debet", None, "Mutasi Kredit", None,
        "Saldo", "Saldo Awal", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]
    headers_bottom = [
        "No", "Bulan", "Nominal (Rp)", "Frek", "Nominal (Rp)", "Frek",
        "(Rp)", "(Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]

    for col_num, value in enumerate(headers_top, start=1):
        if value is not None:
            sheet.cell(row=1, column=col_num, value=value)
    for col_num, value in enumerate(headers_bottom, start=1):
        sheet.cell(row=2, column=col_num, value=value)

    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    sheet.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=4)
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    for col_num in range(7, 14):
        sheet.merge_cells(start_row=1, start_column=col_num, end_row=2, end_column=col_num)

    sorted_rows = sorted(summary_rows, key=lambda row: (row["Year"], row["MonthOrder"], row["Source File"]))
    for row_idx, row in enumerate(sorted_rows, start=3):
        values = [
            row_idx - 2,
            row["Bulan"],
            row["Mutasi Debet Nominal (Rp)"],
            row["Mutasi Debet Frek"],
            row["Mutasi Kredit Nominal (Rp)"],
            row["Mutasi Kredit Frek"],
            row["Saldo (Rp)"],
            row["Saldo Awal (Rp)"],
            row["Adm"],
            row["Pajak"],
            row["Bunga"],
            row["Saldo Min"],
            row["JaGir"],
        ]
        for col_num, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col_num, value=excel_number(value))

    total_row = len(sorted_rows) + 3
    average_row = total_row + 1
    sheet.cell(row=total_row, column=2, value="Total")
    sheet.cell(row=average_row, column=2, value="Rata-Rata")

    data_start = 3
    data_end = total_row - 1
    if data_end >= data_start:
        for col_num in [3, 4, 5, 6, 9, 10, 11, 12, 13]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{data_start}:{col_letter}{data_end})")
        for col_num in [3, 4, 5, 6, 7]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=average_row, column=col_num, value=f"=AVERAGE({col_letter}{data_start}:{col_letter}{data_end})")

    style_summary_sheet(sheet)
    auto_fit_columns(sheet)


def style_summary_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            if cell.row in (1, 2):
                cell.fill = header_fill
                cell.font = Font(bold=True)
            elif cell.row >= 3 and cell.column in (3, 5, 7, 8, 9, 10, 11, 12, 13):
                cell.number_format = "#,##0.00"
            elif cell.row >= 3 and cell.column in (4, 6):
                cell.number_format = "0"

    for row_num in (sheet.max_row - 1, sheet.max_row):
        for col_num in range(1, sheet.max_column + 1):
            sheet.cell(row=row_num, column=col_num).font = Font(bold=True)


def process_pdf(pdf_file):
    print("Membaca summary dari isi PDF...")
    df_summary = build_summary_from_text(pdf_file)

    print("Mengambil transaksi asli dari isi PDF...")
    df_merged = parse_transactions_from_text(pdf_file)

    print("Finalisasi DB/CR dan saldo...")
    df_final = finalize_transactions(df_merged)
    report_extraction_balance(df_final, df_summary, pdf_file)

    return df_final, df_summary


def export_to_excel(df_final, df_summary, output_file, pdf_file=None):
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        if pdf_file is None:
            dataframe_for_excel(df_final).to_excel(writer, sheet_name="Transaksi", index=False)
            style_monthly_sheet(writer.sheets["Transaksi"])
        else:
            write_monthly_transaction_sheet(writer, df_final, df_summary, pdf_file, "Transaksi")

        dataframe_for_excel(df_summary).to_excel(writer, sheet_name="SummaryRaw", index=False)
        auto_fit_columns(writer.sheets["SummaryRaw"])


def export_year_workbook(extracted_files, output_file):
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        used_names = set()
        write_summary_sheet(writer, [item["summary_row"] for item in extracted_files])

        for item in extracted_files:
            sheet_name = unique_sheet_name(item["sheet_name"], used_names)
            write_monthly_transaction_sheet(
                writer,
                item["transactions"],
                item["summary"],
                item["pdf_file"],
                sheet_name,
            )


def discover_pdf_groups(input_path):
    input_path = Path(input_path)
    pdf_files = [input_path] if input_path.is_file() else sorted(input_path.rglob("*.pdf"))
    groups = {}

    for pdf_file in pdf_files:
        metadata = extract_pdf_metadata(pdf_file)
        key = (metadata["account"], metadata["year"])
        if key not in groups:
            groups[key] = {
                "account": metadata["account"],
                "year": metadata["year"],
                "folder": metadata["year"],
                "files": [],
            }
        groups[key]["files"].append({"path": pdf_file, "metadata": metadata})

    return groups


def ensure_xlsx_name(name):
    name = clean_text(name) or "{account}_{year}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def output_name_for_group(group_info, output_name_template=None, force_unique=False):
    account = group_info["account"]
    year = group_info["year"]
    folder_name = group_info.get("folder", year)

    if not output_name_template:
        return f"{account}_{year}.xlsx"

    output_name = ensure_xlsx_name(output_name_template)
    has_year_or_folder = any(token in output_name for token in ("{year}", "{folder}"))
    try:
        output_name = output_name.format(account=account, year=year, folder=folder_name)
    except (KeyError, ValueError) as exc:
        raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc

    if force_unique and not has_year_or_folder:
        output_path = Path(output_name)
        output_name = f"{output_path.stem}_{year}{output_path.suffix}"

    return output_name


def run_single_file(pdf_file, output_file):
    df_final, df_summary = process_pdf(pdf_file)
    print("Menulis ke Excel...")
    export_to_excel(df_final, df_summary, output_file, pdf_file)
    print(f"Data successfully written to {output_file}")


def run_folder(input_path, output_dir, output_name_template=None, preserve_relative_folders=True):
    groups = discover_pdf_groups(input_path)
    written_files = []

    if not groups:
        raise FileNotFoundError(f"No PDF files found in folder: {input_path}")

    for _, group_info in sorted(groups.items(), key=lambda item: (item[0][1], item[0][0])):
        file_items = group_info["files"]
        print(
            f"\nMemproses rekening {group_info['account']} tahun {group_info['year']} "
            f"({len(file_items)} file PDF)"
        )
        extracted_files = []

        for file_item in sorted(file_items, key=lambda item: (item["metadata"]["month_order"], item["path"].name)):
            pdf_file = file_item["path"]
            metadata = file_item["metadata"]
            print(f"Memproses file: {pdf_file.name}")
            df_final, df_summary = process_pdf(pdf_file)

            extracted_files.append({
                "pdf_file": pdf_file,
                "transactions": df_final,
                "summary": df_summary,
                "summary_row": build_month_summary(df_final, df_summary, pdf_file, metadata),
                "sheet_name": metadata["month_label"].replace("-", "_"),
                "month_order": metadata["month_order"],
                "year": metadata["year"],
            })

        relative_folder = Path(group_info["year"]) if preserve_relative_folders else Path()
        output_folder = output_dir / relative_folder
        output_file = output_folder / output_name_for_group(
            group_info,
            output_name_template,
            force_unique=len(groups) > 1,
        )
        export_year_workbook(extracted_files, output_file)
        written_files.append(output_file)
        print(f"Workbook selesai dibuat: {output_file}")

    return written_files


def choose_batch_options_with_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except ImportError:
        return None, None, None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    messagebox.showinfo(
        "Bank Jakarta PDF to Excel",
        "Choose the Bank Jakarta folder that contains your monthly PDF files. Subfolders will be included.",
        parent=root,
    )
    input_dir = filedialog.askdirectory(
        title="Choose Bank Jakarta PDF folder",
        initialdir=str(BANK_JAKARTA_DIR.resolve()) if BANK_JAKARTA_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not input_dir:
        root.destroy()
        return None, None, None

    messagebox.showinfo(
        "Bank Jakarta PDF to Excel",
        "Choose the folder where the Excel output should be saved.",
        parent=root,
    )
    output_dir = filedialog.askdirectory(
        title="Choose Excel output folder",
        initialdir=str(EXCEL_DIR.resolve()) if EXCEL_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not output_dir:
        root.destroy()
        return None, None, None

    output_name = simpledialog.askstring(
        "Bank Jakarta PDF to Excel",
        "Enter Excel filename pattern:\nUse {account}, {year}, or {folder} if needed.",
        initialvalue="{account}_{year}.xlsx",
        parent=root,
    )
    if output_name is None:
        root.destroy()
        return None, None, None

    root.destroy()
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name)


def choose_batch_options_with_console():
    input_dir = input("Choose input folder path: ").strip()
    output_dir = input("Choose output folder path: ").strip()
    output_name = input("Excel filename pattern [{account}_{year}.xlsx]: ").strip()
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name or "{account}_{year}.xlsx")


def show_gui_result(message, is_error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
    except ImportError:
        return

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    if is_error:
        messagebox.showerror("Bank Jakarta PDF to Excel", message, parent=root)
    else:
        messagebox.showinfo("Bank Jakarta PDF to Excel", message, parent=root)

    root.destroy()


def parse_args():
    parser = argparse.ArgumentParser(description="Convert Bank Jakarta / DKI mutasi PDF files to Excel.")
    parser.add_argument(
        "input",
        nargs="?",
        default=None,
        help="PDF filename without extension, PDF file path, or folder path. No arguments opens folder dialogs.",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        default=str(EXCEL_DIR),
        help="Output folder for generated Excel files. Default: excel_file/Bank Jakarta",
    )
    parser.add_argument(
        "--output-name",
        default=None,
        help="Workbook filename pattern. Supports {account}, {year}, {folder}. Default: {account}_{year}.xlsx",
    )
    parser.add_argument(
        "--flat-output",
        action="store_true",
        help="Save all yearly workbooks directly in output-dir instead of year subfolders.",
    )
    return parser.parse_args()


def resolve_input_path(input_value):
    input_path = Path(input_value)

    if input_path.exists():
        return input_path

    candidate_pdf = Path("pdf_file") / f"{input_path}.pdf"
    candidate_folder = Path("pdf_file") / input_path
    if candidate_pdf.exists():
        return candidate_pdf
    if candidate_folder.exists():
        return candidate_folder

    raise FileNotFoundError(f"Input not found: {input_path}")


def main():
    args = parse_args()
    use_gui = len(sys.argv) == 1

    if use_gui:
        selected = choose_batch_options_with_gui()
        if selected == (None, None, None):
            selected = choose_batch_options_with_console()
        input_path, excel_dir, output_name_template = selected
    else:
        input_path = resolve_input_path(args.input or BANK_JAKARTA_DIR)
        excel_dir = Path(args.output_dir)
        output_name_template = args.output_name

    excel_dir.mkdir(parents=True, exist_ok=True)

    if input_path.is_file():
        output_name = ensure_xlsx_name(output_name_template or f"{input_path.stem}.xlsx")
        df_final, df_summary = process_pdf(input_path)
        try:
            output_name = output_name.format(
                account=extract_account_number(df_summary, input_path),
                year=extract_month_year(input_path)[2],
                folder=input_path.parent.name,
            )
        except (KeyError, ValueError) as exc:
            raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc
        print("Menulis ke Excel...")
        export_to_excel(df_final, df_summary, excel_dir / output_name, input_path)
        print(f"Data successfully written to {excel_dir / output_name}")
    else:
        written_files = run_folder(
            input_path,
            excel_dir,
            output_name_template,
            preserve_relative_folders=not getattr(args, "flat_output", False),
        )
        print("\nSemua proses selesai.")
        for output_file in written_files:
            print(f"- {output_file}")

        if use_gui:
            show_gui_result("Done.\n\nSaved files:\n" + "\n".join(str(path) for path in written_files))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}")
        if len(sys.argv) == 1:
            show_gui_result(str(exc), is_error=True)
            sys.exit(1)
        raise
