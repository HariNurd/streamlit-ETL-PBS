from services.pdf_statement_adapter import extract_frames, summary_metrics
import argparse
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path

import fitz
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MANDIRI_DIR = Path("pdf_file") / "Mandiri"
EXCEL_DIR = Path("excel_file") / "Mandiri"

MONTH_ALIASES = {
    "JAN": ("Jan", 1), "JANUARY": ("Jan", 1), "JANUARI": ("Jan", 1),
    "FEB": ("Feb", 2), "FEBRUARY": ("Feb", 2), "FEBRUARI": ("Feb", 2),
    "MAR": ("Mar", 3), "MARCH": ("Mar", 3), "MARET": ("Mar", 3),
    "APR": ("Apr", 4), "APRIL": ("Apr", 4),
    "MAY": ("May", 5), "MEI": ("May", 5),
    "JUN": ("Jun", 6), "JUNE": ("Jun", 6), "JUNI": ("Jun", 6),
    "JUL": ("Jul", 7), "JULY": ("Jul", 7), "JULI": ("Jul", 7),
    "AUG": ("Aug", 8), "AGU": ("Aug", 8), "AGUSTUS": ("Aug", 8), "AUGUST": ("Aug", 8),
    "SEP": ("Sep", 9), "SEPT": ("Sep", 9), "SEPTEMBER": ("Sep", 9),
    "OCT": ("Oct", 10), "OKT": ("Oct", 10), "OCTOBER": ("Oct", 10), "OKTOBER": ("Oct", 10),
    "NOV": ("Nov", 11), "NOVEMBER": ("Nov", 11),
    "DEC": ("Dec", 12), "DES": ("Dec", 12), "DECEMBER": ("Dec", 12), "DESEMBER": ("Dec", 12),
}

MONTH_NAMES = {order: name for name, order in [value for value in MONTH_ALIASES.values()]}
SUMMARY_COLUMNS = [
    "No", "Bulan", "Mutasi Debet Nominal (Rp)", "Mutasi Debet Frek",
    "Mutasi Kredit Nominal (Rp)", "Mutasi Kredit Frek", "Saldo (Rp)",
    "Saldo Awal (Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
]


def clean_text(value):
    if pd.isna(value):
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def append_text(base, extra):
    base = clean_text(base)
    extra = clean_text(extra)
    if not extra:
        return base
    if not base:
        return extra
    return f"{base} {extra}"


def parse_number(value):
    text = clean_text(value).upper()
    if text in ("", "NAN", "-"):
        return pd.NA
    text = text.replace(",", "").replace("+", "").strip()
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return pd.NA


def excel_number(value):
    if pd.isna(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def dataframe_for_excel(df):
    result = df.copy()
    for col in result.columns:
        result[col] = result[col].apply(excel_number)
    return result


def is_money_line(text):
    return bool(re.fullmatch(r"\d[\d,]*\.\d{2}", clean_text(text)))


def is_date_start(text):
    return bool(re.match(r"^\d{2}/\d{2}/\d{4}\s+\d{2}:(?:\d{2}(?::\d{0,2})?)?$", clean_text(text)))


def parse_posting_datetime(first_line, second_line=""):
    text = clean_text(first_line)
    tail = clean_text(second_line)
    if text.endswith(":") and tail:
        text = text + tail
    elif tail:
        text = f"{text}:{tail}"

    for fmt in ("%d/%m/%Y %H:%M:%S",):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


@lru_cache(maxsize=None)
def read_pdf_text_for_metadata(pdf_file):
    pdf_file = Path(pdf_file)
    text_parts = []
    try:
        with fitz.open(pdf_file) as doc:
            for page in doc:
                text_parts.append(page.get_text("text"))
    except Exception:
        return ""
    return "\n".join(text_parts)


def pdf_text_lines(pdf_file):
    text = read_pdf_text_for_metadata(Path(pdf_file))
    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def normalize_metadata_text(text):
    return re.sub(r"\s+", " ", clean_text(text)).upper()


def extract_account_from_text(text):
    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]
    for idx, line in enumerate(lines):
        if line.upper() == "ACCOUNT NO." and idx + 1 < len(lines):
            match = re.search(r"\d{6,}", lines[idx + 1])
            if match:
                return match.group(0)

    normalized = normalize_metadata_text(text)
    match = re.search(r"ACCOUNT\s+NO\.?\s*(\d{6,})", normalized)
    if match:
        return match.group(1)

    matches = re.findall(r"\b\d{8,}\b", normalized)
    return matches[0] if matches else None


def parse_month_name(month_text):
    return MONTH_ALIASES.get(month_text.upper())


def extract_period_from_text(text):
    normalized = normalize_metadata_text(text)
    month_pattern = "|".join(sorted(MONTH_ALIASES, key=len, reverse=True))
    match = re.search(
        rf"PERIOD\s+(\d{{1,2}})\s+({month_pattern})\s+(20\d{{2}})\s*-\s*"
        rf"(\d{{1,2}})\s+({month_pattern})\s+(20\d{{2}})",
        normalized,
    )
    if not match:
        return None

    _, start_month_key, start_year, _, end_month_key, end_year = match.groups()
    start_month_name, start_month_order = MONTH_ALIASES[start_month_key]
    end_month_name, end_month_order = MONTH_ALIASES[end_month_key]
    return {
        "month_label": f"{start_month_name}-{start_year[-2:]}",
        "month_order": start_month_order,
        "year": start_year,
        "start_year": start_year,
        "start_month_order": start_month_order,
        "end_year": end_year,
        "end_month_order": end_month_order,
        "period_label": (
            f"{start_month_name}-{start_year[-2:]}"
            if start_month_order == end_month_order and start_year == end_year
            else f"{start_month_name}-{start_year[-2:]} to {end_month_name}-{end_year[-2:]}"
        ),
    }


def extract_period_from_filename(pdf_file):
    stem = pdf_file.stem.upper()
    year_match = re.search(r"(20\d{2}|\b\d{2}\b)", stem)
    year = None
    if year_match:
        token = year_match.group(1)
        year = f"20{token}" if len(token) == 2 else token

    month_label = None
    month_order = 99
    for token in re.split(r"[^A-Z]+", stem):
        if token in MONTH_ALIASES:
            month_name, month_order = MONTH_ALIASES[token]
            month_label = f"{month_name}-{year[-2:]}" if year else month_name
            break

    if year is None and re.fullmatch(r"20\d{2}", pdf_file.parent.name):
        year = pdf_file.parent.name
        if month_label and "-" not in month_label:
            month_label = f"{month_label}-{year[-2:]}"

    year = year or "Unknown"
    return {
        "month_label": month_label or pdf_file.stem[:20],
        "month_order": month_order,
        "year": year,
        "start_year": year,
        "start_month_order": month_order,
        "end_year": year,
        "end_month_order": month_order,
        "period_label": month_label or pdf_file.stem[:20],
    }


def extract_pdf_metadata(pdf_file):
    pdf_file = Path(pdf_file)
    text = read_pdf_text_for_metadata(pdf_file)
    period = extract_period_from_text(text) or extract_period_from_filename(pdf_file)
    account = extract_account_from_text(text)
    if account is None:
        match = re.search(r"(\d{6,})", pdf_file.stem)
        account = match.group(1) if match else pdf_file.stem.split()[0]

    period["account"] = account
    return period


def extract_month_year(pdf_file):
    metadata = extract_pdf_metadata(pdf_file)
    return metadata["month_label"], metadata["month_order"], metadata["year"]


def extract_account_number(pdf_file):
    return extract_pdf_metadata(pdf_file)["account"]


def summary_value(df_summary, label):
    if df_summary.empty:
        return pd.NA
    mask = df_summary["Keterangan"].str.upper().eq(label.upper())
    if not mask.any():
        return pd.NA
    return df_summary.loc[mask, "Amount"].iloc[0]


def value_after_label(header_lines, label):
    for idx, line in enumerate(header_lines):
        if line.upper() == label.upper():
            for candidate in header_lines[idx + 1:idx + 5]:
                if is_money_line(candidate) or re.fullmatch(r"\d+", candidate):
                    return candidate
    return ""


def build_summary_from_text(pdf_file):
    lines = pdf_text_lines(pdf_file)
    first_table_idx = next((idx for idx, line in enumerate(lines) if line.upper() == "POSTING DATE"), len(lines))
    header_lines = lines[:first_table_idx]

    label_map = {
        "Opening Balance": "Saldo Awal",
        "Total Amount Debited": "Total Pengeluaran",
        "Total Amount Credited": "Total Pemasukan",
        "Closing Balance": "Saldo Akhir",
        "No. of Debit": "Frek Pengeluaran",
        "No. of Credit": "Frek Pemasukan",
    }
    rows = []
    for pdf_label, output_label in label_map.items():
        value = value_after_label(header_lines, pdf_label)
        parsed = parse_number(value)
        rows.append({
            "Keterangan": output_label,
            "Amount": abs(parsed) if not pd.isna(parsed) else pd.NA,
        })
    return pd.DataFrame(rows, columns=["Keterangan", "Amount"])


def is_page_noise(line):
    upper = clean_text(line).upper()
    return (
        upper == "ACCOUNT STATEMENT"
        or upper.startswith("CREATED ")
        or re.fullmatch(r"PAGE\s+\d+\s+OF\s+\d+", upper) is not None
        or upper.startswith("FOR FURTHER QUESTIONS")
        or upper in {
            "POSTING DATE", "REMARK", "REFERENCE NO.", "DEBIT", "CREDIT", "BALANCE",
            "ACCOUNT STATEMENT SUMMARY", "ACCOUNT NO.", "ACCOUNT NAME", "ALIAS",
            "PERIOD", "CURRENCY", "BRANCH", "OPENING BALANCE", "NO. OF DEBIT",
            "TOTAL AMOUNT DEBITED", "CLOSING BALANCE", "NO. OF CREDIT",
            "TOTAL AMOUNT CREDITED",
        }
    )


def transaction_blocks(pdf_file):
    lines = pdf_text_lines(pdf_file)
    start = next((idx + 6 for idx, line in enumerate(lines) if line.upper() == "POSTING DATE"), 0)
    blocks = []
    current = []

    for line in lines[start:]:
        if is_page_noise(line):
            continue
        if is_date_start(line):
            if current:
                blocks.append(current)
            current = [line]
        elif current:
            current.append(line)

    if current:
        blocks.append(current)
    return blocks


def parse_transaction_block(block):
    if len(block) < 5:
        return None

    second = block[1] if len(block) > 1 and re.fullmatch(r"\d{2}:\d{2}|\d{2}", block[1]) else ""
    posted_at = parse_posting_datetime(block[0], second)
    if posted_at is None:
        return None

    content_start = 2 if second else 1
    money_positions = [(idx, line) for idx, line in enumerate(block) if is_money_line(line)]
    if len(money_positions) < 3:
        return None

    debit_idx, debit_text = money_positions[-3]
    credit_idx, credit_text = money_positions[-2]
    balance_idx, balance_text = money_positions[-1]
    if not (debit_idx < credit_idx < balance_idx):
        return None

    description_parts = [
        line for idx, line in enumerate(block[content_start:debit_idx], start=content_start)
        if clean_text(line) != "-"
    ]
    description = clean_text(" ".join(description_parts))

    debit = parse_number(debit_text)
    credit = parse_number(credit_text)
    balance = parse_number(balance_text)
    if pd.isna(balance):
        return None

    return {
        "PostingDate": posted_at,
        "Tanggal": f"{posted_at.day:02d}/{posted_at.month:02d}",
        "Keterangan": append_text(description, f"({posted_at:%H:%M:%S})"),
        "DB": debit if not pd.isna(debit) and debit != 0 else pd.NA,
        "CR": credit if not pd.isna(credit) and credit != 0 else pd.NA,
        "Saldo": balance,
    }


def parse_transactions_from_text(pdf_file):
    records = []
    for block in transaction_blocks(pdf_file):
        record = parse_transaction_block(block)
        if record is not None:
            records.append(record)
    return pd.DataFrame(records, columns=["PostingDate", "Tanggal", "Keterangan", "DB", "CR", "Saldo"])


def opening_from_first_transaction(month_df):
    if month_df.empty:
        return pd.NA
    first = month_df.iloc[0]
    opening = first["Saldo"]
    if not pd.isna(first["DB"]):
        opening += first["DB"]
    if not pd.isna(first["CR"]):
        opening -= first["CR"]
    return opening


def reconcile_transactions_with_balance(clean_df, df_summary):
    if clean_df.empty:
        return clean_df

    result = clean_df.copy()
    last_balance = summary_value(df_summary, "Saldo Awal")
    for idx, row in result.iterrows():
        current_balance = row["Saldo"]
        if pd.isna(current_balance) or pd.isna(last_balance):
            last_balance = current_balance
            continue

        delta = current_balance - last_balance
        if delta < 0:
            result.at[idx, "DB"] = abs(delta)
            result.at[idx, "CR"] = pd.NA
        elif delta > 0:
            result.at[idx, "DB"] = pd.NA
            result.at[idx, "CR"] = delta
        last_balance = current_balance

    return result


def process_pdf(pdf_file):
    """Parse through the validated shared PDF entry point."""
    return extract_frames(pdf_file, "Mandiri")


def sum_amount(df, column):
    values = df[column].dropna()
    if values.empty:
        return Decimal("0.00")
    return sum(values, Decimal("0.00"))


def count_amount(df, column):
    return int(df[column].dropna().count())


def sum_by_description(df, amount_column, pattern):
    if df.empty:
        return pd.NA
    mask = df["Keterangan"].str.upper().str.contains(pattern, regex=True, na=False)
    values = df.loc[mask, amount_column].dropna()
    if values.empty:
        return pd.NA
    return sum(values, Decimal("0.00"))


def report_extraction_balance(clean_df, df_summary, pdf_file):
    extracted_db = sum_amount(clean_df, "DB")
    extracted_cr = sum_amount(clean_df, "CR")
    summary_db = summary_value(df_summary, "Total Pengeluaran")
    summary_cr = summary_value(df_summary, "Total Pemasukan")
    summary_db_freq = summary_value(df_summary, "Frek Pengeluaran")
    summary_cr_freq = summary_value(df_summary, "Frek Pemasukan")

    mismatches = []
    if not pd.isna(summary_db) and extracted_db != summary_db:
        mismatches.append(f"DB extracted {extracted_db} but PDF summary {summary_db}")
    if not pd.isna(summary_cr) and extracted_cr != summary_cr:
        mismatches.append(f"CR extracted {extracted_cr} but PDF summary {summary_cr}")
    if not pd.isna(summary_db_freq) and count_amount(clean_df, "DB") != int(summary_db_freq):
        mismatches.append(f"DB freq extracted {count_amount(clean_df, 'DB')} but PDF summary {int(summary_db_freq)}")
    if not pd.isna(summary_cr_freq) and count_amount(clean_df, "CR") != int(summary_cr_freq):
        mismatches.append(f"CR freq extracted {count_amount(clean_df, 'CR')} but PDF summary {int(summary_cr_freq)}")

    if mismatches:
        print(f"PERINGATAN: {Path(pdf_file).name}: " + "; ".join(mismatches))
    else:
        print(f"Validasi total cocok dengan summary PDF: {Path(pdf_file).name}")


def monthly_groups(clean_df):
    if clean_df.empty:
        return []
    result = []
    for (year, month), month_df in clean_df.groupby([clean_df["PostingDate"].dt.year, clean_df["PostingDate"].dt.month]):
        month_df = month_df.sort_values("PostingDate", kind="stable").reset_index(drop=True)
        month_name = MONTH_NAMES.get(month, f"{month:02d}")
        result.append({
            "year": str(year),
            "month_order": month,
            "month_label": f"{month_name}-{str(year)[-2:]}",
            "sheet_name": f"{month_name}_{str(year)[-2:]}",
            "transactions": month_df[["Tanggal", "Keterangan", "DB", "CR", "Saldo"]].copy(),
            "opening_balance": opening_from_first_transaction(month_df),
        })
    return sorted(result, key=lambda item: (item["year"], item["month_order"]))


def build_month_summary(month_info, pdf_file, metadata):
    clean_df = month_info["transactions"]
    return {
        "Source File": pdf_file.name,
        "Account": metadata["account"],
        "Year": month_info["year"],
        "MonthOrder": month_info["month_order"],
        "Bulan": month_info["month_label"],
        "Mutasi Debet Nominal (Rp)": sum_amount(clean_df, "DB"),
        "Mutasi Debet Frek": count_amount(clean_df, "DB"),
        "Mutasi Kredit Nominal (Rp)": sum_amount(clean_df, "CR"),
        "Mutasi Kredit Frek": count_amount(clean_df, "CR"),
        "Saldo (Rp)": clean_df["Saldo"].iloc[-1] if not clean_df.empty else pd.NA,
        "Saldo Awal (Rp)": month_info["opening_balance"],
        **summary_metrics(clean_df),
    }


def build_monthly_display_dataframe(month_info):
    clean_df = month_info["transactions"]
    opening_row = pd.DataFrame([{
        "Tanggal": f"01/{month_info['month_order']:02d}",
        "Keterangan": "SALDO AWAL",
        "DB": pd.NA,
        "CR": pd.NA,
        "Saldo": month_info["opening_balance"],
    }])
    return pd.concat([opening_row, clean_df], ignore_index=True)


def auto_fit_columns(sheet):
    for col_idx, col_cells in enumerate(sheet.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        sheet.column_dimensions[col_letter].width = min(max_length + 2, 95)


def add_monthly_sheet_totals(sheet):
    first_data_row = 2
    last_data_row = sheet.max_row
    total_row = last_data_row + 3
    freq_row = total_row + 1
    average_row = total_row + 2

    for row_num, label in {total_row: "Total", freq_row: "Frekuensi", average_row: "Rata2"}.items():
        sheet.cell(row=row_num, column=2, value=label)

    for col_num in (3, 4):
        col_letter = get_column_letter(col_num)
        sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=freq_row, column=col_num, value=f"=COUNT({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=average_row, column=col_num, value=f"=IFERROR(AVERAGE({col_letter}{first_data_row}:{col_letter}{last_data_row}),0)")


def style_monthly_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary_start = sheet.max_row - 2

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for col_letter, width in {"A": 10, "B": 95, "C": 18, "D": 18, "E": 18}.items():
        sheet.column_dimensions[col_letter].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column in (3, 4, 5):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row_num in range(summary_start, sheet.max_row + 1):
        label_cell = sheet.cell(row=row_num, column=2)
        label_cell.fill = header_fill
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_num in (3, 4):
            value_cell = sheet.cell(row=row_num, column=col_num)
            value_cell.number_format = "#,##0.00" if row_num != summary_start + 1 else "0"
            value_cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_num in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 18


def write_monthly_transaction_sheet(writer, month_info, sheet_name):
    display_df = dataframe_for_excel(build_monthly_display_dataframe(month_info))
    display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    sheet = writer.sheets[sheet_name]
    add_monthly_sheet_totals(sheet)
    style_monthly_sheet(sheet)


def style_summary_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row in (1, 2):
                cell.fill = header_fill
                cell.font = Font(bold=True)
            elif cell.row >= 3 and cell.column in (3, 5, 7, 8, 9, 10, 11, 12, 13):
                cell.number_format = "#,##0.00"
            elif cell.row >= 3 and cell.column in (4, 6):
                cell.number_format = "0"

    for row_num in (sheet.max_row - 1, sheet.max_row):
        for col_num in range(1, sheet.max_column + 1):
            sheet.cell(row=row_num, column=col_num).font = Font(bold=True)


def write_summary_sheet(writer, summary_rows):
    workbook = writer.book
    sheet = workbook.create_sheet("Summary", 0)
    writer.sheets["Summary"] = sheet

    headers_top = [
        "No", "Bulan", "Mutasi Debet", None, "Mutasi Kredit", None,
        "Saldo", "Saldo Awal", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]
    headers_bottom = [
        "No", "Bulan", "Nominal (Rp)", "Frek", "Nominal (Rp)", "Frek",
        "(Rp)", "(Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]
    for col_num, value in enumerate(headers_top, start=1):
        if value is not None:
            sheet.cell(row=1, column=col_num, value=value)
    for col_num, value in enumerate(headers_bottom, start=1):
        sheet.cell(row=2, column=col_num, value=value)

    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    sheet.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=4)
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    for col_num in range(7, 14):
        sheet.merge_cells(start_row=1, start_column=col_num, end_row=2, end_column=col_num)

    sorted_rows = sorted(summary_rows, key=lambda row: (row["Year"], row["MonthOrder"], row["Source File"]))
    for row_idx, row in enumerate(sorted_rows, start=3):
        values = [
            row_idx - 2, row["Bulan"], row["Mutasi Debet Nominal (Rp)"], row["Mutasi Debet Frek"],
            row["Mutasi Kredit Nominal (Rp)"], row["Mutasi Kredit Frek"], row["Saldo (Rp)"],
            row["Saldo Awal (Rp)"], row["Adm"], row["Pajak"], row["Bunga"], row["Saldo Min"], row["JaGir"],
        ]
        for col_num, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col_num, value=excel_number(value))

    total_row = len(sorted_rows) + 3
    average_row = total_row + 1
    sheet.cell(row=total_row, column=2, value="Total")
    sheet.cell(row=average_row, column=2, value="Rata-Rata")

    data_start = 3
    data_end = total_row - 1
    if data_end >= data_start:
        for col_num in [3, 4, 5, 6, 9, 10, 11, 12, 13]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{data_start}:{col_letter}{data_end})")
        for col_num in [3, 4, 5, 6, 7]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=average_row, column=col_num, value=f"=AVERAGE({col_letter}{data_start}:{col_letter}{data_end})")

    style_summary_sheet(sheet)
    auto_fit_columns(sheet)


def safe_sheet_name(name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean_text(name)).strip()
    return cleaned[:31] or "Sheet"


def unique_sheet_name(base_name, used_names):
    base_name = safe_sheet_name(base_name)
    candidate = base_name
    counter = 2
    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def export_year_workbook(month_items, output_file):
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        used_names = set()
        write_summary_sheet(writer, [item["summary_row"] for item in month_items])
        for item in sorted(month_items, key=lambda row: (row["year"], row["month_order"], row["source_file"].name)):
            sheet_name = unique_sheet_name(item["sheet_name"], used_names)
            write_monthly_transaction_sheet(writer, item, sheet_name)


def discover_pdf_groups_by_content(input_path):
    input_path = Path(input_path)
    pdf_files = [input_path] if input_path.is_file() else sorted(input_path.rglob("*.pdf"))
    groups = {}
    for pdf_file in pdf_files:
        metadata = extract_pdf_metadata(pdf_file)
        key = (metadata["account"], metadata["year"])
        groups.setdefault(key, {
            "account": metadata["account"],
            "year": metadata["year"],
            "folder": metadata["year"],
            "files": [],
        })
        groups[key]["files"].append({"path": pdf_file, "metadata": metadata})
    return groups


def ensure_xlsx_name(name):
    name = clean_text(name) or "{account}_{year}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def output_name_for_group(group_info, output_name_template=None, force_unique=False):
    account = group_info["account"]
    year = group_info["year"]
    folder_name = group_info.get("folder", year)
    if not output_name_template:
        return f"{account}_{year}.xlsx"

    output_name = ensure_xlsx_name(output_name_template)
    has_year_or_folder = any(token in output_name for token in ("{year}", "{folder}"))
    try:
        output_name = output_name.format(account=account, year=year, folder=folder_name)
    except (KeyError, ValueError) as exc:
        raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc
    if force_unique and not has_year_or_folder:
        output_path = Path(output_name)
        output_name = f"{output_path.stem}_{year}{output_path.suffix}"
    return output_name


def run_folder(input_path, output_dir, output_name_template=None, preserve_relative_folders=True):
    groups = discover_pdf_groups_by_content(input_path)
    written_files = []
    if not groups:
        raise FileNotFoundError(f"No PDF files found in folder: {input_path}")

    for _, group_info in sorted(groups.items(), key=lambda item: (item[0][1], item[0][0])):
        print(f"\nMemproses rekening {group_info['account']} tahun {group_info['year']} ({len(group_info['files'])} file PDF)")
        month_items = []
        for file_item in sorted(group_info["files"], key=lambda item: (item["metadata"]["month_order"], item["path"].name)):
            pdf_file = file_item["path"]
            metadata = file_item["metadata"]
            print(f"Memproses file: {pdf_file.name}")
            clean_df, df_summary = process_pdf(pdf_file)
            for month_info in monthly_groups(clean_df):
                month_info["source_file"] = pdf_file
                month_info["summary_row"] = build_month_summary(month_info, pdf_file, metadata)
                month_items.append(month_info)

        relative_folder = Path(group_info["year"]) if preserve_relative_folders else Path()
        output_file = output_dir / relative_folder / output_name_for_group(
            group_info,
            output_name_template,
            force_unique=len(groups) > 1,
        )
        export_year_workbook(month_items, output_file)
        written_files.append(output_file)
        print(f"Workbook selesai dibuat: {output_file}")

    return written_files


def export_to_excel(df_final, df_summary, output_file, pdf_file=None):
    metadata = extract_pdf_metadata(pdf_file) if pdf_file else {"account": "", "year": ""}
    month_items = []
    for month_info in monthly_groups(df_final):
        month_info["source_file"] = Path(pdf_file) if pdf_file else Path("")
        month_info["summary_row"] = build_month_summary(month_info, Path(pdf_file), metadata)
        month_items.append(month_info)
    export_year_workbook(month_items, output_file)


def convert_pdf(pdf_file, output_file):
    """
    Convert one Mandiri PDF statement into one Excel workbook.

    Kept as a small wrapper around the existing parsing/export pipeline so CLI
    behavior and workbook formatting stay the same.
    """
    pdf_file = Path(pdf_file)
    output_file = Path(output_file)
    df_final, df_summary = process_pdf(pdf_file)
    print("Menulis ke Excel...")
    output_file.parent.mkdir(parents=True, exist_ok=True)
    export_to_excel(df_final, df_summary, output_file, pdf_file)
    print(f"Data successfully written to {output_file}")
    return output_file


def run_single_file(pdf_file, output_file):
    convert_pdf(pdf_file, output_file)


def choose_batch_options_with_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except ImportError:
        return None, None, None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    messagebox.showinfo(
        "Mandiri PDF to Excel",
        "Choose the Mandiri folder that contains your PDF files. Subfolders will be included.",
        parent=root,
    )
    input_dir = filedialog.askdirectory(
        title="Choose Mandiri PDF folder",
        initialdir=str(MANDIRI_DIR.resolve()) if MANDIRI_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not input_dir:
        root.destroy()
        return None, None, None

    messagebox.showinfo("Mandiri PDF to Excel", "Choose the folder where the Excel output should be saved.", parent=root)
    output_dir = filedialog.askdirectory(
        title="Choose Excel output folder",
        initialdir=str(EXCEL_DIR.resolve()) if EXCEL_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not output_dir:
        root.destroy()
        return None, None, None

    output_name = simpledialog.askstring(
        "Mandiri PDF to Excel",
        "Enter Excel filename pattern:\nUse {account}, {year}, or {folder} if needed.",
        initialvalue="{account}_{year}.xlsx",
        parent=root,
    )
    root.destroy()
    if output_name is None:
        return None, None, None
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name)


def choose_batch_options_with_console():
    input_dir = input("Choose input folder path: ").strip()
    output_dir = input("Choose output folder path: ").strip()
    output_name = input("Excel filename pattern [{account}_{year}.xlsx]: ").strip()
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name or "{account}_{year}.xlsx")


def show_gui_result(message, is_error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
    except ImportError:
        return

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    if is_error:
        messagebox.showerror("Mandiri PDF to Excel", message, parent=root)
    else:
        messagebox.showinfo("Mandiri PDF to Excel", message, parent=root)
    root.destroy()


def parse_args():
    parser = argparse.ArgumentParser(description="Convert Bank Mandiri account statement PDF files to Excel.")
    parser.add_argument(
        "input",
        nargs="?",
        default=None,
        help="PDF filename without extension, PDF file path, or folder path. No arguments opens folder dialogs.",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        default=str(EXCEL_DIR),
        help="Output folder for generated Excel files. Default: excel_file/Mandiri",
    )
    parser.add_argument(
        "--output-name",
        default=None,
        help="Workbook filename pattern. Supports {account}, {year}, {folder}. Default: {account}_{year}.xlsx",
    )
    parser.add_argument(
        "--flat-output",
        action="store_true",
        help="Save yearly workbooks directly in output-dir instead of year subfolders.",
    )
    return parser.parse_args()


def resolve_input_path(input_value):
    input_path = Path(input_value)
    if input_path.exists():
        return input_path

    candidate_pdf = Path("pdf_file") / f"{input_path}.pdf"
    candidate_folder = Path("pdf_file") / input_path
    if candidate_pdf.exists():
        return candidate_pdf
    if candidate_folder.exists():
        return candidate_folder
    raise FileNotFoundError(f"Input not found: {input_path}")


def main():
    args = parse_args()
    use_gui = len(sys.argv) == 1
    if use_gui:
        selected = choose_batch_options_with_gui()
        if selected == (None, None, None):
            selected = choose_batch_options_with_console()
        input_path, excel_dir, output_name_template = selected
    else:
        input_path = resolve_input_path(args.input or MANDIRI_DIR)
        excel_dir = Path(args.output_dir)
        output_name_template = args.output_name

    excel_dir.mkdir(parents=True, exist_ok=True)
    if input_path.is_file():
        output_name = ensure_xlsx_name(output_name_template or f"{input_path.stem}.xlsx")
        metadata = extract_pdf_metadata(input_path)
        try:
            output_name = output_name.format(account=metadata["account"], year=metadata["year"], folder=input_path.parent.name)
        except (KeyError, ValueError) as exc:
            raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc
        run_single_file(input_path, excel_dir / output_name)
    else:
        written_files = run_folder(
            input_path,
            excel_dir,
            output_name_template,
            preserve_relative_folders=not getattr(args, "flat_output", False),
        )
        print("\nSemua proses selesai.")
        for output_file in written_files:
            print(f"- {output_file}")
        if use_gui:
            show_gui_result("Done.\n\nSaved files:\n" + "\n".join(str(path) for path in written_files))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}")
        if len(sys.argv) == 1:
            show_gui_result(str(exc), is_error=True)
            sys.exit(1)
        raise
