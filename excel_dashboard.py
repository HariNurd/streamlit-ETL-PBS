import re
from datetime import date, datetime

import pandas as pd
from openpyxl.chart import BarChart, PieChart, ScatterChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


FALLBACK_KETERANGAN = "Tidak ada Fasilitas Aktif"
IDENTITY_COLUMN = "NIK/NPWP"

SOURCE_COLUMNS = [
    "Nama",
    IDENTITY_COLUMN,
    "Bank",
    "Penggunaan",
    "Plafond (Rp)",
    "Baki Debet (Rp)",
    "Kol.",
    "Awal",
    "Jatuh Tempo",
    "Rate",
    "No Laporan",
    "Keterangan",
]

CALCULATED_COLUMNS = [
    "Utilization Ratio",
    "Remaining Days",
    "Jatuh Tempo Year",
    "Jatuh Tempo Month",
    "Maturity Bucket",
    "Estimated Annual Interest Cost",
    "Risk Flags",
]

DASHBOARD_COLUMNS = SOURCE_COLUMNS + CALCULATED_COLUMNS

MATURITY_BUCKET_ORDER = [
    "Overdue",
    "0-3 Months",
    "3-6 Months",
    "6-12 Months",
    ">12 Months",
    "Unknown",
]

MONTH_MAP = {
    "januari": "01",
    "februari": "02",
    "maret": "03",
    "april": "04",
    "mei": "05",
    "juni": "06",
    "juli": "07",
    "agustus": "08",
    "september": "09",
    "oktober": "10",
    "november": "11",
    "desember": "12",
}

DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F3F6F9"
BORDER_GREY = "B7C9D6"
WHITE = "FFFFFF"

CURRENCY_FORMAT = "#,##0"
PERCENT_FORMAT = "0.00%"
DATE_FORMAT = "dd-mm-yyyy"
INTEGER_FORMAT = "#,##0"

# Internal dataframe values stay in their original names because the calculation
# logic and chart references depend on them. These maps only control what is
# displayed in Excel.
DISPLAY_LABELS = {
    "Utilization Ratio": "Rasio Utilisasi",
    "Remaining Days": "Sisa Hari",
    "Jatuh Tempo Year": "Tahun Jatuh Tempo",
    "Jatuh Tempo Month": "Bulan Jatuh Tempo",
    "Maturity Bucket": "Bucket Jatuh Tempo",
    "Estimated Annual Interest Cost": "Estimasi Beban Bunga Tahunan",
    "Risk Flags": "Indikator Risiko",
    "Facility Count": "Jumlah Fasilitas",
    "Number of Facilities": "Jumlah Fasilitas",
    "Number of Debtors": "Jumlah Debitur",
    "Number of Banks": "Jumlah Bank",
    "Banks / Financial Institutions": "Jumlah Bank / Lembaga Keuangan",
    "Total Plafond": "Total Plafond",
    "Total Baki Debet": "Total Baki Debet",
    "Average Rate": "Rata-rata Rate",
    "Highest Rate": "Rate Tertinggi",
    "Worst Kolektibilitas": "Kolektibilitas Terburuk",
    "Worst Kol.": "Kol. Terburuk",
    "Highest Kol.": "Kol. Tertinggi",
    "Active Facilities": "Fasilitas Aktif",
    "Written-Off Facilities": "Fasilitas Hapus Buku",
    "Problem Loan Exposure": "Exposure Kredit Bermasalah",
    "Problem Facilities": "Jumlah Fasilitas Bermasalah",
    "High Rate Facilities": "Fasilitas Rate Tinggi",
    "Outstanding Problem Loans": "Kredit Bermasalah Berbaki Debet",
    "Annual Interest Cost": "Estimasi Beban Bunga Tahunan",
    "Baki Debet Due 0-3 Months": "Baki Debet Jatuh Tempo 0-3 Bulan",
    "Facilities Due 0-3 Months": "Fasilitas Jatuh Tempo 0-3 Bulan",
    "Overdue Baki Debet": "Baki Debet Lewat Jatuh Tempo",
    "Overdue Facilities": "Fasilitas Lewat Jatuh Tempo",
    "Total Baki Debet by Kol.": "Total Baki Debet berdasarkan Kolektibilitas",
    "Facility Count by Kol.": "Jumlah Fasilitas berdasarkan Kolektibilitas",
    "Total Baki Debet by Keterangan": "Total Baki Debet berdasarkan Keterangan",
    "Facility Count by Keterangan": "Jumlah Fasilitas berdasarkan Keterangan",
    "Top 10 Banks by Total Baki Debet": "10 Bank Terbesar berdasarkan Total Baki Debet",
    "Top 10 Debtors by Total Baki Debet": "10 Debitur Terbesar berdasarkan Total Baki Debet",
    "Top 10 Banks by Total Plafond": "10 Bank Terbesar berdasarkan Total Plafond",
    "Summary per Debtor": "Ringkasan per Debitur",
    "Facilities with Kol. 3, 4, or 5": "Fasilitas dengan Kol. 3, 4, atau 5",
    "Written-off Facilities": "Fasilitas Hapus Buku",
    "Facilities with Rate > 30%": "Fasilitas dengan Rate > 30%",
    "Facilities with Multiple Risk Flags": "Fasilitas dengan Beberapa Indikator Risiko",
    "Debtor Risk Ranking": "Peringkat Risiko Debitur",
    "Heatmap: Nama vs Kol. by Baki Debet": "Heatmap: Nama vs Kol. berdasarkan Baki Debet",
    "Heatmap: Bank vs Kol. by Baki Debet": "Heatmap: Bank vs Kol. berdasarkan Baki Debet",
    "Maturity Summary by Bucket": "Ringkasan Jatuh Tempo berdasarkan Bucket",
    "Maturity Summary by Year": "Ringkasan Jatuh Tempo berdasarkan Tahun",
    "Average Rate by Bank": "Rata-rata Rate berdasarkan Bank",
    "Rate Distribution": "Distribusi Rate",
    "Facilities Maturing Within the Next 3 Months": "Fasilitas Jatuh Tempo dalam 3 Bulan",
    "Top 10 Highest Rate Facilities": "10 Fasilitas dengan Rate Tertinggi",
    "Top 10 Facilities by Estimated Annual Interest Cost": "10 Fasilitas berdasarkan Estimasi Beban Bunga Tahunan",
    "High Rate and High Outstanding Facilities": "Fasilitas Rate Tinggi dan Baki Debet Tinggi",
    "Scatter Data: Rate vs Baki Debet": "Data Scatter: Rate vs Baki Debet",
    "Scatter Data: Rate vs Baki Debet by Kol.": "Data Scatter: Rate vs Baki Debet berdasarkan Kol.",
    "Written-Off Count": "Jumlah Hapus Buku",
    "High Rate Count": "Jumlah Rate Tinggi",
    "Rate Band": "Kelompok Rate",
}

CHART_LABELS = {
    "Facility Count by Kol.": "Jumlah Fasilitas berdasarkan Kol.",
    "Total Baki Debet by Kol.": "Total Baki Debet berdasarkan Kol.",
    "Top Banks by Baki Debet": "Bank Terbesar berdasarkan Baki Debet",
    "Top Debtors by Baki Debet": "Debitur Terbesar berdasarkan Baki Debet",
    "Top Banks by Plafond": "Bank Terbesar berdasarkan Plafond",
    "Total Baki Debet by Maturity Bucket": "Total Baki Debet berdasarkan Bucket Jatuh Tempo",
    "Total Baki Debet by Jatuh Tempo Year": "Total Baki Debet berdasarkan Tahun Jatuh Tempo",
    "Rate vs Baki Debet": "Rate vs Baki Debet",
}

MATURITY_BUCKET_DISPLAY = {
    "Overdue": "Lewat Jatuh Tempo",
    "0-3 Months": "0-3 Bulan",
    "3-6 Months": "3-6 Bulan",
    "6-12 Months": "6-12 Bulan",
    ">12 Months": ">12 Bulan",
    "Unknown": "Tidak Diketahui",
}

RISK_FLAG_DISPLAY = {
    "Current": "Lancar",
    "Special Mention": "Dalam Perhatian Khusus",
    "Problem Loan": "Kredit Bermasalah",
    "Outstanding Problem Loan": "Kredit Bermasalah Berbaki Debet",
    "Written-Off": "Hapus Buku",
    "Charged-Off": "Hapus Tagih",
    "High Rate": "Rate Tinggi",
    "Near Maturity": "Jatuh Tempo Dekat",
    "Overdue": "Lewat Jatuh Tempo",
}

THIN_SIDE = Side(style="thin", color=BORDER_GREY)
TABLE_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

FACILITY_COLUMNS = [
    "Nama",
    IDENTITY_COLUMN,
    "Bank",
    "Penggunaan",
    "Plafond (Rp)",
    "Baki Debet (Rp)",
    "Kol.",
    "Awal",
    "Jatuh Tempo",
    "Rate",
    "Keterangan",
    "Risk Flags",
]

MATURITY_COLUMNS = [
    "Nama",
    IDENTITY_COLUMN,
    "Bank",
    "Penggunaan",
    "Baki Debet (Rp)",
    "Kol.",
    "Jatuh Tempo",
    "Remaining Days",
    "Maturity Bucket",
    "Rate",
    "Estimated Annual Interest Cost",
    "Keterangan",
]


def _is_missing(value):
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _clean_text(value):
    if value is None or _is_missing(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\r", " ").replace("\n", " ")).strip()


def display_label(label):
    return DISPLAY_LABELS.get(str(label), str(label))


def chart_label(label):
    return CHART_LABELS.get(str(label), display_label(label))


def _display_risk_flags(value):
    flags = []
    for flag in _clean_text(value).split(","):
        flag = flag.strip()
        if flag:
            flags.append(RISK_FLAG_DISPLAY.get(flag, flag))
    return ", ".join(flags)


def _display_cell_value(value, column_name):
    column_name = str(column_name)
    if column_name == "Maturity Bucket":
        return MATURITY_BUCKET_DISPLAY.get(_clean_text(value), _clean_text(value))
    if column_name == "Risk Flags":
        return _display_risk_flags(value)
    return value


def localize_dataframe_columns(df):
    return df.rename(columns=lambda column: display_label(column))


def _localize_dataframe_for_excel(df):
    display_df = df.copy()
    for column in display_df.columns:
        display_df[column] = display_df[column].apply(lambda value: _display_cell_value(value, column))
    return localize_dataframe_columns(display_df)


def _safe_sheet_name(name):
    return re.sub(r"[\[\]\:\*\?\/\\]", "_", name)[:31]


def _safe_table_name(name):
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not cleaned or not cleaned[0].isalpha():
        cleaned = f"T_{cleaned}"
    return cleaned[:250]


def _parse_amount(value):
    if value is None or _is_missing(value):
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)

    text = _clean_text(value)
    if not text:
        return 0.0

    text = text.replace("Rp", "").replace("rp", "").replace(" ", "")
    text = re.sub(r"[^0-9,.\-]", "", text)
    if not text:
        return 0.0

    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(".", "")

    try:
        return float(text)
    except ValueError:
        return 0.0


def _parse_rate(value):
    if value is None or _is_missing(value):
        return pd.NA
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        return number / 100 if abs(number) > 1 else number

    text = _clean_text(value)
    if not text:
        return pd.NA

    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", text)
    if not match:
        return pd.NA

    try:
        number = float(match.group(0).replace(",", "."))
    except ValueError:
        return pd.NA
    return number / 100 if abs(number) > 1 else number


def _parse_kol(value):
    if value is None or _is_missing(value):
        return pd.NA
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)

    match = re.search(r"\d+", _clean_text(value))
    return int(match.group(0)) if match else pd.NA


def _normalize_indonesian_date(text):
    text = _clean_text(text)
    match = re.match(r"^(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})$", text)
    if not match:
        return text

    day, month_name, year = match.groups()
    month = MONTH_MAP.get(month_name.lower())
    if not month:
        return text
    return f"{int(day):02d}-{month}-{year}"


def _parse_date(value):
    if value is None or _is_missing(value):
        return pd.NaT
    if isinstance(value, pd.Timestamp):
        return value.date() if not pd.isna(value) else pd.NaT
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _normalize_indonesian_date(value)
    if not text:
        return pd.NaT

    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    return parsed.date() if not pd.isna(parsed) else pd.NaT


def _as_float(value, default=None):
    if value is None or _is_missing(value):
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _as_int(value, default=None):
    number = _as_float(value)
    return int(number) if number is not None else default


def add_maturity_bucket(row):
    remaining_days = _as_int(row.get("Remaining Days"))
    if remaining_days is None:
        return "Unknown"
    if remaining_days < 0:
        return "Overdue"
    if remaining_days <= 90:
        return "0-3 Months"
    if remaining_days <= 180:
        return "3-6 Months"
    if remaining_days <= 365:
        return "6-12 Months"
    return ">12 Months"


def add_risk_flags(row):
    flags = []
    kol = _as_int(row.get("Kol."))
    baki_debet = _as_float(row.get("Baki Debet (Rp)"), 0) or 0
    rate = _as_float(row.get("Rate"))
    remaining_days = _as_int(row.get("Remaining Days"))
    keterangan = _clean_text(row.get("Keterangan")).lower()

    if kol == 1:
        flags.append("Current")
    elif kol == 2:
        flags.append("Special Mention")
    elif kol is not None and kol >= 3:
        flags.append("Problem Loan")

    if kol is not None and kol >= 3 and baki_debet > 0:
        flags.append("Outstanding Problem Loan")
    if "hapus tagih" in keterangan:
        flags.append("Charged-Off")
    if "dihapusbukukan" in keterangan:
        flags.append("Written-Off")
    if rate is not None and rate > 0.30:
        flags.append("High Rate")
    if remaining_days is not None and 0 <= remaining_days <= 90:
        flags.append("Near Maturity")
    if remaining_days is not None and remaining_days < 0:
        flags.append("Overdue")

    return ", ".join(flags)


def _fallback_mask(df):
    if df.empty or "Keterangan" not in df.columns:
        return pd.Series(False, index=df.index)
    return df["Keterangan"].apply(_clean_text).str.casefold() == FALLBACK_KETERANGAN.casefold()


def prepare_dashboard_data(facilities_df):
    """Create the clean dashboard source table without SLIK subtotal rows."""
    if facilities_df is None:
        df = pd.DataFrame(columns=SOURCE_COLUMNS)
    else:
        df = facilities_df.copy()

    if IDENTITY_COLUMN not in df.columns and "NIK" in df.columns:
        df[IDENTITY_COLUMN] = df["NIK"]

    for column in SOURCE_COLUMNS:
        if column not in df.columns:
            df[column] = pd.NA
    df = df[SOURCE_COLUMNS].copy()

    if not df.empty:
        bank_text = df["Bank"].apply(_clean_text)
        df = df[bank_text.str.casefold() != "jumlah"].copy()

    if df.empty:
        return pd.DataFrame(columns=DASHBOARD_COLUMNS)

    df["Nama"] = df["Nama"].apply(_clean_text).replace("", pd.NA).ffill().fillna("")
    for column in [IDENTITY_COLUMN, "Bank", "Penggunaan", "No Laporan", "Keterangan"]:
        df[column] = df[column].apply(_clean_text)

    df = df[~_fallback_mask(df)].copy()
    if df.empty:
        return pd.DataFrame(columns=DASHBOARD_COLUMNS)

    for column in ["Plafond (Rp)", "Baki Debet (Rp)"]:
        df[column] = pd.to_numeric(df[column].apply(_parse_amount), errors="coerce").fillna(0)

    df["Kol."] = pd.to_numeric(df["Kol."].apply(_parse_kol), errors="coerce").astype("Int64")
    df["Awal"] = df["Awal"].apply(_parse_date)
    df["Jatuh Tempo"] = df["Jatuh Tempo"].apply(_parse_date)
    df["Rate"] = pd.to_numeric(df["Rate"].apply(_parse_rate), errors="coerce")

    maturity_dates = pd.to_datetime(df["Jatuh Tempo"], errors="coerce")
    today = pd.Timestamp(date.today())
    remaining_days = (maturity_dates - today).dt.days
    df["Remaining Days"] = remaining_days.astype("Int64")
    df["Jatuh Tempo Year"] = maturity_dates.dt.year.astype("Int64")
    df["Jatuh Tempo Month"] = maturity_dates.dt.month.astype("Int64")

    df["Utilization Ratio"] = 0.0
    plafond_mask = df["Plafond (Rp)"] > 0
    df.loc[plafond_mask, "Utilization Ratio"] = (
        df.loc[plafond_mask, "Baki Debet (Rp)"] / df.loc[plafond_mask, "Plafond (Rp)"]
    )

    df["Maturity Bucket"] = df.apply(add_maturity_bucket, axis=1)
    df["Estimated Annual Interest Cost"] = df["Baki Debet (Rp)"] * df["Rate"].fillna(0)
    df["Risk Flags"] = df.apply(add_risk_flags, axis=1)

    return df[DASHBOARD_COLUMNS]


def _to_excel_value(value):
    if value is None or _is_missing(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date() if value.time() == datetime.min.time() else value.to_pydatetime()
    if hasattr(value, "item"):
        try:
            return value.item()
        except (AttributeError, ValueError):
            pass
    return value


def _is_currency_column(column_name):
    name = str(column_name)
    currency_terms = [
        "Plafond",
        "Baki Debet",
        "Exposure",
        "Interest Cost",
        "Beban Bunga",
    ]
    return any(term in name for term in currency_terms)


def _is_percentage_column(column_name):
    name = str(column_name)
    return name in {
        "Rate",
        "Average Rate",
        "Avg Rate",
        "Highest Rate",
        "Utilization Ratio",
        "Rata-rata Rate",
        "Rate Tertinggi",
        "Rasio Utilisasi",
    }


def _is_date_column(column_name):
    return str(column_name) in {"Awal", "Jatuh Tempo", "Jadwal"}


def _is_integer_column(column_name):
    name = str(column_name)
    integer_terms = [
        "Count",
        "Number",
        "Facilities",
        "Kol",
        "Days",
        "Year",
        "Month",
        "Jumlah",
        "Fasilitas",
        "Sisa Hari",
        "Tahun",
        "Bulan",
    ]
    return any(term in name for term in integer_terms)


def _apply_number_format(cell, column_name):
    if _is_currency_column(column_name):
        cell.number_format = CURRENCY_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    elif _is_percentage_column(column_name):
        cell.number_format = PERCENT_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    elif _is_date_column(column_name):
        cell.number_format = DATE_FORMAT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    elif _is_integer_column(column_name):
        cell.number_format = INTEGER_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _style_header_cell(cell):
    cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
    cell.font = Font(color=WHITE, bold=True)
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_section_title(cell):
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.font = Font(color="17365D", bold=True)
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _style_data_cell(cell, column_name):
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    _apply_number_format(cell, column_name)


def _add_excel_table(sheet, table_name, first_row, first_col, last_row, last_col):
    if last_row <= first_row:
        return

    ref = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    table = Table(displayName=_safe_table_name(table_name), ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    sheet.add_table(table)


def _auto_adjust_column_widths(sheet, min_width=10, max_width=36):
    for column_idx in range(1, sheet.max_column + 1):
        max_length = 0
        for row_idx in range(1, sheet.max_row + 1):
            value = sheet.cell(row=row_idx, column=column_idx).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        sheet.column_dimensions[get_column_letter(column_idx)].width = min(
            max(max_length + 2, min_width), max_width
        )


def _create_dashboard_sheet(writer, sheet_name, title):
    workbook = writer.book
    safe_name = _safe_sheet_name(sheet_name)
    sheet = workbook.create_sheet(safe_name)
    writer.sheets[safe_name] = sheet
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)
    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
    title_cell.font = Font(color=WHITE, bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)
    subtitle_cell = sheet.cell(
        row=2,
        column=1,
        value=f"Dibuat dari data fasilitas kredit yang telah dibersihkan pada {date.today():%d-%m-%Y}",
    )
    subtitle_cell.fill = PatternFill("solid", fgColor=LIGHT_GREY)
    subtitle_cell.font = Font(color="44546A", italic=True)
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")

    for column_idx in range(1, 15):
        sheet.cell(row=1, column=column_idx).border = TABLE_BORDER
        sheet.cell(row=2, column=column_idx).border = TABLE_BORDER

    return sheet


def _format_kpi_value(cell, kind):
    if kind == "currency":
        cell.number_format = CURRENCY_FORMAT
    elif kind == "percent":
        cell.number_format = PERCENT_FORMAT
    elif kind == "number":
        cell.number_format = INTEGER_FORMAT


def _write_kpi_cards(sheet, cards, start_row=3, start_col=1, cards_per_row=5, card_width=3):
    for idx, card in enumerate(cards):
        label, value, kind = card
        row = start_row + (idx // cards_per_row) * 4
        col = start_col + (idx % cards_per_row) * (card_width + 1)
        end_col = col + card_width - 1

        sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end_col)
        sheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=end_col)

        label_cell = sheet.cell(row=row, column=col, value=display_label(label))
        label_cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        label_cell.font = Font(color=WHITE, bold=True, size=9)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        value_cell = sheet.cell(row=row + 1, column=col, value=_to_excel_value(value))
        value_cell.fill = PatternFill("solid", fgColor=LIGHT_GREY)
        value_cell.font = Font(color="17365D", bold=True, size=13)
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _format_kpi_value(value_cell, kind)

        for row_idx in range(row, row + 3):
            for col_idx in range(col, end_col + 1):
                sheet.cell(row=row_idx, column=col_idx).border = TABLE_BORDER


def _safe_dataframe(df, columns=None):
    if df is None:
        return pd.DataFrame(columns=columns or [])
    if columns:
        for column in columns:
            if column not in df.columns:
                df[column] = pd.NA
        return df[columns].copy()
    return df.copy()


def _write_dataframe(
    sheet,
    df,
    start_row,
    start_col,
    title=None,
    table_name=None,
    max_rows=None,
    section_width=None,
):
    df = _safe_dataframe(df)
    if max_rows is not None:
        df = df.head(max_rows)

    columns = [str(column) for column in df.columns]
    if not columns:
        columns = ["Message"]
        df = pd.DataFrame(columns=columns)

    title_row = start_row if title else None
    header_row = start_row + 1 if title else start_row
    first_data_row = header_row + 1
    last_col = start_col + len(columns) - 1

    if title:
        end_col = max(last_col, start_col + (section_width or len(columns)) - 1)
        sheet.merge_cells(start_row=title_row, start_column=start_col, end_row=title_row, end_column=end_col)
        _style_section_title(sheet.cell(row=title_row, column=start_col, value=display_label(title)))
        for col_idx in range(start_col, end_col + 1):
            sheet.cell(row=title_row, column=col_idx).border = TABLE_BORDER

    for col_offset, column_name in enumerate(columns):
        cell = sheet.cell(row=header_row, column=start_col + col_offset, value=display_label(column_name))
        _style_header_cell(cell)

    if df.empty:
        message_cell = sheet.cell(row=first_data_row, column=start_col, value="Tidak ada data")
        message_cell.fill = PatternFill("solid", fgColor=WHITE)
        message_cell.border = TABLE_BORDER
        message_cell.alignment = Alignment(horizontal="left", vertical="center")
        for col_idx in range(start_col + 1, last_col + 1):
            sheet.cell(row=first_data_row, column=col_idx).border = TABLE_BORDER
        return {
            "header_row": header_row,
            "first_data_row": first_data_row,
            "last_data_row": header_row,
            "last_written_row": first_data_row,
            "start_col": start_col,
            "end_col": last_col,
            "columns": columns,
            "row_count": 0,
        }

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = first_data_row + row_offset
        for col_offset, column_name in enumerate(columns):
            cell = sheet.cell(
                row=excel_row,
                column=start_col + col_offset,
                value=_to_excel_value(_display_cell_value(row[column_name], column_name)),
            )
            _style_data_cell(cell, column_name)

    last_data_row = first_data_row + len(df) - 1
    if table_name:
        _add_excel_table(sheet, table_name, header_row, start_col, last_data_row, last_col)

    return {
        "header_row": header_row,
        "first_data_row": first_data_row,
        "last_data_row": last_data_row,
        "last_written_row": last_data_row,
        "start_col": start_col,
        "end_col": last_col,
        "columns": columns,
        "row_count": len(df),
    }


def _column_index(meta, column_name):
    if column_name not in meta["columns"]:
        return None
    return meta["start_col"] + meta["columns"].index(column_name)


def _has_chart_data(meta):
    return meta["row_count"] > 0 and meta["last_data_row"] >= meta["first_data_row"]


def _add_bar_chart(
    sheet,
    meta,
    label_col,
    value_col,
    anchor,
    title,
    horizontal=True,
    width=14,
    height=7,
):
    if not _has_chart_data(meta):
        return

    label_idx = _column_index(meta, label_col)
    value_idx = _column_index(meta, value_col)
    if label_idx is None or value_idx is None:
        return

    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.style = 10
    chart.title = chart_label(title)
    chart.y_axis.title = display_label(label_col if horizontal else value_col)
    chart.x_axis.title = display_label(value_col if horizontal else label_col)
    chart.width = width
    chart.height = height
    chart.legend = None

    data = Reference(sheet, min_col=value_idx, min_row=meta["header_row"], max_row=meta["last_data_row"])
    labels = Reference(sheet, min_col=label_idx, min_row=meta["first_data_row"], max_row=meta["last_data_row"])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    sheet.add_chart(chart, anchor)


def _add_pie_chart(sheet, meta, label_col, value_col, anchor, title, width=10, height=7):
    if not _has_chart_data(meta):
        return

    label_idx = _column_index(meta, label_col)
    value_idx = _column_index(meta, value_col)
    if label_idx is None or value_idx is None:
        return

    chart = PieChart()
    chart.title = chart_label(title)
    chart.width = width
    chart.height = height
    data = Reference(sheet, min_col=value_idx, min_row=meta["header_row"], max_row=meta["last_data_row"])
    labels = Reference(sheet, min_col=label_idx, min_row=meta["first_data_row"], max_row=meta["last_data_row"])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    sheet.add_chart(chart, anchor)


def _apply_heatmap(sheet, meta):
    if not _has_chart_data(meta) or meta["end_col"] <= meta["start_col"]:
        return

    start_col = meta["start_col"] + 1
    data_range = (
        f"{get_column_letter(start_col)}{meta['first_data_row']}:"
        f"{get_column_letter(meta['end_col'])}{meta['last_data_row']}"
    )
    rule = ColorScaleRule(
        start_type="min",
        start_color=WHITE,
        mid_type="percentile",
        mid_value=50,
        mid_color=LIGHT_BLUE,
        end_type="max",
        end_color=DARK_BLUE,
    )
    sheet.conditional_formatting.add(data_range, rule)


def _written_off_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return df["Keterangan"].fillna("").astype(str).str.contains(
        r"Dihapusbukukan|Hapus\s+Tagih",
        case=False,
        na=False,
        regex=True,
    )


def _active_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return df["Keterangan"].fillna("").astype(str).str.contains("Aktif", case=False, na=False)


def _problem_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return (df["Kol."] >= 3).fillna(False)


def _high_rate_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return df["Rate"].fillna(0) > 0.30


def _outstanding_problem_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return _problem_mask(df) & (df["Baki Debet (Rp)"] > 0)


def _multiple_risk_flags_mask(df):
    if df.empty:
        return pd.Series(False, index=df.index)
    return df["Risk Flags"].fillna("").astype(str).apply(lambda value: len([x for x in value.split(",") if x.strip()]) > 1)


def _safe_mean(series):
    values = pd.to_numeric(series, errors="coerce").dropna()
    return values.mean() if not values.empty else None


def _top_sum(df, group_col, value_col, output_col, top=10):
    columns = [group_col, output_col]
    if df.empty:
        return pd.DataFrame(columns=columns)

    work = df.copy()
    work[group_col] = work[group_col].apply(lambda value: _clean_text(value) or "Unknown")
    result = (
        work.groupby(group_col, dropna=False)[value_col]
        .sum()
        .reset_index(name=output_col)
        .sort_values(output_col, ascending=False)
        .head(top)
    )
    return result[columns]


def _count_by_kol(df):
    columns = ["Kol.", "Facility Count"]
    if df.empty or df["Kol."].dropna().empty:
        return pd.DataFrame(columns=columns)

    result = df.dropna(subset=["Kol."]).groupby("Kol.").size().reset_index(name="Facility Count")
    result["Kol."] = result["Kol."].astype(int)
    return result.sort_values("Kol.")[columns]


def _sum_by_kol(df):
    columns = ["Kol.", "Total Baki Debet"]
    if df.empty or df["Kol."].dropna().empty:
        return pd.DataFrame(columns=columns)

    result = (
        df.dropna(subset=["Kol."])
        .groupby("Kol.")["Baki Debet (Rp)"]
        .sum()
        .reset_index(name="Total Baki Debet")
    )
    result["Kol."] = result["Kol."].astype(int)
    return result.sort_values("Kol.")[columns]


def _summary_per_debtor(df):
    columns = [
        "Nama",
        "Total Plafond",
        "Total Baki Debet",
        "Number of Facilities",
        "Number of Banks",
        "Worst Kolektibilitas",
        "Average Rate",
        "Utilization Ratio",
        "Active Facilities",
        "Written-Off Facilities",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    for nama, group in df.groupby("Nama", sort=False, dropna=False):
        total_plafond = group["Plafond (Rp)"].sum()
        total_baki = group["Baki Debet (Rp)"].sum()
        rows.append(
            {
                "Nama": _clean_text(nama) or "Unknown",
                "Total Plafond": total_plafond,
                "Total Baki Debet": total_baki,
                "Number of Facilities": len(group),
                "Number of Banks": group["Bank"].replace("", pd.NA).dropna().nunique(),
                "Worst Kolektibilitas": _as_int(group["Kol."].max(), ""),
                "Average Rate": _safe_mean(group["Rate"]),
                "Utilization Ratio": total_baki / total_plafond if total_plafond else 0,
                "Active Facilities": int(_active_mask(group).sum()),
                "Written-Off Facilities": int(_written_off_mask(group).sum()),
            }
        )

    return pd.DataFrame(rows, columns=columns).sort_values("Total Baki Debet", ascending=False)


def _risk_flags_for_group(group):
    flags = []
    for value in group["Risk Flags"].fillna(""):
        for flag in str(value).split(","):
            flag = flag.strip()
            if flag and flag not in flags:
                flags.append(flag)
    return ", ".join(flags)


def _debtor_risk_ranking(df):
    columns = [
        "Nama",
        "Total Baki Debet",
        "Worst Kol.",
        "Problem Loan Exposure",
        "Written-Off Count",
        "High Rate Count",
        "Risk Flags",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    for nama, group in df.groupby("Nama", sort=False, dropna=False):
        problem_group = group[_problem_mask(group)]
        rows.append(
            {
                "Nama": _clean_text(nama) or "Unknown",
                "Total Baki Debet": group["Baki Debet (Rp)"].sum(),
                "Worst Kol.": _as_int(group["Kol."].max(), ""),
                "Problem Loan Exposure": problem_group["Baki Debet (Rp)"].sum(),
                "Written-Off Count": int(_written_off_mask(group).sum()),
                "High Rate Count": int(_high_rate_mask(group).sum()),
                "Risk Flags": _risk_flags_for_group(group),
            }
        )

    return (
        pd.DataFrame(rows, columns=columns)
        .sort_values(["Worst Kol.", "Problem Loan Exposure", "Total Baki Debet"], ascending=[False, False, False])
        .reset_index(drop=True)
    )


def _keterangan_summary(df, value_col=None, output_col="Facility Count", top=10):
    columns = ["Keterangan", output_col]
    if df.empty:
        return pd.DataFrame(columns=columns)

    work = df.copy()
    work["Keterangan"] = work["Keterangan"].apply(lambda value: _clean_text(value) or "Unknown")
    if value_col:
        result = work.groupby("Keterangan")[value_col].sum().reset_index(name=output_col)
    else:
        result = work.groupby("Keterangan").size().reset_index(name=output_col)
    return result.sort_values(output_col, ascending=False).head(top)[columns]


def _heatmap_pivot(df, index_col):
    if df.empty or df["Kol."].dropna().empty:
        return pd.DataFrame(columns=[index_col])

    pivot = pd.pivot_table(
        df,
        index=index_col,
        columns="Kol.",
        values="Baki Debet (Rp)",
        aggfunc="sum",
        fill_value=0,
    )
    if pivot.empty:
        return pd.DataFrame(columns=[index_col])

    pivot.columns = [f"Kol. {int(column)}" if not _is_missing(column) else "Unknown Kol." for column in pivot.columns]
    return pivot.reset_index()


def _maturity_bucket_summary(df):
    columns = [
        "Maturity Bucket",
        "Facility Count",
        "Total Baki Debet",
        "Total Plafond",
        "Average Rate",
        "Estimated Annual Interest Cost",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    present_buckets = set(df["Maturity Bucket"].dropna().astype(str))
    for bucket in MATURITY_BUCKET_ORDER:
        if bucket == "Unknown" and bucket not in present_buckets:
            continue
        group = df[df["Maturity Bucket"] == bucket]
        rows.append(
            {
                "Maturity Bucket": bucket,
                "Facility Count": len(group),
                "Total Baki Debet": group["Baki Debet (Rp)"].sum(),
                "Total Plafond": group["Plafond (Rp)"].sum(),
                "Average Rate": _safe_mean(group["Rate"]),
                "Estimated Annual Interest Cost": group["Estimated Annual Interest Cost"].sum(),
            }
        )
    return pd.DataFrame(rows, columns=columns)


def _maturity_year_summary(df):
    columns = [
        "Jatuh Tempo Year",
        "Facility Count",
        "Total Baki Debet",
        "Average Rate",
        "Estimated Annual Interest Cost",
    ]
    if df.empty or df["Jatuh Tempo Year"].dropna().empty:
        return pd.DataFrame(columns=columns)

    rows = []
    for year, group in df.dropna(subset=["Jatuh Tempo Year"]).groupby("Jatuh Tempo Year"):
        rows.append(
            {
                "Jatuh Tempo Year": int(year),
                "Facility Count": len(group),
                "Total Baki Debet": group["Baki Debet (Rp)"].sum(),
                "Average Rate": _safe_mean(group["Rate"]),
                "Estimated Annual Interest Cost": group["Estimated Annual Interest Cost"].sum(),
            }
        )
    return pd.DataFrame(rows, columns=columns).sort_values("Jatuh Tempo Year")


def _average_rate_by_bank(df):
    columns = ["Bank", "Average Rate"]
    if df.empty:
        return pd.DataFrame(columns=columns)

    work = df[df["Bank"].apply(_clean_text) != ""].copy()
    if work.empty:
        return pd.DataFrame(columns=columns)

    result = work.groupby("Bank")["Rate"].mean().reset_index(name="Average Rate").dropna()
    return result.sort_values("Average Rate", ascending=False).head(15)[columns]


def _rate_distribution(df):
    columns = ["Rate Band", "Facility Count", "Total Baki Debet", "Average Rate"]
    bands = [
        ("0-10%", 0, 0.10),
        ("10-20%", 0.10, 0.20),
        ("20-30%", 0.20, 0.30),
        ("30-40%", 0.30, 0.40),
        (">40%", 0.40, None),
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    rates = df["Rate"].fillna(0)
    for label, lower, upper in bands:
        if upper is None:
            mask = rates > lower
        elif lower == 0:
            mask = (rates >= lower) & (rates <= upper)
        else:
            mask = (rates > lower) & (rates <= upper)
        group = df[mask]
        rows.append(
            {
                "Rate Band": label,
                "Facility Count": len(group),
                "Total Baki Debet": group["Baki Debet (Rp)"].sum(),
                "Average Rate": _safe_mean(group["Rate"]),
            }
        )
    return pd.DataFrame(rows, columns=columns)


def _style_clean_data_sheet(sheet, clean_df):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    for cell in sheet[1]:
        _style_header_cell(cell)

    for row_idx in range(2, sheet.max_row + 1):
        for col_idx, column_name in enumerate(clean_df.columns, start=1):
            _style_data_cell(sheet.cell(row=row_idx, column=col_idx), column_name)

    if sheet.max_row > 1:
        _add_excel_table(sheet, "DataCleanTable", 1, 1, sheet.max_row, sheet.max_column)

    _auto_adjust_column_widths(sheet, min_width=11, max_width=34)


def write_clean_data_sheet(writer, clean_df):
    clean_df = _safe_dataframe(clean_df, DASHBOARD_COLUMNS)
    _localize_dataframe_for_excel(clean_df).to_excel(writer, sheet_name="Data_Clean", index=False)
    _style_clean_data_sheet(writer.sheets["Data_Clean"], clean_df)


def write_dashboard_summary_sheet(writer, clean_df):
    clean_df = _safe_dataframe(clean_df, DASHBOARD_COLUMNS)
    sheet = _create_dashboard_sheet(writer, "Dashboard_Summary", "Ringkasan Eksekutif")

    total_plafond = clean_df["Plafond (Rp)"].sum() if not clean_df.empty else 0
    total_baki = clean_df["Baki Debet (Rp)"].sum() if not clean_df.empty else 0
    average_rate = _safe_mean(clean_df["Rate"]) if not clean_df.empty else None
    worst_kol = _as_int(clean_df["Kol."].max(), "") if not clean_df.empty else ""
    utilization_ratio = total_baki / total_plafond if total_plafond else 0

    cards = [
        ("Total Plafond", total_plafond, "currency"),
        ("Total Baki Debet", total_baki, "currency"),
        ("Number of Debtors", clean_df["Nama"].replace("", pd.NA).dropna().nunique(), "number"),
        ("Number of Facilities", len(clean_df), "number"),
        ("Banks / Financial Institutions", clean_df["Bank"].replace("", pd.NA).dropna().nunique(), "number"),
        ("Active Facilities", int(_active_mask(clean_df).sum()), "number"),
        ("Written-Off Facilities", int(_written_off_mask(clean_df).sum()), "number"),
        ("Worst Kolektibilitas", worst_kol, "number"),
        ("Average Rate", average_rate, "percent"),
        ("Utilization Ratio", utilization_ratio, "percent"),
    ]
    _write_kpi_cards(sheet, cards)

    kol_count = _count_by_kol(clean_df)
    top_debtors = _top_sum(clean_df, "Nama", "Baki Debet (Rp)", "Total Baki Debet", top=10)
    top_banks_baki = _top_sum(clean_df, "Bank", "Baki Debet (Rp)", "Total Baki Debet", top=10)
    top_banks_plafond = _top_sum(clean_df, "Bank", "Plafond (Rp)", "Total Plafond", top=10)
    debtor_summary = _summary_per_debtor(clean_df)

    row = 12
    kol_meta = _write_dataframe(sheet, kol_count, row, 1, "Facility Count by Kol.", "SummaryKolCount")
    _add_pie_chart(sheet, kol_meta, "Kol.", "Facility Count", "D12", "Facility Count by Kol.")

    row = max(kol_meta["last_written_row"] + 4, 28)
    bank_baki_meta = _write_dataframe(sheet, top_banks_baki, row, 1, "Top 10 Banks by Total Baki Debet", "TopBanksBaki")
    _add_bar_chart(sheet, bank_baki_meta, "Bank", "Total Baki Debet", f"D{row}", "Top Banks by Baki Debet")

    row = bank_baki_meta["last_written_row"] + 4
    debtor_baki_meta = _write_dataframe(sheet, top_debtors, row, 1, "Top 10 Debtors by Total Baki Debet", "TopDebtorsBaki")
    _add_bar_chart(sheet, debtor_baki_meta, "Nama", "Total Baki Debet", f"D{row}", "Top Debtors by Baki Debet")

    row = debtor_baki_meta["last_written_row"] + 4
    bank_plafond_meta = _write_dataframe(sheet, top_banks_plafond, row, 1, "Top 10 Banks by Total Plafond", "TopBanksPlafond")
    _add_bar_chart(sheet, bank_plafond_meta, "Bank", "Total Plafond", f"D{row}", "Top Banks by Plafond")

    row = bank_plafond_meta["last_written_row"] + 4
    _write_dataframe(sheet, debtor_summary, row, 1, "Summary per Debtor", "DebtorSummary")

    _auto_adjust_column_widths(sheet)


def write_dashboard_risk_sheet(writer, clean_df):
    clean_df = _safe_dataframe(clean_df, DASHBOARD_COLUMNS)
    sheet = _create_dashboard_sheet(writer, "Dashboard_Risk", "Analisis Risiko Kredit")

    problem_mask = _problem_mask(clean_df)
    high_rate_mask = _high_rate_mask(clean_df)
    outstanding_problem_mask = _outstanding_problem_mask(clean_df)

    cards = [
        ("Problem Loan Exposure", clean_df.loc[problem_mask, "Baki Debet (Rp)"].sum(), "currency"),
        ("Problem Facilities", int(problem_mask.sum()), "number"),
        ("Written-Off Facilities", int(_written_off_mask(clean_df).sum()), "number"),
        ("Highest Kol.", _as_int(clean_df["Kol."].max(), "") if not clean_df.empty else "", "number"),
        ("High Rate Facilities", int(high_rate_mask.sum()), "number"),
        ("Outstanding Problem Loans", int(outstanding_problem_mask.sum()), "number"),
    ]
    _write_kpi_cards(sheet, cards, cards_per_row=3)

    kol_baki = _sum_by_kol(clean_df)
    kol_count = _count_by_kol(clean_df)
    ket_baki = _keterangan_summary(clean_df, value_col="Baki Debet (Rp)", output_col="Total Baki Debet")
    ket_count = _keterangan_summary(clean_df, output_col="Facility Count")

    row = 12
    kol_baki_meta = _write_dataframe(sheet, kol_baki, row, 1, "Total Baki Debet by Kol.", "RiskBakiByKol")
    _add_bar_chart(sheet, kol_baki_meta, "Kol.", "Total Baki Debet", f"D{row}", "Total Baki Debet by Kol.", horizontal=False)

    row = kol_baki_meta["last_written_row"] + 4
    kol_count_meta = _write_dataframe(sheet, kol_count, row, 1, "Facility Count by Kol.", "RiskCountByKol")
    _add_bar_chart(sheet, kol_count_meta, "Kol.", "Facility Count", f"D{row}", "Facility Count by Kol.", horizontal=False)

    row = kol_count_meta["last_written_row"] + 4
    ket_baki_meta = _write_dataframe(sheet, ket_baki, row, 1, "Total Baki Debet by Keterangan", "RiskBakiByKeterangan")
    _add_bar_chart(sheet, ket_baki_meta, "Keterangan", "Total Baki Debet", f"D{row}", "Total Baki Debet by Keterangan")

    row = ket_baki_meta["last_written_row"] + 4
    ket_count_meta = _write_dataframe(sheet, ket_count, row, 1, "Facility Count by Keterangan", "RiskCountByKeterangan")
    _add_bar_chart(sheet, ket_count_meta, "Keterangan", "Facility Count", f"D{row}", "Facility Count by Keterangan")

    problem_facilities = clean_df.loc[problem_mask, FACILITY_COLUMNS].sort_values(
        ["Kol.", "Baki Debet (Rp)"], ascending=[False, False]
    )
    written_off = clean_df.loc[_written_off_mask(clean_df), FACILITY_COLUMNS].sort_values("Baki Debet (Rp)", ascending=False)
    high_rate = clean_df.loc[high_rate_mask, FACILITY_COLUMNS].sort_values("Rate", ascending=False)
    multi_risk = clean_df.loc[_multiple_risk_flags_mask(clean_df), FACILITY_COLUMNS].sort_values(
        ["Kol.", "Baki Debet (Rp)"], ascending=[False, False]
    )
    debtor_risk = _debtor_risk_ranking(clean_df)

    row = ket_count_meta["last_written_row"] + 4
    problem_meta = _write_dataframe(sheet, problem_facilities, row, 1, "Facilities with Kol. 3, 4, or 5", "ProblemFacilities")
    row = problem_meta["last_written_row"] + 3
    written_meta = _write_dataframe(sheet, written_off, row, 1, "Written-off Facilities", "WrittenOffFacilities")
    row = written_meta["last_written_row"] + 3
    high_rate_meta = _write_dataframe(sheet, high_rate, row, 1, "Facilities with Rate > 30%", "HighRateFacilities")
    row = high_rate_meta["last_written_row"] + 3
    multi_risk_meta = _write_dataframe(sheet, multi_risk, row, 1, "Facilities with Multiple Risk Flags", "MultiRiskFacilities")
    row = multi_risk_meta["last_written_row"] + 3
    debtor_risk_meta = _write_dataframe(sheet, debtor_risk, row, 1, "Debtor Risk Ranking", "DebtorRiskRanking")

    row = debtor_risk_meta["last_written_row"] + 4
    debtor_heatmap = _heatmap_pivot(clean_df, "Nama")
    debtor_heatmap_meta = _write_dataframe(sheet, debtor_heatmap, row, 1, "Heatmap: Nama vs Kol. by Baki Debet", "DebtorKolHeatmap")
    _apply_heatmap(sheet, debtor_heatmap_meta)

    row = debtor_heatmap_meta["last_written_row"] + 4
    bank_heatmap = _heatmap_pivot(clean_df, "Bank")
    bank_heatmap_meta = _write_dataframe(sheet, bank_heatmap, row, 1, "Heatmap: Bank vs Kol. by Baki Debet", "BankKolHeatmap")
    _apply_heatmap(sheet, bank_heatmap_meta)

    _auto_adjust_column_widths(sheet)


def _write_scatter_chart(sheet, clean_df, start_row, start_col, anchor):
    work = clean_df.dropna(subset=["Rate", "Kol."]).copy()
    work = work[work["Baki Debet (Rp)"] > 0]
    if work.empty:
        _write_dataframe(
            sheet,
            pd.DataFrame(columns=["Rate", "Baki Debet (Rp)", "Kol."]),
            start_row,
            start_col,
            "Scatter Data: Rate vs Baki Debet",
            "ScatterNoData",
        )
        return

    sheet.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 7)
    _style_section_title(
        sheet.cell(
            row=start_row,
            column=start_col,
            value=display_label("Scatter Data: Rate vs Baki Debet by Kol."),
        )
    )
    for col_idx in range(start_col, start_col + 8):
        sheet.cell(row=start_row, column=col_idx).border = TABLE_BORDER

    chart = ScatterChart()
    chart.title = chart_label("Rate vs Baki Debet")
    chart.x_axis.title = display_label("Rate")
    chart.y_axis.title = display_label("Baki Debet")
    chart.width = 14
    chart.height = 7
    chart.style = 13

    current_col = start_col
    max_written_row = start_row + 1
    for kol, group in work.groupby("Kol."):
        group = group.sort_values("Rate", ascending=True)
        rate_header = sheet.cell(row=start_row + 1, column=current_col, value=f"Kol. {int(kol)} Rate")
        baki_header = sheet.cell(row=start_row + 1, column=current_col + 1, value=f"Kol. {int(kol)} Baki Debet")
        _style_header_cell(rate_header)
        _style_header_cell(baki_header)

        for offset, (_, row) in enumerate(group.iterrows(), start=2):
            rate_cell = sheet.cell(row=start_row + offset, column=current_col, value=_to_excel_value(row["Rate"]))
            baki_cell = sheet.cell(row=start_row + offset, column=current_col + 1, value=_to_excel_value(row["Baki Debet (Rp)"]))
            _style_data_cell(rate_cell, "Rate")
            _style_data_cell(baki_cell, "Baki Debet (Rp)")
            max_written_row = max(max_written_row, start_row + offset)

        if len(group) > 0:
            xvalues = Reference(
                sheet,
                min_col=current_col,
                min_row=start_row + 2,
                max_row=start_row + len(group) + 1,
            )
            yvalues = Reference(
                sheet,
                min_col=current_col + 1,
                min_row=start_row + 2,
                max_row=start_row + len(group) + 1,
            )
            series = Series(yvalues, xvalues, title=f"Kol. {int(kol)}")
            chart.series.append(series)
        current_col += 3

    sheet.add_chart(chart, anchor)
    return max_written_row


def write_dashboard_maturity_rate_sheet(writer, clean_df):
    clean_df = _safe_dataframe(clean_df, DASHBOARD_COLUMNS)
    sheet = _create_dashboard_sheet(writer, "Dashboard_Maturity_Rate", "Analisis Jatuh Tempo & Rate")

    due_soon_mask = clean_df["Remaining Days"].between(0, 90).fillna(False) if not clean_df.empty else pd.Series(False, index=clean_df.index)
    overdue_mask = (clean_df["Remaining Days"] < 0).fillna(False) if not clean_df.empty else pd.Series(False, index=clean_df.index)
    average_rate = _safe_mean(clean_df["Rate"]) if not clean_df.empty else None
    highest_rate = clean_df["Rate"].max() if not clean_df.empty and clean_df["Rate"].notna().any() else None

    cards = [
        ("Baki Debet Due 0-3 Months", clean_df.loc[due_soon_mask, "Baki Debet (Rp)"].sum(), "currency"),
        ("Facilities Due 0-3 Months", int(due_soon_mask.sum()), "number"),
        ("Overdue Baki Debet", clean_df.loc[overdue_mask, "Baki Debet (Rp)"].sum(), "currency"),
        ("Overdue Facilities", int(overdue_mask.sum()), "number"),
        ("Average Rate", average_rate, "percent"),
        ("Highest Rate", highest_rate, "percent"),
        ("Annual Interest Cost", clean_df["Estimated Annual Interest Cost"].sum() if not clean_df.empty else 0, "currency"),
    ]
    _write_kpi_cards(sheet, cards, cards_per_row=4)

    maturity_bucket = _maturity_bucket_summary(clean_df)
    maturity_year = _maturity_year_summary(clean_df)
    avg_rate_bank = _average_rate_by_bank(clean_df)
    rate_distribution = _rate_distribution(clean_df)

    row = 12
    bucket_meta = _write_dataframe(sheet, maturity_bucket, row, 1, "Maturity Summary by Bucket", "MaturityBucketSummary")
    _add_bar_chart(
        sheet,
        bucket_meta,
        "Maturity Bucket",
        "Total Baki Debet",
        f"H{row}",
        "Total Baki Debet by Maturity Bucket",
        horizontal=False,
    )

    row = bucket_meta["last_written_row"] + 4
    year_meta = _write_dataframe(sheet, maturity_year, row, 1, "Maturity Summary by Year", "MaturityYearSummary")
    _add_bar_chart(
        sheet,
        year_meta,
        "Jatuh Tempo Year",
        "Total Baki Debet",
        f"H{row}",
        "Total Baki Debet by Jatuh Tempo Year",
        horizontal=False,
    )

    row = year_meta["last_written_row"] + 4
    avg_rate_meta = _write_dataframe(sheet, avg_rate_bank, row, 1, "Average Rate by Bank", "AverageRateByBank")
    _add_bar_chart(sheet, avg_rate_meta, "Bank", "Average Rate", f"D{row}", "Average Rate by Bank")

    row = avg_rate_meta["last_written_row"] + 4
    rate_dist_meta = _write_dataframe(sheet, rate_distribution, row, 1, "Rate Distribution", "RateDistribution")
    _add_bar_chart(
        sheet,
        rate_dist_meta,
        "Rate Band",
        "Facility Count",
        f"D{row}",
        "Rate Distribution",
        horizontal=False,
    )

    row = rate_dist_meta["last_written_row"] + 4
    maturing_soon = clean_df.loc[due_soon_mask, MATURITY_COLUMNS].sort_values("Jatuh Tempo")
    overdue = clean_df.loc[overdue_mask, MATURITY_COLUMNS].sort_values("Remaining Days")
    highest_rate_facilities = clean_df.dropna(subset=["Rate"]).sort_values("Rate", ascending=False)[MATURITY_COLUMNS].head(10)
    top_interest_cost = clean_df.sort_values("Estimated Annual Interest Cost", ascending=False)[MATURITY_COLUMNS].head(10)
    high_rate_outstanding = clean_df.loc[_high_rate_mask(clean_df) & (clean_df["Baki Debet (Rp)"] > 0), MATURITY_COLUMNS].sort_values(
        "Baki Debet (Rp)", ascending=False
    )

    soon_meta = _write_dataframe(sheet, maturing_soon, row, 1, "Facilities Maturing Within the Next 3 Months", "MaturingSoon")
    row = soon_meta["last_written_row"] + 3
    overdue_meta = _write_dataframe(sheet, overdue, row, 1, "Overdue Facilities", "OverdueFacilities")
    row = overdue_meta["last_written_row"] + 3
    highest_rate_meta = _write_dataframe(sheet, highest_rate_facilities, row, 1, "Top 10 Highest Rate Facilities", "TopHighestRate")
    row = highest_rate_meta["last_written_row"] + 3
    interest_meta = _write_dataframe(sheet, top_interest_cost, row, 1, "Top 10 Facilities by Estimated Annual Interest Cost", "TopInterestCost")
    row = interest_meta["last_written_row"] + 3
    high_outstanding_meta = _write_dataframe(
        sheet,
        high_rate_outstanding,
        row,
        1,
        "High Rate and High Outstanding Facilities",
        "HighRateOutstanding",
    )

    row = high_outstanding_meta["last_written_row"] + 4
    _write_scatter_chart(sheet, clean_df, row, 1, f"H{row}")

    _auto_adjust_column_widths(sheet)


def write_dashboard_sheets(writer, facilities_df):
    clean_df = prepare_dashboard_data(facilities_df)
    write_clean_data_sheet(writer, clean_df)
    write_dashboard_summary_sheet(writer, clean_df)
    write_dashboard_risk_sheet(writer, clean_df)
    write_dashboard_maturity_rate_sheet(writer, clean_df)
