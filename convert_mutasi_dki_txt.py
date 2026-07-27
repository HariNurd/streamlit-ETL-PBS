import argparse
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TXT_DIR = Path("txt_file")
EXCEL_DIR = Path("excel_file")

MONTH_LABELS = {
    1: "Jan",
    2: "Feb",
    3: "Mar",
    4: "Apr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Aug",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dec",
}

DATE_PATTERN = r"\d{1,2}/\d{2}/\d{2}"
MONEY_PATTERN = r"\d[\d,]*\.\d{2}(?:\s+DB)?"
TX_PATTERN = re.compile(
    rf"^\s*(?P<Tanggal>{DATE_PATTERN})\s+"
    rf"(?P<body>.*?)\s+"
    rf"(?P<Mutasi>{MONEY_PATTERN})"
    rf"(?:\s+(?P<Saldo>{MONEY_PATTERN}))?\s*$",
    flags=re.IGNORECASE,
)
ACCOUNT_PATTERN = re.compile(r"No\.\s*Rekening\s*:\s*(\d+)", flags=re.IGNORECASE)
PERIOD_PATTERN = re.compile(
    r"Periode\s+Tgl\.\s*:\s*(?P<start>\d{1,2}/\d{1,2}/\d{2})\s+To\s+(?P<end>\d{1,2}/\d{1,2}/\d{2})",
    flags=re.IGNORECASE,
)
SALDO_AWAL_PATTERN = re.compile(r"SALDO\s+AWAL", flags=re.IGNORECASE)
SALDO_AKHIR_PATTERN = re.compile(r"SALDO\s+AKHIR", flags=re.IGNORECASE)
PINDAHAN_PATTERN = re.compile(r"PINDAHAN", flags=re.IGNORECASE)

TRANSACTION_COLUMNS = [
    "Account",
    "Tanggal",
    "TanggalExcel",
    "Day",
    "MonthOrder",
    "Year",
    "Keterangan",
    "Kolom3",
    "Amount",
    "Direction",
    "DB",
    "CR",
    "Saldo",
    "SaldoReported",
    "StatementBlock",
    "SourceLine",
    "SourceSequence",
    "ValidationIssue",
]


def clean_text(value):
    if pd.isna(value):
        return ""
    text = str(value).replace("\x00", " ").replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\x1b(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])", " ", text)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]", " ", text)
    text = text.strip()
    return re.sub(r"\s+", " ", text)


def append_text(base, extra):
    base = clean_text(base)
    extra = clean_text(extra)
    if not extra:
        return base
    if not base:
        return extra
    return f"{base} {extra}"


def clean_transaction_description(value):
    text = clean_text(value)
    # Some TXT exports leak page/column artifacts into wrapped descriptions,
    # for example "BIAYA ADMINISTRASI 2 B 2 2". They are not transaction text.
    text = re.sub(r"\s+\d+\s+B\s+\d+\s+\d+\s*$", "", text, flags=re.IGNORECASE)
    return clean_text(text)


def parse_number(value):
    text = clean_text(value).upper()
    if text in ("", "NAN", "NONE"):
        return pd.NA

    is_debit_balance = "DB" in text
    text = text.replace("DB", "").replace(",", "").replace("+", "").strip()
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return pd.NA

    return -number if is_debit_balance else number


def parse_amount(value):
    text = clean_text(value).upper()
    if text in ("", "NAN", "NONE"):
        return pd.NA

    text = text.replace("DB", "").replace(",", "").replace("+", "").strip()
    try:
        return abs(Decimal(text))
    except (InvalidOperation, ValueError):
        return pd.NA


def excel_number(value):
    if pd.isna(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def dataframe_for_excel(df):
    result = df.copy()
    for col in result.columns:
        result[col] = result[col].apply(excel_number)
    return result


def parse_statement_date(value):
    match = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{2})\s*$", clean_text(value))
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    year += 2000
    return day, month, year


def statement_date_value(value):
    parsed = parse_statement_date(value)
    if not parsed:
        return None
    day, month, year = parsed
    try:
        return date(year, month, day)
    except ValueError:
        return None


def display_date(value):
    parsed = parse_statement_date(value)
    if not parsed:
        return clean_text(value)
    day, month, _ = parsed
    return f"{day:02d}/{month:02d}"


def month_label(month, year):
    return f"{MONTH_LABELS.get(month, month)}-{str(year)[-2:]}"


def sheet_label(month, year):
    return f"{MONTH_LABELS.get(month, month)}_{str(year)[-2:]}"


def split_body_keterangan_kol3(body):
    body = clean_text(body)
    if not body:
        return "", ""

    parts = body.rsplit(None, 1)
    if len(parts) == 2:
        left, right = parts
        if re.fullmatch(r"[A-Z0-9/\-]+", right, flags=re.IGNORECASE):
            return left, right

    return body, ""


def is_metadata_line(line):
    text = clean_text(line).upper()
    if not text:
        return True
    if set(text) <= {"*"}:
        return True

    metadata_keywords = [
        "NO. REKENING",
        "NAMA NASABAH",
        "KUASA",
        "ALAMAT",
        "PLAFOND",
        "PERIODE TGL.",
        "HALAMAN KE",
        "BUNGA",
        "BUNGA OD",
        "DENDA",
        "TANGGAL BUKA",
        "TANGGAL AKHIR",
        "TANGGAL CETAK",
    ]
    return any(keyword in text for keyword in metadata_keywords)


def read_text_file(txt_file):
    raw = Path(txt_file).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def opening_balance_from_first_row(row):
    saldo = row["Saldo"]
    amount = row["Amount"]
    if pd.isna(saldo) or pd.isna(amount):
        return pd.NA
    if row["Direction"] == "DB":
        return saldo + amount
    if row["Direction"] == "CR":
        return saldo - amount
    return pd.NA


def parse_txt_file(txt_file):
    transactions = []
    statement_blocks = []
    current_account = None
    current_block = None
    current_tx = None
    source_sequence = 0
    last_line_number = 0

    def finish_tx():
        nonlocal current_tx
        if current_tx is not None:
            transactions.append(current_tx)
            current_tx = None

    def finish_block(end_line):
        nonlocal current_block
        if current_block is None:
            return
        block_rows = [
            row
            for row in transactions
            if row["StatementBlock"] == current_block["statement_block"]
        ]
        if pd.isna(current_block["saldo_awal"]) and block_rows:
            current_block["saldo_awal"] = opening_balance_from_first_row(block_rows[0])
        current_block["source_line_end"] = max(end_line, current_block["source_line_start"])
        current_block["transaction_count"] = len(block_rows)
        statement_blocks.append(current_block)
        current_block = None

    lines = read_text_file(txt_file).splitlines()
    for line_number, raw_line in enumerate(lines, start=1):
        last_line_number = line_number
        line = raw_line.rstrip("\n")
        text = clean_text(line)
        if not text:
            continue

        account_match = ACCOUNT_PATTERN.search(line)
        if account_match:
            finish_tx()
            finish_block(line_number - 1)
            current_account = account_match.group(1)
            block_number = len(statement_blocks) + 1
            current_block = {
                "statement_block": f"{current_account}#{block_number}",
                "account": current_account,
                "saldo_awal": pd.NA,
                "saldo_akhir": pd.NA,
                "period_start": None,
                "period_end": None,
                "source_line_start": line_number,
                "source_line_end": line_number,
                "source_order": block_number,
                "transaction_count": 0,
            }
            continue

        period_match = PERIOD_PATTERN.search(line)
        if period_match and current_block is not None:
            current_block["period_start"] = period_match.group("start")
            current_block["period_end"] = period_match.group("end")
            continue

        if SALDO_AWAL_PATTERN.search(line) and current_block is not None:
            amounts = re.findall(r"\d[\d,]*\.\d{2}(?:\s+DB)?", line, flags=re.IGNORECASE)
            if amounts:
                current_block["saldo_awal"] = parse_number(amounts[-1])
            continue

        if SALDO_AKHIR_PATTERN.search(line) and current_block is not None:
            finish_tx()
            amounts = re.findall(r"\d[\d,]*\.\d{2}(?:\s+DB)?", line, flags=re.IGNORECASE)
            if amounts:
                current_block["saldo_akhir"] = parse_number(amounts[-1])
            continue

        if PINDAHAN_PATTERN.search(line):
            finish_tx()
            if current_block is not None:
                amounts = re.findall(
                    r"\d[\d,]*\.\d{2}(?:\s+DB)?",
                    line,
                    flags=re.IGNORECASE,
                )
                if amounts:
                    current_block["saldo_awal"] = parse_number(amounts[-1])
            continue

        tx_match = TX_PATTERN.match(line)
        if tx_match:
            finish_tx()
            source_sequence += 1

            tanggal = clean_text(tx_match.group("Tanggal"))
            parsed_date = parse_statement_date(tanggal)
            body = clean_text(tx_match.group("body"))
            mutasi_text = clean_text(tx_match.group("Mutasi"))
            saldo_text = clean_text(tx_match.group("Saldo") or "")
            keterangan, kolom3 = split_body_keterangan_kol3(body)
            keterangan = clean_transaction_description(keterangan)
            amount = parse_amount(mutasi_text)
            saldo = parse_number(saldo_text)
            direction = "DB" if "DB" in mutasi_text.upper() else "CR"

            current_tx = {
                "Account": current_account,
                "Tanggal": tanggal,
                "TanggalExcel": display_date(tanggal),
                "Day": parsed_date[0] if parsed_date else 99,
                "MonthOrder": parsed_date[1] if parsed_date else 99,
                "Year": str(parsed_date[2]) if parsed_date else "",
                "Keterangan": keterangan,
                "Kolom3": kolom3,
                "Amount": amount,
                "Direction": direction,
                "DB": amount if direction == "DB" else pd.NA,
                "CR": amount if direction == "CR" else pd.NA,
                "Saldo": saldo,
                "SaldoReported": not pd.isna(saldo),
                "StatementBlock": (
                    current_block["statement_block"]
                    if current_block is not None
                    else ""
                ),
                "SourceLine": line_number,
                "SourceSequence": source_sequence,
                "ValidationIssue": "",
            }
            continue

        if is_metadata_line(line):
            continue

        if current_tx is not None and not is_metadata_line(line):
            if not (SALDO_AWAL_PATTERN.search(line) or SALDO_AKHIR_PATTERN.search(line) or PINDAHAN_PATTERN.search(line)):
                current_tx["Keterangan"] = clean_transaction_description(
                    append_text(current_tx["Keterangan"], text)
                )

    finish_tx()
    finish_block(last_line_number)

    accounts = build_account_metadata(statement_blocks)

    df = pd.DataFrame(transactions)
    if df.empty:
        return df, accounts

    df = df[df["Account"].notna() & (df["Account"].astype(str) != "")]
    df = df[df["MonthOrder"].between(1, 12)]
    df = df[TRANSACTION_COLUMNS]
    return df.reset_index(drop=True), accounts


def build_account_metadata(statement_blocks):
    accounts = {}
    grouped_blocks = {}
    for block in statement_blocks:
        grouped_blocks.setdefault(block["account"], []).append(block)

    for account, blocks in grouped_blocks.items():
        blocks = sorted(blocks, key=lambda item: item["source_order"])
        dated_start_blocks = [
            block
            for block in blocks
            if statement_date_value(block.get("period_start")) is not None
        ]
        dated_end_blocks = [
            block
            for block in blocks
            if statement_date_value(block.get("period_end")) is not None
        ]
        blocks_with_start = [
            block
            for block in dated_start_blocks
            if not pd.isna(block.get("saldo_awal", pd.NA))
        ] or dated_start_blocks
        blocks_with_end = [
            block
            for block in dated_end_blocks
            if not pd.isna(block.get("saldo_akhir", pd.NA))
        ] or dated_end_blocks

        opening_block = (
            min(
                blocks_with_start,
                key=lambda item: (
                    statement_date_value(item.get("period_start")),
                    item["source_order"],
                ),
            )
            if blocks_with_start
            else blocks[0]
        )
        ending_block = (
            max(
                blocks_with_end,
                key=lambda item: (
                    statement_date_value(item.get("period_end")),
                    item["source_order"],
                ),
            )
            if blocks_with_end
            else blocks[-1]
        )

        accounts[account] = {
            "account": account,
            "saldo_awal": opening_block.get("saldo_awal", pd.NA),
            "saldo_akhir": ending_block.get("saldo_akhir", pd.NA),
            "period_start": opening_block.get("period_start"),
            "period_end": ending_block.get("period_end"),
            "blocks": blocks,
            "statement_block_count": len(blocks),
            "duplicates_removed": 0,
            "validation_issues": [],
        }

    return accounts


def sum_amount(df, column):
    values = [value for value in df[column].dropna()]
    return sum(values, Decimal("0.00"))


def count_amount(df, column):
    return int(df[column].dropna().shape[0])


def sum_by_description(df, amount_column, pattern):
    if df.empty:
        return Decimal("0.00")
    mask = df["Keterangan"].fillna("").str.contains(pattern, case=False, regex=True)
    values = [value for value in df.loc[mask, amount_column].dropna()]
    return sum(values, Decimal("0.00"))


def append_validation_issue(value, message):
    current = clean_text(value)
    message = clean_text(message)
    if not current:
        return message
    if message in current.split(" | "):
        return current
    return f"{current} | {message}"


def add_account_validation_issue(
    account_info,
    issue_type,
    severity,
    message,
    statement_block=None,
    source_line=None,
):
    issues = account_info.setdefault("validation_issues", [])
    issue = {
        "issue_type": issue_type,
        "severity": severity,
        "account": account_info.get("account"),
        "statement_block": statement_block,
        "source_line": source_line,
        "message": clean_text(message),
    }
    signature = tuple(issue.get(key) for key in [
        "issue_type",
        "severity",
        "statement_block",
        "source_line",
        "message",
    ])
    existing_signatures = {
        tuple(existing.get(key) for key in [
            "issue_type",
            "severity",
            "statement_block",
            "source_line",
            "message",
        ])
        for existing in issues
    }
    if signature not in existing_signatures:
        issues.append(issue)


def block_metadata_by_id(account_info):
    return {
        block.get("statement_block"): block
        for block in account_info.get("blocks", [])
    }


def reconcile_statement_block(block_df, block_info, account_info):
    result = block_df.copy()
    if "ValidationIssue" not in result.columns:
        result["ValidationIssue"] = ""
    if "SaldoReported" not in result.columns:
        result["SaldoReported"] = result["Saldo"].apply(lambda value: not pd.isna(value))

    sort_columns = [
        column
        for column in ["SourceSequence", "Year", "MonthOrder", "Day"]
        if column in result.columns
    ]
    if sort_columns:
        result = result.sort_values(sort_columns, kind="stable")

    last_balance = block_info.get("saldo_awal", pd.NA)
    statement_block = block_info.get("statement_block")

    for idx, row in result.iterrows():
        saldo = row.get("Saldo", pd.NA)
        db_value = row.get("DB", pd.NA)
        cr_value = row.get("CR", pd.NA)
        source_line = row.get("SourceLine")

        if pd.isna(db_value) and pd.isna(cr_value):
            amount = row.get("Amount", pd.NA)
            direction = clean_text(row.get("Direction")).upper()
            if not pd.isna(amount) and direction == "DB":
                db_value = amount
                result.at[idx, "DB"] = amount
            elif not pd.isna(amount) and direction == "CR":
                cr_value = amount
                result.at[idx, "CR"] = amount

        if not pd.isna(db_value) and not pd.isna(cr_value):
            message = "Transaksi memiliki nilai DB dan CR sekaligus; nilai sumber dipertahankan."
            result.at[idx, "ValidationIssue"] = append_validation_issue(
                result.at[idx, "ValidationIssue"],
                message,
            )
            add_account_validation_issue(
                account_info,
                "db_cr_both_present",
                "WARNING",
                message,
                statement_block=statement_block,
                source_line=source_line,
            )

        if pd.isna(saldo):
            if not pd.isna(last_balance) and not pd.isna(db_value) and pd.isna(cr_value):
                saldo = last_balance - db_value
                result.at[idx, "Saldo"] = saldo
            elif not pd.isna(last_balance) and pd.isna(db_value) and not pd.isna(cr_value):
                saldo = last_balance + cr_value
                result.at[idx, "Saldo"] = saldo

            if not pd.isna(saldo):
                last_balance = saldo
            continue

        if not pd.isna(last_balance):
            expected_balance = None
            if not pd.isna(db_value) and pd.isna(cr_value):
                expected_balance = last_balance - db_value
            elif pd.isna(db_value) and not pd.isna(cr_value):
                expected_balance = last_balance + cr_value

            if expected_balance is not None and expected_balance != saldo:
                message = (
                    f"Saldo transaksi {saldo} tidak sesuai dengan saldo sebelumnya "
                    f"{last_balance} dan mutasi sumber; saldo seharusnya {expected_balance}. "
                    "Nilai DB/CR sumber tidak diubah."
                )
                result.at[idx, "ValidationIssue"] = append_validation_issue(
                    result.at[idx, "ValidationIssue"],
                    message,
                )
                add_account_validation_issue(
                    account_info,
                    "balance_mismatch",
                    "WARNING",
                    message,
                    statement_block=statement_block,
                    source_line=source_line,
                )

        last_balance = saldo

    ending_balance = block_info.get("saldo_akhir", pd.NA)
    if (
        not pd.isna(ending_balance)
        and not pd.isna(last_balance)
        and ending_balance != last_balance
    ):
        message = (
            f"Saldo akhir blok {ending_balance} tidak sama dengan saldo transaksi "
            f"terakhir {last_balance}."
        )
        add_account_validation_issue(
            account_info,
            "ending_balance_mismatch",
            "WARNING",
            message,
            statement_block=statement_block,
            source_line=block_info.get("source_line_end"),
        )

    return result


def transaction_fingerprint(row, include_balance=True):
    fingerprint = (
        clean_text(row.get("Account")),
        clean_text(row.get("Year")),
        int(row.get("MonthOrder", 0) or 0),
        int(row.get("Day", 0) or 0),
        clean_text(row.get("Keterangan")).upper(),
        clean_text(row.get("Kolom3")).upper(),
        clean_text(row.get("Direction")).upper(),
        stable_signature_value(row.get("Amount")),
    )
    if include_balance:
        return fingerprint + (stable_signature_value(row.get("Saldo")),)
    return fingerprint


def block_preference_key(block):
    start = statement_date_value(block.get("period_start"))
    end = statement_date_value(block.get("period_end"))
    span = (end - start).days if start is not None and end is not None else -1
    start_order = start.toordinal() if start is not None else date.max.toordinal()
    return (-span, start_order, int(block.get("source_order", 0) or 0))


def add_overlap_conflict_issues(account_df, account_info):
    core_groups = {}
    for idx, row in account_df.iterrows():
        core_groups.setdefault(transaction_fingerprint(row, include_balance=False), []).append(idx)

    for indices in core_groups.values():
        block_sets = {}
        for idx in indices:
            row = account_df.loc[idx]
            block_id = row.get("StatementBlock")
            block_sets.setdefault(block_id, set()).add(
                transaction_fingerprint(row, include_balance=True)
            )
        if len(block_sets) < 2:
            continue

        strict_sets = list(block_sets.values())
        has_covering_block = any(
            all(other_set.issubset(candidate_set) for other_set in strict_sets)
            for candidate_set in strict_sets
        )
        if has_covering_block:
            continue

        block_names = ", ".join(clean_text(block_id) for block_id in block_sets)
        message = (
            "Transaksi yang tampak sama memiliki saldo berbeda pada blok yang "
            f"tumpang tindih ({block_names}); semua versi dipertahankan untuk ditinjau."
        )
        add_account_validation_issue(
            account_info,
            "overlap_conflict",
            "WARNING",
            message,
        )
        for idx in indices:
            account_df.at[idx, "ValidationIssue"] = append_validation_issue(
                account_df.at[idx, "ValidationIssue"],
                message,
            )


def deduplicate_overlapping_blocks(account_df, account_info):
    result = account_df.copy().reset_index(drop=True)
    if result.empty or "StatementBlock" not in result.columns:
        return result

    add_overlap_conflict_issues(result, account_info)
    block_map = block_metadata_by_id(account_info)
    fingerprint_groups = {}
    for idx, row in result.iterrows():
        fingerprint_groups.setdefault(
            transaction_fingerprint(row, include_balance=True),
            [],
        ).append(idx)

    drop_indices = set()
    for indices in fingerprint_groups.values():
        rows_by_block = {}
        for idx in indices:
            block_id = result.at[idx, "StatementBlock"]
            rows_by_block.setdefault(block_id, []).append(idx)
        if len(rows_by_block) < 2:
            continue

        keep_count = max(len(block_rows) for block_rows in rows_by_block.values())
        preferred_indices = sorted(
            indices,
            key=lambda idx: (
                block_preference_key(block_map.get(result.at[idx, "StatementBlock"], {})),
                int(result.at[idx, "SourceSequence"] or 0),
            ),
        )
        drop_indices.update(preferred_indices[keep_count:])

    if drop_indices:
        account_info["duplicates_removed"] = len(drop_indices)
        result = result.drop(index=sorted(drop_indices)).reset_index(drop=True)
        add_account_validation_issue(
            account_info,
            "exact_cross_block_duplicates",
            "INFO",
            (
                f"{len(drop_indices)} transaksi duplikat identik antarblok "
                "dihapus dengan mempertahankan jumlah kemunculan maksimum dalam satu blok."
            ),
        )

    return result


def reconcile_transactions_with_balance(account_df, account_info):
    result = account_df.copy()
    account_info.setdefault("validation_issues", [])
    account_info["duplicates_removed"] = 0

    block_map = block_metadata_by_id(account_info)
    if "StatementBlock" in result.columns and result["StatementBlock"].nunique() > 1:
        add_account_validation_issue(
            account_info,
            "multiple_statement_blocks",
            "INFO",
            (
                f"{result['StatementBlock'].nunique()} blok rekening terdeteksi; "
                "setiap blok direkonsiliasi secara terpisah sebelum data digabungkan."
            ),
        )

    reconciled_frames = []
    if "StatementBlock" in result.columns and result["StatementBlock"].notna().any():
        for block_id, block_df in result.groupby("StatementBlock", sort=False):
            block_info = block_map.get(
                block_id,
                {
                    "statement_block": block_id,
                    "saldo_awal": account_info.get("saldo_awal", pd.NA),
                    "saldo_akhir": account_info.get("saldo_akhir", pd.NA),
                },
            )
            reconciled_frames.append(
                reconcile_statement_block(block_df, block_info, account_info)
            )
    else:
        fallback_block = {
            "statement_block": None,
            "saldo_awal": account_info.get("saldo_awal", pd.NA),
            "saldo_akhir": account_info.get("saldo_akhir", pd.NA),
        }
        reconciled_frames.append(
            reconcile_statement_block(result, fallback_block, account_info)
        )

    result = pd.concat(reconciled_frames, ignore_index=True)
    result = deduplicate_overlapping_blocks(result, account_info)
    sort_columns = [
        column
        for column in ["Year", "MonthOrder", "Day", "SourceSequence"]
        if column in result.columns
    ]
    if sort_columns:
        result = result.sort_values(sort_columns, kind="stable").reset_index(drop=True)
    return result


def monthly_groups(account_df, account_info):
    result = []
    sort_columns = [
        column
        for column in ["Year", "MonthOrder", "Day", "SourceSequence"]
        if column in account_df.columns
    ]
    ordered = account_df.sort_values(sort_columns, kind="stable").reset_index(drop=True)

    for (year, month), month_df in ordered.groupby(["Year", "MonthOrder"], sort=True):
        month_df = month_df.copy()
        month_df = month_df[["TanggalExcel", "Keterangan", "DB", "CR", "Saldo"]]
        month_df = month_df.rename(columns={"TanggalExcel": "Tanggal"})

        first_row = ordered[(ordered["Year"] == year) & (ordered["MonthOrder"] == month)].iloc[0]
        opening = opening_balance_from_first_row(first_row)
        if pd.isna(opening) and str(year) == str(account_info.get("period_start", ""))[-2:]:
            opening = account_info.get("saldo_awal", pd.NA)

        result.append({
            "year": str(year),
            "month_order": int(month),
            "month_label": month_label(int(month), str(year)),
            "sheet_name": sheet_label(int(month), str(year)),
            "opening": opening,
            "transactions": month_df.reset_index(drop=True),
        })

    return result


def build_month_summary(account, month_info, source_file):
    df = month_info["transactions"]
    return {
        "Source File": source_file.name,
        "Account": account,
        "Year": month_info["year"],
        "MonthOrder": month_info["month_order"],
        "Bulan": month_info["month_label"],
        "Mutasi Debet Nominal (Rp)": sum_amount(df, "DB"),
        "Mutasi Debet Frek": count_amount(df, "DB"),
        "Mutasi Kredit Nominal (Rp)": sum_amount(df, "CR"),
        "Mutasi Kredit Frek": count_amount(df, "CR"),
        "Saldo (Rp)": df["Saldo"].dropna().iloc[-1] if not df["Saldo"].dropna().empty else pd.NA,
        "Saldo Awal (Rp)": month_info["opening"],
        "Adm": sum_by_description(
            df,
            "DB",
            r"\bADM\b|ADMIN|BIAYA\s+ADMIN|FEE|\bPROVISI\b",
        ),
        "Pajak": sum_by_description(df, "DB", r"\bPPH\b|PAJAK|TAX"),
        "Bunga": sum_by_description(df, "DB", r"BUNGA|INTEREST|\bBNG\b"),
        "Saldo Min": sum_by_description(df, "DB", r"SALDO\s+MIN"),
        "JaGir": sum_by_description(df, "CR", r"JASA\s+GIRO|JAGIR"),
    }


def report_statement_balance(account, account_df, account_info, source_file):
    saldo_awal = account_info.get("saldo_awal", pd.NA)
    saldo_akhir = account_info.get("saldo_akhir", pd.NA)
    if pd.isna(saldo_awal) or pd.isna(saldo_akhir):
        return

    db_total = sum_amount(account_df, "DB")
    cr_total = sum_amount(account_df, "CR")
    expected = saldo_awal + cr_total - db_total
    if expected != saldo_akhir:
        print(
            f"PERINGATAN: {source_file.name} rekening {account}: "
            f"saldo awal + CR - DB = {expected}, tetapi SALDO AKHIR = {saldo_akhir}"
        )


def report_account_validation_issues(account_info, source_file):
    for issue in account_info.get("validation_issues", []):
        prefix = "INFO" if issue.get("severity") == "INFO" else "PERINGATAN"
        print(
            f"{prefix}: {source_file.name} rekening {account_info.get('account')}: "
            f"{issue.get('message')}"
        )


def safe_sheet_name(name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean_text(name)).strip()
    return cleaned[:31] or "Sheet"


def unique_sheet_name(base_name, used_names):
    base_name = safe_sheet_name(base_name)
    candidate = base_name
    counter = 2
    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def build_monthly_display_dataframe(month_info):
    opening_row = pd.DataFrame([{
        "Tanggal": f"01/{month_info['month_order']:02d}",
        "Keterangan": "SALDO AWAL",
        "DB": pd.NA,
        "CR": pd.NA,
        "Saldo": month_info["opening"],
    }])
    return pd.concat([opening_row, month_info["transactions"]], ignore_index=True)


def auto_fit_columns(sheet):
    for col_idx, col_cells in enumerate(sheet.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        sheet.column_dimensions[col_letter].width = max_length + 2


def add_monthly_sheet_totals(sheet):
    first_data_row = 2
    last_data_row = sheet.max_row
    total_row = last_data_row + 3
    freq_row = total_row + 1
    average_row = total_row + 2

    for row_num, label in {total_row: "Total", freq_row: "Frekuensi", average_row: "Rata2"}.items():
        sheet.cell(row=row_num, column=2, value=label)

    for col_num in (3, 4):
        col_letter = get_column_letter(col_num)
        sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=freq_row, column=col_num, value=f"=COUNT({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=average_row, column=col_num, value=f"=IFERROR({col_letter}{total_row}/{col_letter}{freq_row},0)")


def style_monthly_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary_start = sheet.max_row - 2

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for col_letter, width in {"A": 10, "B": 95, "C": 16, "D": 16, "E": 16}.items():
        sheet.column_dimensions[col_letter].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column in (3, 4, 5):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row_num in range(summary_start, sheet.max_row + 1):
        label_cell = sheet.cell(row=row_num, column=2)
        label_cell.fill = header_fill
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_num in (3, 4):
            value_cell = sheet.cell(row=row_num, column=col_num)
            value_cell.number_format = "0" if row_num == summary_start + 1 else "#,##0.00"
            value_cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_num in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 18


def write_monthly_transaction_sheet(writer, month_info, sheet_name):
    display_df = dataframe_for_excel(build_monthly_display_dataframe(month_info))
    display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    sheet = writer.sheets[sheet_name]
    add_monthly_sheet_totals(sheet)
    style_monthly_sheet(sheet)


def style_summary_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row in (1, 2):
                cell.fill = header_fill
                cell.font = Font(bold=True)
            elif cell.row >= 3 and cell.column in (3, 5, 7, 8, 9, 10, 11, 12, 13):
                cell.number_format = "#,##0.00"
            elif cell.row >= 3 and cell.column in (4, 6):
                cell.number_format = "0"

    for row_num in (sheet.max_row - 1, sheet.max_row):
        for col_num in range(1, sheet.max_column + 1):
            sheet.cell(row=row_num, column=col_num).font = Font(bold=True)


def write_summary_sheet(writer, summary_rows):
    workbook = writer.book
    sheet = workbook.create_sheet("Summary", 0)
    writer.sheets["Summary"] = sheet

    headers_top = [
        "No", "Bulan", "Mutasi Debet", None, "Mutasi Kredit", None,
        "Saldo", "Saldo Awal", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]
    headers_bottom = [
        "No", "Bulan", "Nominal (Rp)", "Frek", "Nominal (Rp)", "Frek",
        "(Rp)", "(Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]

    for col_num, value in enumerate(headers_top, start=1):
        if value is not None:
            sheet.cell(row=1, column=col_num, value=value)
    for col_num, value in enumerate(headers_bottom, start=1):
        sheet.cell(row=2, column=col_num, value=value)

    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    sheet.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=4)
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    for col_num in range(7, 14):
        sheet.merge_cells(start_row=1, start_column=col_num, end_row=2, end_column=col_num)

    sorted_rows = sorted(summary_rows, key=lambda row: (row["Year"], row["MonthOrder"], row["Source File"]))
    for row_idx, row in enumerate(sorted_rows, start=3):
        values = [
            row_idx - 2,
            row["Bulan"],
            row["Mutasi Debet Nominal (Rp)"],
            row["Mutasi Debet Frek"],
            row["Mutasi Kredit Nominal (Rp)"],
            row["Mutasi Kredit Frek"],
            row["Saldo (Rp)"],
            row["Saldo Awal (Rp)"],
            row["Adm"],
            row["Pajak"],
            row["Bunga"],
            row["Saldo Min"],
            row["JaGir"],
        ]
        for col_num, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col_num, value=excel_number(value))

    total_row = len(sorted_rows) + 3
    average_row = total_row + 1
    sheet.cell(row=total_row, column=2, value="Total")
    sheet.cell(row=average_row, column=2, value="Rata-Rata")

    data_start = 3
    data_end = total_row - 1
    if data_end >= data_start:
        for col_num in [3, 4, 5, 6, 9, 10, 11, 12, 13]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{data_start}:{col_letter}{data_end})")
        for col_num in [3, 4, 5, 6, 7]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=average_row, column=col_num, value=f"=AVERAGE({col_letter}{data_start}:{col_letter}{data_end})")

    style_summary_sheet(sheet)
    auto_fit_columns(sheet)


def data_quality_rows(account_info):
    rows = []
    for issue in account_info.get("validation_issues", []):
        rows.append({
            "Severity": issue.get("severity"),
            "Issue": issue.get("issue_type"),
            "Account": issue.get("account"),
            "Statement Block": issue.get("statement_block"),
            "Source Line": issue.get("source_line"),
            "Message": issue.get("message"),
        })
    if not rows:
        rows.append({
            "Severity": "INFO",
            "Issue": "no_issues",
            "Account": account_info.get("account"),
            "Statement Block": None,
            "Source Line": None,
            "Message": "Tidak ada masalah validasi yang terdeteksi.",
        })
    return rows


def write_data_quality_sheet(writer, rows):
    df = pd.DataFrame(rows, columns=[
        "Severity",
        "Issue",
        "Account",
        "Statement Block",
        "Source Line",
        "Message",
    ])
    df.to_excel(writer, sheet_name="Data_Quality", index=False)
    sheet = writer.sheets["Data_Quality"]
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for col_letter, width in {
        "A": 12,
        "B": 32,
        "C": 18,
        "D": 24,
        "E": 12,
        "F": 100,
    }.items():
        sheet.column_dimensions[col_letter].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_num in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 36


def export_year_workbook(month_items, output_file, quality_rows=None):
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        used_names = set()
        write_summary_sheet(writer, [item["summary_row"] for item in month_items])
        write_data_quality_sheet(writer, quality_rows or [])
        for item in month_items:
            sheet_name = unique_sheet_name(item["sheet_name"], used_names)
            write_monthly_transaction_sheet(writer, item, sheet_name)


def stable_signature_value(value):
    if pd.isna(value):
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    return clean_text(value)


def month_item_signature(account, month_info):
    transaction_rows = []
    for row in month_info["transactions"].to_dict("records"):
        transaction_rows.append(tuple(stable_signature_value(row.get(column)) for column in [
            "Tanggal",
            "Keterangan",
            "DB",
            "CR",
            "Saldo",
        ]))

    return (
        clean_text(account),
        clean_text(month_info.get("year", "")),
        int(month_info.get("month_order", 0) or 0),
        stable_signature_value(month_info.get("opening")),
        tuple(transaction_rows),
    )


def discover_txt_files(input_path):
    input_path = Path(input_path)
    if input_path.is_file():
        return [input_path]
    return sorted(input_path.rglob("*.txt"))


def source_folder_name(input_path):
    input_path = Path(input_path)
    return input_path.stem if input_path.is_file() else input_path.name


def run_path(input_path, output_dir=EXCEL_DIR, output_name_template=None):
    input_path = Path(input_path)
    groups = {}
    seen_month_signatures = set()

    for txt_file in discover_txt_files(input_path):
        print(f"Memproses {txt_file}...")
        df_raw, accounts = parse_txt_file(txt_file)
        if df_raw.empty:
            print(f"PERINGATAN: tidak ada transaksi valid di {txt_file}")
            continue

        for account, account_df in df_raw.groupby("Account", sort=True):
            account_info = accounts.get(account, {"saldo_awal": pd.NA, "saldo_akhir": pd.NA})
            account_df = reconcile_transactions_with_balance(account_df, account_info)
            report_statement_balance(account, account_df, account_info, txt_file)
            report_account_validation_issues(account_info, txt_file)

            for month_info in monthly_groups(account_df, account_info):
                signature = month_item_signature(account, month_info)
                if signature in seen_month_signatures:
                    print(
                        "PERINGATAN: duplikat bulan dilewati: "
                        f"rekening {account} {month_info['month_label']} dari {txt_file.name}"
                    )
                    continue
                seen_month_signatures.add(signature)

                key = (account, month_info["year"])
                groups.setdefault(key, {
                    "account": account,
                    "year": month_info["year"],
                    "items": [],
                    "data_quality_rows": data_quality_rows(account_info),
                })
                month_info["summary_row"] = build_month_summary(account, month_info, txt_file)
                groups[key]["items"].append(month_info)

    folder_name = source_folder_name(input_path)
    output_files = []
    for group in groups.values():
        if output_name_template:
            output_name = output_name_template
            if not output_name.lower().endswith(".xlsx"):
                output_name += ".xlsx"
            output_name = output_name.format(account=group["account"], year=group["year"], folder=folder_name)
        else:
            output_name = f"{group['account']}_{group['year']}.xlsx"

        output_file = Path(output_dir) / folder_name / group["year"] / output_name
        export_year_workbook(
            group["items"],
            output_file,
            quality_rows=group.get("data_quality_rows"),
        )
        output_files.append(output_file)
        print(f"Workbook selesai dibuat: {output_file}")

    return output_files


def convert_dki_txt(input_path, output_dir, output_name_template=None):
    """
    Convert one TXT file or a folder of TXT files into yearly Excel workbook(s).
    """
    return run_path(input_path, output_dir, output_name_template)


def parse_args():
    parser = argparse.ArgumentParser(description="Convert Bank Jakarta / DKI mutasi TXT files to yearly Excel workbooks.")
    parser.add_argument("input", nargs="?", help="TXT file, TXT folder, or filename stem under txt_file.")
    parser.add_argument("-o", "--output-dir", default=str(EXCEL_DIR), help="Output root directory. Default: excel_file")
    parser.add_argument(
        "--output-name",
        help="Workbook filename pattern. Supports {account}, {year}, and {folder}. Default: {account}_{year}.xlsx",
    )
    return parser.parse_args()


def resolve_input_path(input_value):
    if input_value:
        candidate = Path(input_value)
        if candidate.exists():
            return candidate
        txt_candidate = TXT_DIR / input_value
        if txt_candidate.exists():
            return txt_candidate
        if not input_value.lower().endswith(".txt"):
            txt_candidate = TXT_DIR / f"{input_value}.txt"
            if txt_candidate.exists():
                return txt_candidate
        return candidate

    file_name = input("Enter the TXT filename without extension, or a folder path: ").strip()
    return resolve_input_path(file_name)


def main():
    args = parse_args()
    input_path = resolve_input_path(args.input)
    if not input_path.exists():
        raise FileNotFoundError(f"Input not found: {input_path}")

    output_files = run_path(input_path, Path(args.output_dir), args.output_name)
    if not output_files:
        print("Tidak ada workbook dibuat.")


if __name__ == "__main__":
    main()
