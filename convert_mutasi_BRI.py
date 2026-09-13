from services.pdf_statement_adapter import extract_frames, summary_metrics
import argparse
import re
import sys
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path

import fitz
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BRI_DIR = Path("pdf_file") / "BRI"
EXCEL_DIR = Path("excel_file") / "BRI"

MONTH_ALIASES = {
    "JAN": ("Jan", 1), "JANUARI": ("Jan", 1), "JANUARY": ("Jan", 1),
    "FEB": ("Feb", 2), "FEBRUARI": ("Feb", 2), "FEBRUARY": ("Feb", 2),
    "MAR": ("Mar", 3), "MARET": ("Mar", 3), "MARCH": ("Mar", 3),
    "APR": ("Apr", 4), "APRIL": ("Apr", 4),
    "MEI": ("May", 5), "MAY": ("May", 5),
    "JUN": ("Jun", 6), "JUNI": ("Jun", 6), "JUNE": ("Jun", 6),
    "JUL": ("Jul", 7), "JULI": ("Jul", 7), "JULY": ("Jul", 7),
    "AGU": ("Aug", 8), "AGS": ("Aug", 8), "AUG": ("Aug", 8), "AGUST": ("Aug", 8),
    "AGUSTUS": ("Aug", 8), "AUGUST": ("Aug", 8),
    "SEP": ("Sep", 9), "SEPT": ("Sep", 9), "SEPTEMBER": ("Sep", 9),
    "OKT": ("Oct", 10), "OCT": ("Oct", 10), "OKTOBER": ("Oct", 10), "OCTOBER": ("Oct", 10),
    "NOV": ("Nov", 11), "NOVEMBER": ("Nov", 11),
    "DES": ("Dec", 12), "DEC": ("Dec", 12), "DESEMBER": ("Dec", 12), "DECEMBER": ("Dec", 12),
}

MONTH_NAMES = {order: name for name, order in [value for value in MONTH_ALIASES.values()]}
SUMMARY_COLUMNS = [
    "No", "Bulan", "Mutasi Debet Nominal (Rp)", "Mutasi Debet Frek",
    "Mutasi Kredit Nominal (Rp)", "Mutasi Kredit Frek", "Saldo (Rp)",
    "Saldo Awal (Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
]


def clean_text(value):
    if pd.isna(value):
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def parse_number(value):
    text = clean_text(value).upper()
    if text in ("", "NAN", "-"):
        return pd.NA
    text = text.replace(",", "").replace("+", "").strip()
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return pd.NA


def excel_number(value):
    if pd.isna(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def dataframe_for_excel(df):
    result = df.copy()
    for col in result.columns:
        result[col] = result[col].apply(excel_number)
    return result


def is_money(text):
    return bool(re.fullmatch(r"\d[\d,]*\.\d{2}", clean_text(text)))


def is_transaction_date(text):
    return bool(re.fullmatch(r"\d{2}/\d{2}/\d{2}\s+\d{2}:\d{2}:\d{2}", clean_text(text)))


def date_to_day_month(text):
    match = re.match(r"^(\d{2})/(\d{2})/\d{2}\s+\d{2}:\d{2}:\d{2}$", clean_text(text))
    if not match:
        return clean_text(text)
    return f"{match.group(1)}/{match.group(2)}"


def normalize_metadata_text(text):
    return re.sub(r"\s+", " ", clean_text(text)).upper()


@lru_cache(maxsize=None)
def read_pdf_text(pdf_file):
    pdf_file = Path(pdf_file)
    text_parts = []
    try:
        with fitz.open(pdf_file) as doc:
            for page in doc:
                text_parts.append(page.get_text("text"))
    except Exception:
        return ""
    return "\n".join(text_parts)


def pdf_text_lines(pdf_file):
    return [clean_text(line) for line in read_pdf_text(Path(pdf_file)).splitlines() if clean_text(line)]


def extract_account_from_text(text):
    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]
    for idx, line in enumerate(lines):
        if line.upper() in {"NO. REKENING", "ACCOUNT NO"} and idx + 1 < len(lines):
            match = re.search(r"\d{6,}", lines[idx + 1])
            if match:
                return match.group(0)
        if line == ":" and idx > 0 and idx + 1 < len(lines) and lines[idx - 1].upper() in {"ACCOUNT NO", "NO. REKENING"}:
            match = re.search(r"\d{6,}", lines[idx + 1])
            if match:
                return match.group(0)

    normalized = normalize_metadata_text(text)
    match = re.search(r"(?:NO\.?\s*REKENING|ACCOUNT\s+NO)\s*:?\s*(\d{6,})", normalized)
    if match:
        return match.group(1)
    matches = re.findall(r"\b\d{8,}\b", normalized)
    return matches[0] if matches else None


def extract_account_from_filename(pdf_file):
    stem = Path(pdf_file).stem
    matches = re.findall(r"\b\d{6,}\b", stem)
    return matches[0] if matches else "Unknown"


def extract_period_from_text(text):
    normalized = normalize_metadata_text(text)
    match = re.search(r"PERIODE\s+TRANSAKSI\s+TRANSACTION\s+PERIODE?\s*:?\s*(\d{2})/(\d{2})/(\d{2,4})\s*-\s*(\d{2})/(\d{2})/(\d{2,4})", normalized)
    if not match:
        match = re.search(r"(\d{2})/(\d{2})/(\d{2,4})\s*-\s*(\d{2})/(\d{2})/(\d{2,4})", normalized)
    if not match:
        return None

    _, start_month, start_year, _, end_month, end_year = match.groups()
    year = f"20{start_year}" if len(start_year) == 2 else start_year
    end_full_year = f"20{end_year}" if len(end_year) == 2 else end_year
    month_order = int(start_month)
    month_name = MONTH_NAMES.get(month_order, start_month)
    end_month_order = int(end_month)
    end_month_name = MONTH_NAMES.get(end_month_order, end_month)
    return {
        "month_label": f"{month_name}-{year[-2:]}",
        "month_order": month_order,
        "year": year,
        "start_year": year,
        "start_month_order": month_order,
        "end_year": end_full_year,
        "end_month_order": end_month_order,
        "period_label": (
            f"{month_name}-{year[-2:]}"
            if month_order == end_month_order and year == end_full_year
            else f"{month_name}-{year[-2:]} to {end_month_name}-{end_full_year[-2:]}"
        ),
    }


def extract_period_from_filename(pdf_file):
    stem = Path(pdf_file).stem.upper()
    year = None
    year_match = re.search(r"\b(20\d{2}|\d{2})\b", stem)
    if year_match:
        token = year_match.group(1)
        year = f"20{token}" if len(token) == 2 else token

    month_label = None
    month_order = 99
    for token in re.split(r"[^A-Z]+", stem):
        if token in MONTH_ALIASES:
            month_name, month_order = MONTH_ALIASES[token]
            month_label = f"{month_name}-{year[-2:]}" if year else month_name
            break

    year = year or "Unknown"
    return {
        "month_label": month_label or Path(pdf_file).stem[:20],
        "month_order": month_order,
        "year": year,
        "start_year": year,
        "start_month_order": month_order,
        "end_year": year,
        "end_month_order": month_order,
        "period_label": month_label or Path(pdf_file).stem[:20],
    }


def extract_pdf_metadata(pdf_file):
    pdf_file = Path(pdf_file)
    text = read_pdf_text(pdf_file)
    period = extract_period_from_text(text) or extract_period_from_filename(pdf_file)
    account = extract_account_from_text(text) or extract_account_from_filename(pdf_file)
    period["account"] = account
    period["source_file"] = pdf_file.name
    return period


def build_summary_from_text(pdf_file):
    lines = pdf_text_lines(pdf_file)
    labels = ["Saldo Awal", "Total Pengeluaran", "Total Pemasukan", "Saldo Akhir"]
    summary_values = {}

    for idx, line in enumerate(lines):
        if line.upper() != "SALDO AWAL":
            continue
        window = lines[idx:idx + 24]
        if not any("TOTAL TRANSAKSI DEBET" in item.upper() for item in window):
            continue
        amounts = [parse_number(item) for item in window if is_money(item)]
        amounts = [amount for amount in amounts if not pd.isna(amount)]
        if len(amounts) >= 4:
            summary_values = dict(zip(labels, amounts[:4]))
            break

    return pd.DataFrame(
        [{"Keterangan": label, "Amount": summary_values.get(label, pd.NA)} for label in labels],
        columns=["Keterangan", "Amount"],
    )


def summary_value(df_summary, label):
    if df_summary.empty:
        return pd.NA
    mask = df_summary["Keterangan"].str.upper().eq(label.upper())
    if not mask.any():
        return pd.NA
    return df_summary.loc[mask, "Amount"].iloc[0]


def parse_transaction_block(block):
    date_text = block[0]
    # BRI repeats page footers/headers between transactions. Use the first valid
    # teller/debit/credit/balance group after the transaction date.
    for idx in range(4, len(block)):
        if is_money(block[idx - 2]) and is_money(block[idx - 1]) and is_money(block[idx]):
            if idx == 4:
                teller = ""
                description_lines = block[1:idx - 2]
            else:
                teller = block[idx - 3]
                description_lines = block[1:idx - 3]
            description = " ".join(description_lines).strip()
            debit = parse_number(block[idx - 2])
            credit = parse_number(block[idx - 1])
            balance = parse_number(block[idx])
            if not description or pd.isna(debit) or pd.isna(credit) or pd.isna(balance):
                continue
            return {
                "Tanggal": date_to_day_month(date_text),
                "Keterangan": description,
                "Teller": teller,
                "DB": debit if debit != 0 else pd.NA,
                "CR": credit if credit != 0 else pd.NA,
                "Saldo": balance,
                "RawDate": date_text,
            }
    return None


def read_transactions_from_text(pdf_file):
    lines = pdf_text_lines(pdf_file)
    if not lines:
        raise ValueError(
            f"{Path(pdf_file).name} has no extractable PDF text. It appears to be an image-only scan; OCR is required."
        )

    transactions = []
    idx = 0
    while idx < len(lines):
        if not is_transaction_date(lines[idx]):
            idx += 1
            continue
        next_idx = idx + 1
        while next_idx < len(lines) and not is_transaction_date(lines[next_idx]):
            next_idx += 1
        row = parse_transaction_block(lines[idx:next_idx])
        if row:
            transactions.append(row)
        idx = next_idx

    df = pd.DataFrame(transactions, columns=["Tanggal", "Keterangan", "Teller", "DB", "CR", "Saldo", "RawDate"])
    if df.empty:
        raise ValueError(f"No BRI transaction rows found in {Path(pdf_file).name}")
    return df


def reconcile_transactions_with_balance(clean_df, df_summary):
    opening = summary_value(df_summary, "Saldo Awal")
    if pd.isna(opening):
        return clean_df

    rows = clean_df.copy()
    last_balance = opening
    for idx, row in rows.iterrows():
        db = row["DB"] if not pd.isna(row["DB"]) else Decimal("0")
        cr = row["CR"] if not pd.isna(row["CR"]) else Decimal("0")
        expected_credit_balance = last_balance + cr
        expected_debit_balance = last_balance - db
        actual = row["Saldo"]

        if db == 0 and cr == 0:
            diff = actual - last_balance
            if diff > 0:
                rows.at[idx, "CR"] = diff
            elif diff < 0:
                rows.at[idx, "DB"] = abs(diff)
        elif not pd.isna(actual) and expected_credit_balance == actual and not pd.isna(row["DB"]):
            rows.at[idx, "CR"] = db
            rows.at[idx, "DB"] = pd.NA
        elif not pd.isna(actual) and expected_debit_balance == actual and not pd.isna(row["CR"]):
            rows.at[idx, "DB"] = cr
            rows.at[idx, "CR"] = pd.NA

        last_balance = actual if not pd.isna(actual) else last_balance

    return rows


def count_amount(df, column):
    return int(df[column].apply(lambda value: not pd.isna(value) and value != 0).sum())


def sum_amount(df, column):
    values = [value for value in df[column] if not pd.isna(value)]
    return sum(values, Decimal("0"))


def sum_by_description(df, amount_column, pattern):
    if df.empty:
        return pd.NA

    mask = df["Keterangan"].str.upper().str.contains(pattern, regex=True, na=False)
    values = df.loc[mask, amount_column].dropna()

    if values.empty:
        return pd.NA
    return sum(values, Decimal("0"))


def report_extraction_balance(clean_df, df_summary, pdf_file):
    summary_db = summary_value(df_summary, "Total Pengeluaran")
    summary_cr = summary_value(df_summary, "Total Pemasukan")
    extracted_db = sum_amount(clean_df, "DB")
    extracted_cr = sum_amount(clean_df, "CR")
    mismatches = []
    if not pd.isna(summary_db) and extracted_db != summary_db:
        mismatches.append(f"DB extracted {extracted_db} but PDF summary {summary_db}")
    if not pd.isna(summary_cr) and extracted_cr != summary_cr:
        mismatches.append(f"CR extracted {extracted_cr} but PDF summary {summary_cr}")
    if mismatches:
        print(f"PERINGATAN: {Path(pdf_file).name}: " + "; ".join(mismatches))
    else:
        print(f"Validasi total cocok dengan summary PDF: {Path(pdf_file).name}")


def process_pdf(pdf_file):
    """Parse through the validated shared PDF entry point."""
    return extract_frames(pdf_file, "BRI")


def safe_sheet_name(name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean_text(name)).strip()
    return cleaned[:31] or "Sheet"


def unique_sheet_name(base_name, used_names):
    base_name = safe_sheet_name(base_name)
    candidate = base_name
    counter = 2
    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def build_monthly_display_dataframe(clean_df, df_summary, month_order=None):
    saldo_awal = summary_value(df_summary, "Saldo Awal")
    opening_row = {
        "Tanggal": f"01/{month_order:02d}" if month_order else "",
        "Keterangan": "SALDO AWAL",
        "DB": pd.NA,
        "CR": pd.NA,
        "Saldo": saldo_awal,
    }
    display_df = clean_df[["Tanggal", "Keterangan", "DB", "CR", "Saldo"]].copy()
    return pd.concat([pd.DataFrame([opening_row]), display_df], ignore_index=True)


def build_month_summary(clean_df, df_summary, pdf_file, metadata):
    return {
        "No": None,
        "Bulan": metadata["month_label"],
        "Mutasi Debet Nominal (Rp)": summary_value(df_summary, "Total Pengeluaran"),
        "Mutasi Debet Frek": count_amount(clean_df, "DB"),
        "Mutasi Kredit Nominal (Rp)": summary_value(df_summary, "Total Pemasukan"),
        "Mutasi Kredit Frek": count_amount(clean_df, "CR"),
        "Saldo (Rp)": summary_value(df_summary, "Saldo Akhir"),
        "Saldo Awal (Rp)": summary_value(df_summary, "Saldo Awal"),
        **summary_metrics(clean_df),
        "Year": metadata["year"],
        "MonthOrder": metadata["month_order"],
        "Source File": Path(pdf_file).name,
    }


def monthly_metrics(clean_df):
    db_total = sum_amount(clean_df, "DB")
    cr_total = sum_amount(clean_df, "CR")
    db_freq = count_amount(clean_df, "DB")
    cr_freq = count_amount(clean_df, "CR")
    return [
        ("Total DB", db_total), ("Frekuensi DB", db_freq), ("Rata-rata DB", db_total / db_freq if db_freq else pd.NA),
        ("Total CR", cr_total), ("Frekuensi CR", cr_freq), ("Rata-rata CR", cr_total / cr_freq if cr_freq else pd.NA),
    ]


def auto_fit_columns(sheet):
    for col_idx, col_cells in enumerate(sheet.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        sheet.column_dimensions[col_letter].width = max_length + 2


def add_monthly_sheet_totals(sheet):
    first_data_row = 2
    last_data_row = sheet.max_row
    total_row = last_data_row + 3
    freq_row = total_row + 1
    average_row = total_row + 2

    for row_num, label in {total_row: "Total", freq_row: "Frekuensi", average_row: "Rata2"}.items():
        sheet.cell(row=row_num, column=2, value=label)

    for col_num in (3, 4):
        col_letter = get_column_letter(col_num)
        sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=freq_row, column=col_num, value=f"=COUNT({col_letter}{first_data_row}:{col_letter}{last_data_row})")
        sheet.cell(row=average_row, column=col_num, value=f"=IFERROR(AVERAGE({col_letter}{first_data_row}:{col_letter}{last_data_row}),0)")


def style_monthly_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary_start = sheet.max_row - 2

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    for col_letter, width in {"A": 10, "B": 95, "C": 16, "D": 16, "E": 16}.items():
        sheet.column_dimensions[col_letter].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column in (3, 4, 5):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row_num in range(summary_start, sheet.max_row + 1):
        label_cell = sheet.cell(row=row_num, column=2)
        label_cell.fill = header_fill
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_num in (3, 4):
            value_cell = sheet.cell(row=row_num, column=col_num)
            value_cell.number_format = "#,##0.00" if row_num != summary_start + 1 else "0"
            value_cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_num in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 18


def write_monthly_transaction_sheet(writer, clean_df, df_summary, sheet_name, month_order=None):
    display_df = build_monthly_display_dataframe(clean_df, df_summary, month_order)
    display_df = dataframe_for_excel(display_df)
    display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    sheet = writer.sheets[sheet_name]
    add_monthly_sheet_totals(sheet)
    style_monthly_sheet(sheet)


def write_summary_sheet(writer, summary_rows):
    workbook = writer.book
    sheet = workbook.create_sheet("Summary", 0)
    writer.sheets["Summary"] = sheet

    headers_top = [
        "No", "Bulan", "Mutasi Debet", None, "Mutasi Kredit", None,
        "Saldo", "Saldo Awal", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]
    headers_bottom = [
        "No", "Bulan", "Nominal (Rp)", "Frek", "Nominal (Rp)", "Frek",
        "(Rp)", "(Rp)", "Adm", "Pajak", "Bunga", "Saldo Min", "JaGir",
    ]

    for col_num, value in enumerate(headers_top, start=1):
        if value is not None:
            sheet.cell(row=1, column=col_num, value=value)
    for col_num, value in enumerate(headers_bottom, start=1):
        sheet.cell(row=2, column=col_num, value=value)

    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    sheet.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=4)
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    for col_num in range(7, 14):
        sheet.merge_cells(start_row=1, start_column=col_num, end_row=2, end_column=col_num)

    sorted_rows = sorted(summary_rows, key=lambda row: (row["Year"], row["MonthOrder"], row["Source File"]))
    for row_idx, row in enumerate(sorted_rows, start=3):
        values = [
            row_idx - 2,
            row["Bulan"],
            row["Mutasi Debet Nominal (Rp)"],
            row["Mutasi Debet Frek"],
            row["Mutasi Kredit Nominal (Rp)"],
            row["Mutasi Kredit Frek"],
            row["Saldo (Rp)"],
            row["Saldo Awal (Rp)"],
            row["Adm"],
            row["Pajak"],
            row["Bunga"],
            row["Saldo Min"],
            row["JaGir"],
        ]
        for col_num, value in enumerate(values, start=1):
            sheet.cell(row=row_idx, column=col_num, value=excel_number(value))

    total_row = len(sorted_rows) + 3
    average_row = total_row + 1
    sheet.cell(row=total_row, column=2, value="Total")
    sheet.cell(row=average_row, column=2, value="Rata-Rata")

    data_start = 3
    data_end = total_row - 1
    if data_end >= data_start:
        for col_num in [3, 4, 5, 6, 9, 10, 11, 12, 13]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=total_row, column=col_num, value=f"=SUM({col_letter}{data_start}:{col_letter}{data_end})")
        for col_num in [3, 4, 5, 6, 7]:
            col_letter = get_column_letter(col_num)
            sheet.cell(row=average_row, column=col_num, value=f"=AVERAGE({col_letter}{data_start}:{col_letter}{data_end})")

    style_summary_sheet(sheet)
    auto_fit_columns(sheet)


def style_transaction_sheet(sheet, summary_start=None):
    header_fill = PatternFill("solid", fgColor="BDD7EE")
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx in (3, 4, 5):
            row[col_idx - 1].number_format = "#,##0.00"

    if summary_start:
        for row_num in range(summary_start, sheet.max_row + 1):
            sheet.cell(row=row_num, column=2).font = Font(bold=True)
            sheet.cell(row=row_num, column=3).number_format = "#,##0.00" if "Frekuensi" not in str(sheet.cell(row=row_num, column=2).value) else "0"

    widths = {"A": 12, "B": 60, "C": 18, "D": 18, "E": 18}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"


def style_summary_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            if cell.row in (1, 2):
                cell.fill = header_fill
                cell.font = Font(bold=True)
            elif cell.row >= 3 and cell.column in (3, 5, 7, 8, 9, 10, 11, 12, 13):
                cell.number_format = "#,##0.00"
            elif cell.row >= 3 and cell.column in (4, 6):
                cell.number_format = "0"

    for row_num in (sheet.max_row - 1, sheet.max_row):
        for col_num in range(1, sheet.max_column + 1):
            sheet.cell(row=row_num, column=col_num).font = Font(bold=True)


def export_year_workbook(extracted_files, output_file):
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    used_names = set()
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        write_summary_sheet(writer, [item["summary_row"] for item in extracted_files])
        for item in sorted(extracted_files, key=lambda row: (row["month_order"], row["pdf_file"].name)):
            sheet_name = unique_sheet_name(item["sheet_name"], used_names)
            write_monthly_transaction_sheet(
                writer,
                item["transactions"],
                item["summary"],
                sheet_name,
                item["month_order"],
            )


def export_to_excel(df_final, df_summary, output_file, pdf_file=None):
    metadata = extract_pdf_metadata(pdf_file) if pdf_file else {"month_label": "Transaksi", "month_order": 99, "year": ""}
    export_year_workbook([{
        "pdf_file": Path(pdf_file) if pdf_file else Path(""),
        "transactions": df_final,
        "summary": df_summary,
        "summary_row": build_month_summary(df_final, df_summary, pdf_file or "", metadata),
        "sheet_name": metadata["month_label"].replace("-", "_"),
        "month_order": metadata["month_order"],
        "year": metadata["year"],
    }], output_file)


def discover_pdf_groups_by_content(input_path):
    input_path = Path(input_path)
    pdf_files = [input_path] if input_path.is_file() else sorted(input_path.rglob("*.pdf"))
    groups = {}
    for pdf_file in pdf_files:
        metadata = extract_pdf_metadata(pdf_file)
        key = (metadata["account"], metadata["year"])
        groups.setdefault(key, {"account": metadata["account"], "year": metadata["year"], "files": []})
        groups[key]["files"].append({"path": pdf_file, "metadata": metadata})
    return groups


def ensure_xlsx_name(name):
    name = clean_text(name) or "{account}_{year}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def output_name_for_group(group_info, output_name_template=None, force_unique=False):
    account = group_info["account"]
    year = group_info["year"]
    if not output_name_template:
        return f"{account}_{year}.xlsx"

    output_name = ensure_xlsx_name(output_name_template)
    has_year = "{year}" in output_name
    try:
        output_name = output_name.format(account=account, year=year, folder=year)
    except (KeyError, ValueError) as exc:
        raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc
    if force_unique and not has_year:
        path = Path(output_name)
        output_name = f"{path.stem}_{year}{path.suffix}"
    return output_name


def run_folder(input_path, output_dir, output_name_template=None, preserve_relative_folders=True):
    groups = discover_pdf_groups_by_content(input_path)
    written_files = []
    skipped_files = []
    if not groups:
        raise FileNotFoundError(f"No PDF files found in folder: {input_path}")

    for _, group_info in sorted(groups.items(), key=lambda item: (item[0][1], item[0][0])):
        print(f"\nMemproses rekening {group_info['account']} tahun {group_info['year']} ({len(group_info['files'])} file PDF)")
        extracted_files = []
        for file_item in sorted(group_info["files"], key=lambda item: (item["metadata"]["month_order"], item["path"].name)):
            pdf_file = file_item["path"]
            metadata = file_item["metadata"]
            print(f"Memproses file: {pdf_file.name}")
            try:
                df_final, df_summary = process_pdf(pdf_file)
            except ValueError as exc:
                print(f"PERINGATAN: melewati {pdf_file.name}: {exc}")
                skipped_files.append(pdf_file)
                continue
            extracted_files.append({
                "pdf_file": pdf_file,
                "transactions": df_final,
                "summary": df_summary,
                "summary_row": build_month_summary(df_final, df_summary, pdf_file, metadata),
                "sheet_name": metadata["month_label"].replace("-", "_"),
                "month_order": metadata["month_order"],
                "year": metadata["year"],
            })

        if not extracted_files:
            continue

        relative_folder = Path(group_info["year"]) if preserve_relative_folders else Path()
        output_file = output_dir / relative_folder / output_name_for_group(
            group_info,
            output_name_template,
            force_unique=len(groups) > 1,
        )
        export_year_workbook(extracted_files, output_file)
        written_files.append(output_file)
        print(f"Workbook selesai dibuat: {output_file}")

    if skipped_files:
        print("\nPDF image-only yang dilewati karena tidak ada teks yang bisa diekstrak:")
        for pdf_file in skipped_files:
            print(f"- {pdf_file}")
    return written_files


def convert_pdf(pdf_file, output_file):
    """
    Convert one BRI PDF statement into one Excel workbook.

    Kept as a small wrapper around the existing parsing/export pipeline so CLI
    behavior and workbook formatting stay the same.
    """
    pdf_file = Path(pdf_file)
    output_file = Path(output_file)
    df_final, df_summary = process_pdf(pdf_file)
    print("Menulis ke Excel...")
    output_file.parent.mkdir(parents=True, exist_ok=True)
    export_to_excel(df_final, df_summary, output_file, pdf_file)
    print(f"Data successfully written to {output_file}")
    return output_file


def run_single_file(pdf_file, output_file):
    convert_pdf(pdf_file, output_file)


def choose_batch_options_with_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except ImportError:
        return None, None, None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    messagebox.showinfo("BRI PDF to Excel", "Choose the BRI folder that contains your PDF files.", parent=root)
    input_dir = filedialog.askdirectory(
        title="Choose BRI PDF folder",
        initialdir=str(BRI_DIR.resolve()) if BRI_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not input_dir:
        root.destroy()
        return None, None, None
    messagebox.showinfo("BRI PDF to Excel", "Choose the folder where the Excel output should be saved.", parent=root)
    output_dir = filedialog.askdirectory(
        title="Choose Excel output folder",
        initialdir=str(EXCEL_DIR.resolve()) if EXCEL_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not output_dir:
        root.destroy()
        return None, None, None
    output_name = simpledialog.askstring(
        "BRI PDF to Excel",
        "Enter Excel filename pattern:\nUse {account}, {year}, or {folder} if needed.",
        initialvalue="{account}_{year}.xlsx",
        parent=root,
    )
    root.destroy()
    if output_name is None:
        return None, None, None
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name)


def choose_batch_options_with_console():
    input_dir = input("Choose input folder path: ").strip()
    output_dir = input("Choose output folder path: ").strip()
    output_name = input("Excel filename pattern [{account}_{year}.xlsx]: ").strip()
    return Path(input_dir), Path(output_dir), ensure_xlsx_name(output_name or "{account}_{year}.xlsx")


def show_gui_result(message, is_error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
    except ImportError:
        return
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    if is_error:
        messagebox.showerror("BRI PDF to Excel", message, parent=root)
    else:
        messagebox.showinfo("BRI PDF to Excel", message, parent=root)
    root.destroy()


def parse_args():
    parser = argparse.ArgumentParser(description="Convert BRI account statement PDF files to Excel.")
    parser.add_argument("input", nargs="?", default=None, help="PDF file path or folder path. No arguments opens folder dialogs.")
    parser.add_argument("-o", "--output-dir", default=str(EXCEL_DIR), help="Output folder. Default: excel_file/BRI")
    parser.add_argument("--output-name", default=None, help="Workbook filename pattern. Supports {account}, {year}, {folder}.")
    parser.add_argument("--flat-output", action="store_true", help="Save yearly workbooks directly in output-dir.")
    return parser.parse_args()


def resolve_input_path(input_value):
    input_path = Path(input_value)
    if input_path.exists():
        return input_path
    candidate_pdf = Path("pdf_file") / f"{input_path}.pdf"
    candidate_folder = Path("pdf_file") / input_path
    if candidate_pdf.exists():
        return candidate_pdf
    if candidate_folder.exists():
        return candidate_folder
    raise FileNotFoundError(f"Input not found: {input_path}")


def main():
    args = parse_args()
    use_gui = len(sys.argv) == 1
    if use_gui:
        selected = choose_batch_options_with_gui()
        if selected == (None, None, None):
            selected = choose_batch_options_with_console()
        input_path, excel_dir, output_name_template = selected
    else:
        input_path = resolve_input_path(args.input or BRI_DIR)
        excel_dir = Path(args.output_dir)
        output_name_template = args.output_name

    excel_dir.mkdir(parents=True, exist_ok=True)
    if input_path.is_file():
        metadata = extract_pdf_metadata(input_path)
        output_name = ensure_xlsx_name(output_name_template or f"{input_path.stem}.xlsx")
        try:
            output_name = output_name.format(account=metadata["account"], year=metadata["year"], folder=input_path.parent.name)
        except (KeyError, ValueError) as exc:
            raise ValueError("Output name template only supports {account}, {year}, and {folder}.") from exc
        run_single_file(input_path, excel_dir / output_name)
    else:
        written_files = run_folder(
            input_path,
            excel_dir,
            output_name_template,
            preserve_relative_folders=not getattr(args, "flat_output", False),
        )
        print("\nSemua proses selesai.")
        for output_file in written_files:
            print(f"- {output_file}")
        if use_gui:
            show_gui_result("Done.\n\nSaved files:\n" + "\n".join(str(path) for path in written_files))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}")
        if len(sys.argv) == 1:
            show_gui_result(str(exc), is_error=True)
            sys.exit(1)
        raise
