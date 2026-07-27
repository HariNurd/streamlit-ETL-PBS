import argparse
import calendar
import re
import sys
from datetime import date
from pathlib import Path

import fitz
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_dashboard import write_dashboard_sheets


SLIK_DIR = Path("slik")
EXCEL_DIR = Path("excel_file")
FALLBACK_KETERANGAN = "Tidak ada Fasilitas Aktif"
IDENTITY_COLUMN = "NIK/NPWP"

OUTPUT_COLUMNS = [
    "Nama",
    IDENTITY_COLUMN,
    "Bank",
    "Penggunaan",
    "Plafond (Rp)",
    "Baki Debet (Rp)",
    "Kol.",
    "Awal",
    "Jatuh Tempo",
    "Rate",
    "No Laporan",
    "Keterangan",
]

BANK_GARANSI_FACILITY_TYPE = "Garansi Yang Diberikan"
CREDIT_FACILITY_TYPE = "Kredit/Pembiayaan"

BANK_GARANSI_COLUMNS = [
    "Nama",
    IDENTITY_COLUMN,
    "Bank",
    "No Rekening",
    "Jenis Garansi",
    "Tujuan Garansi",
    "Nama Yang Dijamin",
    "Plafond (Rp)",
    "Nominal (Rp)",
    "Setoran Jaminan (Rp)",
    "Kol.",
    "Tanggal Akad Awal",
    "Tanggal Akad Akhir",
    "Tanggal Diterbitkan",
    "Jatuh Tempo",
    "No Akad Awal",
    "No Akad Akhir",
    "No Laporan",
    "Keterangan",
]

FACILITY_DETAIL_COLUMNS = [
    "Jenis Fasilitas",
    "No Rekening",
    "No Akad Awal",
    "No Akad Akhir",
    "Jenis Garansi",
    "Tujuan Garansi",
    "Nama Yang Dijamin",
    "Nominal (Rp)",
    "Setoran Jaminan (Rp)",
    "Tanggal Diterbitkan",
    "Tanggal Akad Akhir",
]

FACILITY_DATA_COLUMNS = OUTPUT_COLUMNS + FACILITY_DETAIL_COLUMNS

KETERANGAN_PRIORITY = [
    "Lunas",
    "Lunas Dengan Diskon",
    "Lunas Karena Pengambilalihan Agunan",
    "Aktif",
    "Dihapusbukukan",
    "Hapus Tagih",
]

MONTH_MAP = {
    "Januari": "01",
    "Februari": "02",
    "Maret": "03",
    "April": "04",
    "Mei": "05",
    "Juni": "06",
    "Juli": "07",
    "Agustus": "08",
    "September": "09",
    "Oktober": "10",
    "November": "11",
    "Desember": "12",
}

FACILITY_START_LABELS = {
    CREDIT_FACILITY_TYPE,
    BANK_GARANSI_FACILITY_TYPE,
}

SECTION_END_LABELS = {
    "Agunan",
    "Penjamin",
    "Irrecovable L/C",
    "Surat Berharga",
    "Fasilitas Lain",
    "Nomor Laporan",
    "Operator",
}

FIELD_LABELS = {
    "Kredit/Pembiayaan",
    "No Rekening",
    "Kualitas",
    "Kualitas / Jumlah",
    "Kualitas / Jumlah Hari",
    "Kualitas / Jumlah Hari Tunggakan",
    "Hari Tunggakan",
    "Pelapor",
    "Cabang",
    "Baki Debet",
    "Tanggal Update",
    "No Akad Akhir",
    "Sebab Macet",
    "Kredit Program Pemerintah",
    "Cara Restrukturisasi",
    "Kab/Kota Lokasi Proyek",
    "Kondisi",
    "Sektor Ekonomi",
    "Tanggal Restrukturisasi Akhir",
    "Suku Bunga/Imbalan",
    "Jenis Suku Bunga/Imbalan",
    "No Akad Awal",
    "Realisasi/Pencairan Bulan Berjalan",
    "Tanggal Awal Kredit",
    "Tunggakan Pokok",
    "Kategori Debitur",
    "Denda",
    "Valuta",
    "Tanggal Kondisi",
    "Jenis Kredit/Pembiayaan",
    "Nilai Proyek",
    "Keterangan",
    "Nama Yang Dijamin",
    "Nama LJK",
    "Nominal",
    "Tanggal Mulai",
    "Tunggakan Bunga",
    "Tanggal Akad Akhir",
    "Tanggal Macet",
    "Tanggal Diterbitkan",
    "Tanggal Wan Prestasi",
    "Sifat Kredit/Pembiayaan",
    "Jumlah Hari Tunggakan",
    "Akad Kredit/Pembiayaan",
    "Plafon Awal",
    "Jenis Penggunaan",
    "Jenis Garansi",
    "Tujuan Garansi",
    "Setoran Jaminan",
    "Frekuensi Restrukturisasi",
    "Frekuensi Perpanjangan Kredit/",
    "Pembiayaan",
    "Plafon",
    "Tanggal Jatuh Tempo",
    "Frekuensi Tunggakan",
    "Tanggal Akad Awal",
    "Nilai dalam Mata Uang Asal",
}

DEBTOR_NAME_EXACT_BLOCKLIST = {
    "INDONESIA",
    "JAKARTA",
    "JAKARTA UTARA",
    "KOTA JAKARTA",
    "KOTA JAKARTA UTARA",
    "PENJARINGAN",
    "KAPUK MUARA",
    "PLUIT",
    "BANDUNG",
    "POSISI DATA TERAKHIR",
    "NOMOR LAPORAN",
    "TANGGAL PERMINTAAN",
    "TANGGAL CETAK",
    "DATA POKOK DEBITUR",
    "RINGKASAN FASILITAS",
    "FASILITAS",
    "KREDIT/PEMBIAYAAN",
    "KREDIT / PEMBIAYAAN",
    "INFORMASI DEBITUR",
    "INFORMASI DIBERIKAN",
    "NAMA",
    "NAMA DEBITUR",
    "NO. IDENTITAS",
    "NO IDENTITAS",
    "NOMOR IDENTITAS",
    "IDENTITAS",
    "NIK",
    "NIK /",
    "NPWP",
    "JENIS KELAMIN",
    "TEMPAT LAHIR",
    "TANGGAL LAHIR",
    "PELAPOR",
    "OPERATOR",
}

DEBTOR_NAME_BLOCKED_PHRASES = [
    "POSISI DATA",
    "NOMOR LAPORAN",
    "TANGGAL PERMINTAAN",
    "TANGGAL CETAK",
    "DATA POKOK DEBITUR",
    "RINGKASAN FASILITAS",
    "DATA TIDAK DITEMUKAN",
    "KREDIT/PEMBIAYAAN",
    "KREDIT / PEMBIAYAAN",
    "INFORMASI DEBITUR",
    "INFORMASI DIBERIKAN",
    "SISTEM LAYANAN",
    "KEUANGAN DENGAN",
    "OTORITAS JASA KEUANGAN",
    "SISTEM LAYANAN INFORMASI KEUANGAN",
    "KODE REF",
    "KATA KUNCI PENCARIAN",
    "TEMPAT PENDIRIAN",
    "TANGGAL AKTE PENDIRIAN",
    "TANGGAL DIBENTUK",
]

DEBTOR_NAME_BLOCKED_TOKENS = {
    "BANK",
    "FINANCE",
    "IDEB",
    "SLIK",
    "NIK",
    "NPWP",
    "LAKI-LAKI",
    "PEREMPUAN",
    "PULAU",
    "KEPULAUAN",
    "KONSULTAN",
    "S-1",
    "RAHASIA",
    "PERDAGANGAN",
    "COMMANDITER",
    "PERSEROAN",
    "VENOTSCHAP",
    "WIRASWASTA",
    "LAIN-LAIN",
    "RUMAH TANGGA",
    "TANPA GELAR",
    "WIL",
    "LJK",
}


def clean_text(value):
    text = str(value).replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def clean_lines(text):
    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def parse_idr_number(value):
    text = clean_text(value)
    if not text:
        return pd.NA

    text = text.replace("Rp", "").replace(" ", "")
    text = re.sub(r"[^0-9,.\-]", "", text)
    if not text:
        return pd.NA

    text = text.replace(".", "").replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return pd.NA

    return int(number) if number.is_integer() else number


def parse_rate_percent(value):
    text = clean_text(value)
    if not text:
        return pd.NA

    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", text)
    if not match:
        return pd.NA

    try:
        return float(match.group(0).replace(",", ".")) / 100
    except ValueError:
        return pd.NA


def normalize_indonesian_date(text):
    text = clean_text(text)
    match = re.match(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    if not match:
        return text

    day, month_name, year = match.groups()
    month = MONTH_MAP.get(month_name)
    if not month:
        return text

    return f"{int(day):02d}-{month}-{year}"


def parse_output_date(value):
    text = normalize_indonesian_date(value)
    match = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{4})$", clean_text(text))
    if not match:
        return None

    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def add_months_clamped(start_date, months):
    month_index = start_date.month - 1 + months
    year = start_date.year + month_index // 12
    month = month_index % 12 + 1
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, min(start_date.day, last_day))


def worst_kolektibilitas(values):
    numeric_values = pd.to_numeric(values, errors="coerce").dropna()
    if numeric_values.empty:
        return ""
    return int(numeric_values.max())


def keterangan_rank(value):
    text = clean_text(value).casefold()
    compact = re.sub(r"[^a-z0-9]+", "", text)
    if not text:
        return 0
    if "hapus tagih" in text:
        return 6
    if "dihapusbukukan" in compact or "hapus buku" in text:
        return 5
    if "aktif" in text:
        return 4
    if "pengambilalihan agunan" in text:
        return 3
    if "diskon" in text:
        return 2
    if "lunas" in text:
        return 1
    return 0


def worst_keterangan(values):
    worst_rank = 0
    worst_value = ""

    for value in values:
        text = clean_text(value)
        rank = keterangan_rank(text)
        if rank > worst_rank:
            worst_rank = rank
            worst_value = text

    if worst_value:
        return worst_value

    return next((clean_text(value) for value in values if clean_text(value)), "")


def is_hapus_tagih_condition(kondisi):
    return "hapus tagih" in clean_text(kondisi).casefold()


def is_dihapusbukukan_condition(kondisi):
    text = clean_text(kondisi).casefold()
    compact = re.sub(r"[^a-z0-9]+", "", text)
    return "dihapusbukukan" in compact or "hapus buku" in text


def is_active_condition(kondisi):
    text = clean_text(kondisi).casefold()
    return text == "aktif" or "fasilitas aktif" in text


def keterangan_from_kondisi(kondisi):
    if is_hapus_tagih_condition(kondisi):
        return "Hapus Tagih"
    if is_dihapusbukukan_condition(kondisi):
        return "Dihapusbukukan"
    if is_active_condition(kondisi):
        return "Aktif"
    return clean_text(kondisi)


def is_currency_line(text):
    return bool(re.search(r"\bRp\s*[-0-9.]+,\d{2}\b", clean_text(text)))


def is_label(text):
    text = clean_text(text)
    return (
        text in FIELD_LABELS
        or text in SECTION_END_LABELS
        or text.startswith("Kualitas / Jumlah")
    )


def read_pdf_page_lines(pdf_file):
    page_lines = []
    with fitz.open(pdf_file) as doc:
        for page in doc:
            page_lines.append(clean_lines(page.get_text("text")))
    return page_lines


def read_pdf_lines(pdf_file):
    return [line for page_lines in read_pdf_page_lines(pdf_file) for line in page_lines]


def is_possible_nik(value):
    digits = re.sub(r"\D", "", clean_text(value))
    if len(digits) != 16 or len(set(digits)) == 1:
        return False

    try:
        region = int(digits[:6])
        raw_day = int(digits[6:8])
        month = int(digits[8:10])
    except ValueError:
        return False

    day = raw_day - 40 if raw_day > 40 else raw_day
    return region > 0 and 1 <= day <= 31 and 1 <= month <= 12


def extract_nik_from_text(value):
    text = clean_text(value)
    for match in re.finditer(r"(?<!\d)(?:\d[\s.\-]?){16}(?!\d)", text):
        digits = re.sub(r"\D", "", match.group(0))
        if is_possible_nik(digits):
            return digits
    return ""


def is_possible_npwp(value):
    digits = re.sub(r"\D", "", clean_text(value))
    return len(digits) in {15, 16} and len(set(digits)) > 1


def extract_npwp_from_text(value):
    text = clean_text(value)
    for match in re.finditer(r"(?<!\d)(?:\d[\s.\-]?){15,16}(?!\d)", text):
        digits = re.sub(r"\D", "", match.group(0))
        if is_possible_npwp(digits):
            return digits
    return ""


def data_pokok_debitur_lines(page1_lines):
    lines = page1_lines
    try:
        start = lines.index("Data Pokok Debitur")
    except ValueError:
        start = 0

    stop = len(lines)
    for marker in ["Ringkasan Fasilitas", "Pemilik / Pengurus", "Nomor Laporan"]:
        for idx in range(start, len(lines)):
            if lines[idx].startswith(marker):
                stop = min(stop, idx)
                break

    return lines[start:stop]


def extract_labeled_nik(lines):
    for idx, line in enumerate(lines):
        if not re.search(r"\bNIK\b", clean_text(line), flags=re.IGNORECASE):
            continue

        nik = extract_nik_from_text(line)
        if nik:
            return nik

        for candidate in lines[idx + 1 : idx + 4]:
            nik = extract_nik_from_text(candidate)
            if nik:
                return nik
            if is_label(candidate):
                break

    return ""


def extract_nik_debitur(page1_lines):
    data_lines = data_pokok_debitur_lines(page1_lines)

    nik = extract_labeled_nik(data_lines)
    if nik:
        return nik

    for line in data_lines:
        nik = extract_nik_from_text(line)
        if nik:
            return nik

    for idx, line in enumerate(page1_lines):
        normalized = clean_text(line).casefold()
        if not re.fullmatch(r"(no\.?|nomor)\s*identitas|nik", normalized):
            continue

        for candidate in page1_lines[idx : idx + 4]:
            nik = extract_nik_from_text(candidate)
            if nik:
                return nik

    return ""


def extract_labeled_npwp(lines):
    for idx, line in enumerate(lines):
        if not re.search(r"\bNPWP\b", clean_text(line), flags=re.IGNORECASE):
            continue

        npwp = extract_npwp_from_text(line)
        if npwp:
            return npwp

        for candidate in lines[idx + 1 : idx + 8]:
            npwp = extract_npwp_from_text(candidate)
            if npwp:
                return npwp
            if is_label(candidate) or candidate in {"Pemilik / Pengurus", "Ringkasan Fasilitas"}:
                break

    return ""


def extract_npwp_debitur(page1_lines):
    data_lines = data_pokok_debitur_lines(page1_lines)

    npwp = extract_labeled_npwp(page1_lines)
    if npwp:
        return npwp

    npwp = extract_labeled_npwp(data_lines)
    if npwp:
        return npwp

    name_label_indexes = [idx for idx, line in enumerate(data_lines) if line == "Nama Debitur"]
    for start in name_label_indexes:
        for candidate in data_lines[start + 1 : start + 24]:
            npwp = extract_npwp_from_text(candidate)
            if npwp:
                return npwp
            if candidate == "Pemilik / Pengurus":
                break

    return ""


def extract_debitur_identity(page1_lines):
    nik = extract_nik_debitur(page1_lines)
    return nik or extract_npwp_debitur(page1_lines)


def is_nik_label_line(line):
    return bool(re.match(r"^NIK\b", clean_text(line), flags=re.IGNORECASE))


def contains_blocked_name_token(text):
    upper = clean_text(text).upper()
    for token in DEBTOR_NAME_BLOCKED_TOKENS:
        if re.search(rf"(?<![A-Z0-9]){re.escape(token)}(?![A-Z0-9])", upper):
            return True
    return False


def is_debtor_name_candidate(line):
    text = clean_text(line)
    upper = text.upper()

    if not text:
        return False

    if upper in DEBTOR_NAME_EXACT_BLOCKLIST:
        return False

    if any(phrase in upper for phrase in DEBTOR_NAME_BLOCKED_PHRASES):
        return False

    if contains_blocked_name_token(text):
        return False

    if re.search(r"\d", text):
        return False

    return bool(re.search(r"[A-Za-z]{3}", text))


def extract_name_from_data_pokok(page1_lines):
    data_lines = data_pokok_debitur_lines(page1_lines)

    # Primary SLIK layout: the debtor name is directly above the NIK label.
    for idx, line in enumerate(data_lines):
        if is_nik_label_line(line) and idx > 0:
            candidate = data_lines[idx - 1]
            if is_debtor_name_candidate(candidate):
                return clean_text(candidate)

    try:
        name_label_index = data_lines.index("Nama Debitur")
    except ValueError:
        name_label_index = -1

    if name_label_index >= 0:
        # Alternate layout: resolve candidates only inside the Nama Debitur area.
        name_area = data_lines[name_label_index + 1 :]

        for idx, line in enumerate(name_area):
            if line.upper() == "INDONESIA":
                for candidate in name_area[idx + 1 : idx + 5]:
                    if is_debtor_name_candidate(candidate):
                        return clean_text(candidate)

        for idx, line in enumerate(name_area):
            if re.match(r"^\d+\s*/", line):
                for candidate in reversed(name_area[:idx]):
                    if is_debtor_name_candidate(candidate):
                        return clean_text(candidate)

        for line in name_area:
            if is_debtor_name_candidate(line):
                return clean_text(line)

    return ""


def extract_nama_debitur_candidate(page1_lines):
    lines = page1_lines

    data_pokok_name = extract_name_from_data_pokok(lines)
    if data_pokok_name:
        return data_pokok_name

    # Header fallback for layouts where the name appears before the metadata body.
    header_stop = len(lines)
    for marker in ["Informasi diberikan", "Data Pokok Debitur"]:
        for idx, line in enumerate(lines):
            if line.startswith(marker):
                header_stop = min(header_stop, idx)

    for line in lines[:header_stop]:
        if is_debtor_name_candidate(line) and line == line.upper() and len(line.split()) >= 2:
            return clean_text(line)

    try:
        start = lines.index("Nama Debitur")
    except ValueError:
        start = -1

    if start >= 0:
        stop = len(lines)
        for marker in ["NPWP", "Pemilik / Pengurus", "Ringkasan Fasilitas"]:
            if marker in lines[start:]:
                stop = min(stop, start + lines[start:].index(marker))

        company_pattern = re.compile(r"^(CV|PT|UD|PD|KOPERASI)\b", flags=re.IGNORECASE)
        for line in lines[start:stop]:
            if company_pattern.search(line) and "BANK" not in line.upper():
                return clean_text(line)

        for idx in range(start + 1, stop):
            if is_debtor_name_candidate(lines[idx]):
                return clean_text(lines[idx])

    # Identity fallback: accept the line immediately before a NIK label only.
    for idx, line in enumerate(lines):
        if is_nik_label_line(line) and idx > 0:
            candidate = lines[idx - 1]
            if is_debtor_name_candidate(candidate):
                return clean_text(candidate)

    return ""


def extract_nama_debitur(page1_lines):
    name = extract_nama_debitur_candidate(page1_lines)
    return clean_text(name) if is_debtor_name_candidate(name) else ""


def extract_nomor_laporan(lines):
    for line in lines:
        match = re.search(r"\b\d+/IDEB/\d+/\d+\b", line)
        if match:
            return match.group(0)
    return ""


def split_facility_blocks(lines):
    starts = [idx for idx, line in enumerate(lines) if line in FACILITY_START_LABELS]
    blocks = []

    for pos, start in enumerate(starts):
        end = starts[pos + 1] if pos + 1 < len(starts) else len(lines)
        block = lines[start:end]

        cut_at = len(block)
        for idx, line in enumerate(block):
            if idx > 0 and line in SECTION_END_LABELS:
                cut_at = min(cut_at, idx)

        blocks.append(block[:cut_at])

    return blocks


def facility_type_from_block(lines):
    if not lines:
        return ""
    facility_type = clean_text(lines[0])
    return facility_type if facility_type in FACILITY_START_LABELS else ""


def next_values_after_label(lines, label, max_values=3):
    values = []
    for idx, line in enumerate(lines):
        if line != label:
            continue

        for candidate in lines[idx + 1 :]:
            if is_label(candidate):
                break
            values.append(candidate)
            if len(values) >= max_values:
                break
        break

    return values


def value_after_label(lines, label):
    values = next_values_after_label(lines, label, max_values=4)
    return clean_text(" ".join(values))


def max_number(*values):
    numbers = [value for value in values if not pd.isna(value)]
    if not numbers:
        return pd.NA
    return max(numbers)


def first_available_date(lines, labels):
    for label in labels:
        value = value_after_label(lines, label)
        if value:
            return normalize_indonesian_date(value)
    return ""


def extract_bank(lines):
    for idx, line in enumerate(lines):
        match = re.match(r"^\d{3}\s*-\s*(.+)$", line)
        if not match:
            continue

        bank = clean_text(match.group(1))
        if bank and "IDEB" not in bank.upper():
            return bank, idx

    for idx, line in enumerate(lines):
        if re.search(r"\b(PT|BANK|BPR|BPRS|FINANCE)\b", line, flags=re.IGNORECASE):
            return clean_text(re.sub(r"^\d{3}\s*-\s*", "", line)), idx

    return "", -1


def extract_baki_debet(lines, bank_index):
    start = max(bank_index + 1, 0)
    stop = len(lines)
    if "No Akad Akhir" in lines[start:]:
        stop = start + lines[start:].index("No Akad Akhir")

    for line in lines[start:stop]:
        if is_currency_line(line):
            return parse_idr_number(line)

    value = value_after_label(lines, "Baki Debet")
    return parse_idr_number(value)


def extract_kolektibilitas(lines):
    value = value_after_label(lines, "Kualitas")
    match = re.search(r"\b(\d+)\b", value)
    return match.group(1) if match else ""


def extract_worst_kolektibilitas(lines):
    values = []

    try:
        start = next(
            idx for idx, line in enumerate(lines)
            if line.startswith("Kualitas / Jumlah")
        )
    except StopIteration:
        start = 0

    stop = len(lines)
    if "No Akad Akhir" in lines[start:]:
        stop = start + lines[start:].index("No Akad Akhir")

    for line in lines[start:stop]:
        if re.fullmatch(r"[1-5]", line):
            values.append(int(line))

    current_kol = extract_kolektibilitas(lines)
    if current_kol.isdigit():
        values.append(int(current_kol))

    return str(max(values)) if values else current_kol


def parse_facility_block(lines, nama_debitur, nomor_laporan, nik_debitur=""):
    facility_type = facility_type_from_block(lines) or CREDIT_FACILITY_TYPE
    kondisi = value_after_label(lines, "Kondisi")
    # Kol. must reflect the current "Kualitas" field, not historical values in
    # "Kualitas / Jumlah Hari Tunggakan".
    kolektibilitas = extract_kolektibilitas(lines)

    is_active = is_active_condition(kondisi)
    is_hapus_tagih = is_hapus_tagih_condition(kondisi)
    is_dihapusbukukan = is_dihapusbukukan_condition(kondisi)
    has_non_normal_collectibility = bool(kolektibilitas) and kolektibilitas != "1"

    if not (is_active or is_dihapusbukukan or is_hapus_tagih or has_non_normal_collectibility):
        return None

    bank, bank_index = extract_bank(lines)
    keterangan = keterangan_from_kondisi(kondisi)
    plafon = parse_idr_number(value_after_label(lines, "Plafon"))
    plafon_awal = parse_idr_number(value_after_label(lines, "Plafon Awal"))
    baki_debet = extract_baki_debet(lines, bank_index)
    awal = first_available_date(
        lines,
        ["Tanggal Akad Awal", "Tanggal Awal Kredit", "Tanggal Mulai"],
    )
    is_bank_garansi = facility_type == BANK_GARANSI_FACILITY_TYPE

    return {
        "Nama": nama_debitur,
        IDENTITY_COLUMN: nik_debitur,
        "Bank": bank,
        "Penggunaan": "Bank Garansi" if is_bank_garansi else value_after_label(lines, "Jenis Penggunaan"),
        "Plafond (Rp)": max_number(plafon, plafon_awal),
        "Baki Debet (Rp)": baki_debet,
        "Kol.": kolektibilitas,
        "Awal": awal,
        "Jatuh Tempo": normalize_indonesian_date(value_after_label(lines, "Tanggal Jatuh Tempo")),
        "Rate": parse_rate_percent(value_after_label(lines, "Suku Bunga/Imbalan")),
        "No Laporan": nomor_laporan,
        "Keterangan": keterangan,
        "Jenis Fasilitas": facility_type,
        "No Rekening": value_after_label(lines, "No Rekening"),
        "No Akad Awal": value_after_label(lines, "No Akad Awal"),
        "No Akad Akhir": value_after_label(lines, "No Akad Akhir"),
        "Jenis Garansi": value_after_label(lines, "Jenis Garansi") if is_bank_garansi else "",
        "Tujuan Garansi": value_after_label(lines, "Tujuan Garansi") if is_bank_garansi else "",
        "Nama Yang Dijamin": value_after_label(lines, "Nama Yang Dijamin") if is_bank_garansi else "",
        "Nominal (Rp)": baki_debet if is_bank_garansi else pd.NA,
        "Setoran Jaminan (Rp)": parse_idr_number(value_after_label(lines, "Setoran Jaminan")) if is_bank_garansi else pd.NA,
        "Tanggal Diterbitkan": normalize_indonesian_date(value_after_label(lines, "Tanggal Diterbitkan")) if is_bank_garansi else "",
        "Tanggal Akad Akhir": normalize_indonesian_date(value_after_label(lines, "Tanggal Akad Akhir")),
    }


def build_no_matching_facility_row(nama_debitur, nik_debitur, nomor_laporan):
    return {
        "Nama": nama_debitur,
        IDENTITY_COLUMN: nik_debitur,
        "Bank": "",
        "Penggunaan": "",
        "Plafond (Rp)": 0,
        "Baki Debet (Rp)": 0,
        "Kol.": 0,
        "Awal": "",
        "Jatuh Tempo": "",
        "Rate": "",
        "No Laporan": nomor_laporan,
        "Keterangan": FALLBACK_KETERANGAN,
    }


def is_fallback_facility_row(row):
    try:
        keterangan = row.get("Keterangan", "")
    except AttributeError:
        keterangan = ""
    return clean_text(keterangan).casefold() == FALLBACK_KETERANGAN.casefold()


def count_real_facility_rows(df):
    if df is None or df.empty or "Keterangan" not in df.columns:
        return 0
    return int((~df.apply(is_fallback_facility_row, axis=1)).sum())


def parse_slik_pdf(pdf_file):
    pages = read_pdf_page_lines(pdf_file)
    lines = [line for page_lines in pages for line in page_lines]
    page1_lines = pages[0] if pages else []
    nama_debitur = extract_nama_debitur(page1_lines)
    nik_debitur = extract_debitur_identity(page1_lines)
    nomor_laporan = extract_nomor_laporan(lines)

    rows = []
    for block in split_facility_blocks(lines):
        row = parse_facility_block(block, nama_debitur, nomor_laporan, nik_debitur)
        if row:
            rows.append(row)

    if not rows and nama_debitur and nomor_laporan:
        rows.append(build_no_matching_facility_row(nama_debitur, nik_debitur, nomor_laporan))

    if not rows:
        return pd.DataFrame(columns=FACILITY_DATA_COLUMNS)

    return pd.DataFrame(rows, columns=FACILITY_DATA_COLUMNS)


def add_total_rows(df):
    rows = []

    for _, group in df.groupby("Nama", sort=False, dropna=False):
        rows.extend(group.to_dict("records"))

        total_row = {
            "Nama": "",
            IDENTITY_COLUMN: "",
            "Bank": "Jumlah",
            "Penggunaan": "",
            "Plafond (Rp)": group["Plafond (Rp)"].dropna().sum(),
            "Baki Debet (Rp)": group["Baki Debet (Rp)"].dropna().sum(),
            "Kol.": worst_kolektibilitas(group["Kol."]),
            "Awal": "",
            "Jatuh Tempo": "",
            "Rate": "",
            "No Laporan": "",
            "Keterangan": worst_keterangan(group["Keterangan"]),
        }
        rows.append(total_row)

    return pd.DataFrame(rows, columns=OUTPUT_COLUMNS)


def ensure_output_columns(df):
    if df is None:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)

    df = df.copy()
    if IDENTITY_COLUMN not in df.columns and "NIK" in df.columns:
        df[IDENTITY_COLUMN] = df["NIK"]
    for column in OUTPUT_COLUMNS:
        if column not in df.columns:
            df[column] = pd.NA
    return df[OUTPUT_COLUMNS].copy()


def ensure_facility_data_columns(df):
    if df is None:
        return pd.DataFrame(columns=FACILITY_DATA_COLUMNS)

    df = df.copy()
    if IDENTITY_COLUMN not in df.columns and "NIK" in df.columns:
        df[IDENTITY_COLUMN] = df["NIK"]
    for column in FACILITY_DATA_COLUMNS:
        if column not in df.columns:
            df[column] = pd.NA
    return df[FACILITY_DATA_COLUMNS].copy()


def filter_non_bank_garansi_facilities(df):
    if df is None or df.empty or "Jenis Fasilitas" not in df.columns:
        return df

    facility_type = df["Jenis Fasilitas"].fillna("").astype(str).str.strip()
    return df[facility_type != BANK_GARANSI_FACILITY_TYPE].copy()


def deduplicate_facilities(df):
    dedupe_columns = [
        "Nama",
        IDENTITY_COLUMN,
        "Jenis Fasilitas",
        "No Rekening",
        "No Akad Akhir",
        "Bank",
        "Penggunaan",
        "Plafond (Rp)",
        "Baki Debet (Rp)",
        "Kol.",
        "Awal",
        "Jatuh Tempo",
        "Rate",
        "Keterangan",
    ]
    dedupe_columns = [column for column in dedupe_columns if column in df.columns]

    return df.drop_duplicates(subset=dedupe_columns, keep="first").reset_index(drop=True)


def find_pdf_files(input_name=None):
    if input_name:
        input_path = Path(input_name)
        pdf_file = input_path if input_path.exists() else SLIK_DIR / input_name
        if pdf_file.suffix.lower() != ".pdf":
            pdf_file = pdf_file.with_suffix(".pdf")

        if pdf_file.exists():
            return [pdf_file]

        matches = sorted(SLIK_DIR.rglob(pdf_file.name))
        if matches:
            return matches

        if not pdf_file.exists():
            raise FileNotFoundError(f"PDF file not found: {pdf_file}")

    pdf_files = sorted(SLIK_DIR.rglob("*.pdf"))
    if not pdf_files:
        raise FileNotFoundError(f"No PDF files found in folder: {SLIK_DIR}")

    return pdf_files


def discover_pdf_groups(input_path):
    input_path = Path(input_path)

    if input_path.is_file():
        return {input_path.parent: [input_path]}

    pdf_files = sorted(input_path.rglob("*.pdf"))
    if not pdf_files:
        raise FileNotFoundError(f"No PDF files found in folder: {input_path}")

    child_dirs = [path for path in sorted(input_path.iterdir()) if path.is_dir()]
    groups = {}

    if child_dirs:
        for pdf_file in pdf_files:
            try:
                relative = pdf_file.relative_to(input_path)
            except ValueError:
                group_dir = pdf_file.parent
            else:
                group_dir = input_path / relative.parts[0] if len(relative.parts) > 1 else input_path
            groups.setdefault(group_dir, []).append(pdf_file)
    else:
        groups[input_path] = pdf_files

    return groups


def ensure_xlsx_name(name):
    name = clean_text(name) or "slik_extracted_table.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def resolve_output_file(output_value):
    if output_value:
        output_path = Path(ensure_xlsx_name(output_value))
        if output_path.parent == Path("."):
            return EXCEL_DIR / output_path.name
        return output_path

    return EXCEL_DIR / "slik_extracted_table.xlsx"


def safe_output_stem(name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean_text(name)).strip()
    return cleaned or "slik_extracted_table"


def output_file_for_group(group_dir, output_dir, output_name_template=None, force_unique=False):
    group_dir = Path(group_dir)
    folder_name = safe_output_stem(group_dir.name)

    if not output_name_template:
        return output_dir / f"{folder_name}.xlsx"

    output_name = ensure_xlsx_name(output_name_template)
    has_folder = "{folder}" in output_name
    try:
        output_name = output_name.format(folder=folder_name)
    except (KeyError, ValueError) as exc:
        raise ValueError("Output name template only supports {folder}.") from exc

    output_path = Path(output_name)
    if force_unique and not has_folder:
        output_path = output_path.with_name(f"{output_path.stem}_{folder_name}{output_path.suffix}")

    if output_path.parent == Path("."):
        return output_dir / output_path.name
    return output_path


def numeric_cell_value(value):
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value

    parsed = parse_idr_number(value)
    if pd.isna(parsed):
        return None
    return parsed


def integer_cell_value(value):
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)

    match = re.search(r"\d+", clean_text(value))
    return int(match.group(0)) if match else None


def apply_slik_total_formulas(sheet):
    group_start_row = 2

    for row_idx in range(2, sheet.max_row + 1):
        bank_value = sheet.cell(row=row_idx, column=3).value
        if bank_value != "Jumlah":
            continue

        first_data_row = group_start_row
        last_data_row = row_idx - 1
        if first_data_row <= last_data_row:
            sheet.cell(row=row_idx, column=5).value = f"=SUM(E{first_data_row}:E{last_data_row})"
            sheet.cell(row=row_idx, column=6).value = f"=SUM(F{first_data_row}:F{last_data_row})"
            sheet.cell(row=row_idx, column=7).value = f"=MAX(G{first_data_row}:G{last_data_row})"
        else:
            sheet.cell(row=row_idx, column=5).value = "=0"
            sheet.cell(row=row_idx, column=6).value = "=0"
            sheet.cell(row=row_idx, column=7).value = "=0"

        group_start_row = row_idx + 1


def normalize_slik_numeric_cells(sheet, row_idx, is_total):
    if not is_total:
        for col_idx in [5, 6]:
            cell = sheet.cell(row=row_idx, column=col_idx)
            numeric_value = numeric_cell_value(cell.value)
            if numeric_value is not None:
                cell.value = numeric_value

        kol_cell = sheet.cell(row=row_idx, column=7)
        kol_value = integer_cell_value(kol_cell.value)
        if kol_value is not None:
            kol_cell.value = kol_value


def style_slik_sheet(sheet):
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="BFBFBF")

    widths = {
        "A": 20,
        "B": 18,
        "C": 34,
        "D": 16,
        "E": 16,
        "F": 18,
        "G": 8,
        "H": 14,
        "I": 14,
        "J": 10,
        "K": 24,
        "L": 20,
    }

    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)

    apply_slik_total_formulas(sheet)

    for row_idx in range(2, sheet.max_row + 1):
        bank_value = sheet.cell(row=row_idx, column=3).value
        is_total = bank_value == "Jumlah"
        normalize_slik_numeric_cells(sheet, row_idx, is_total)

        if is_total:
            for col_idx in range(1, sheet.max_column + 1):
                sheet.cell(row=row_idx, column=col_idx).font = Font(bold=True)

        for col_idx in [5, 6]:
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        sheet.cell(row=row_idx, column=7).number_format = "0"

        rate_cell = sheet.cell(row=row_idx, column=10)
        rate_cell.number_format = "0.00%"

    merge_name_cells(sheet)


def merge_name_cells(sheet):
    start_row = None
    current_name = None

    for row_idx in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row_idx, column=1).value
        bank = sheet.cell(row=row_idx, column=3).value

        if bank == "Jumlah":
            if start_row and current_name and row_idx - start_row > 1:
                sheet.merge_cells(start_row=start_row, start_column=1, end_row=row_idx - 1, end_column=1)
            start_row = None
            current_name = None
            continue

        if name != current_name:
            if start_row and current_name and row_idx - start_row > 1:
                sheet.merge_cells(start_row=start_row, start_column=1, end_row=row_idx - 1, end_column=1)
            start_row = row_idx
            current_name = name


def month_count_inclusive(start_date, end_date):
    if end_date < start_date:
        return 0
    return (end_date.year - start_date.year) * 12 + (end_date.month - start_date.month) + 1


def build_installment_schedule_rows(facilities_df):
    schedule_rows = []
    summary_rows = []

    for facility_index, row in facilities_df.reset_index(drop=True).iterrows():
        start_date = parse_output_date(row.get("Awal", ""))
        end_date = parse_output_date(row.get("Jatuh Tempo", ""))
        plafond = row.get("Plafond (Rp)")

        if not start_date or not end_date or pd.isna(plafond) or plafond <= 0:
            continue

        term_months = month_count_inclusive(start_date, end_date)
        if term_months <= 0:
            continue

        annual_rate = row.get("Rate")
        if pd.isna(annual_rate):
            annual_rate = 0
        monthly_rate = annual_rate / 12
        installment_total = round(plafond / term_months)
        opening_balance = float(plafond)
        facility_no = facility_index + 1

        yearly = {}
        for period in range(1, term_months + 1):
            due_date = add_months_clamped(start_date, period - 1)
            if period == term_months:
                due_date = end_date

            total_payment = int(opening_balance) if period == term_months else installment_total
            interest = 0 if period == 1 else round(opening_balance * monthly_rate)
            principal = total_payment - interest
            ending_balance = max(0, round(opening_balance - total_payment))

            schedule_rows.append(
                {
                    "No": facility_no,
                    "Nama": row.get("Nama", ""),
                    "Bank": row.get("Bank", ""),
                    "Penggunaan": row.get("Penggunaan", ""),
                    "Kol.": row.get("Kol.", ""),
                    "Rate": annual_rate,
                    "Bunga Bulanan": monthly_rate,
                    "Ke": period,
                    "Jadwal": due_date,
                    "Pokok": principal,
                    "Bunga": interest if interest else "",
                    "Total": total_payment,
                    "Baki Debet": ending_balance if ending_balance else "",
                }
            )

            year_info = yearly.setdefault(due_date.year, {"cpltd": 0, "bunga": 0, "ltd": 0})
            year_info["cpltd"] += total_payment
            year_info["bunga"] += interest
            year_info["ltd"] = ending_balance
            opening_balance = ending_balance

        for year, year_info in yearly.items():
            summary_rows.append(
                {
                    "No": facility_no,
                    "Nama": row.get("Nama", ""),
                    "Bank": row.get("Bank", ""),
                    "Tahun": f"Des {year}",
                    "CPLTD": year_info["cpltd"],
                    "LTD": year_info["ltd"] if year_info["ltd"] else "",
                    "Bunga": year_info["bunga"] if year_info["bunga"] else "",
                }
            )

    return pd.DataFrame(schedule_rows), pd.DataFrame(summary_rows)


def style_angsuran_sheet(sheet):
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    summary_fill = PatternFill("solid", fgColor="FFF200")

    widths = {
        "A": 7,
        "B": 22,
        "C": 30,
        "D": 16,
        "E": 8,
        "F": 10,
        "G": 14,
        "H": 7,
        "I": 14,
        "J": 15,
        "K": 15,
        "L": 15,
        "M": 15,
        "O": 22,
        "P": 30,
        "Q": 12,
        "R": 15,
        "S": 15,
        "T": 15,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)

    for row_idx in range(2, sheet.max_row + 1):
        for col_idx in [6, 7]:
            sheet.cell(row=row_idx, column=col_idx).number_format = "0.00%"
        sheet.cell(row=row_idx, column=9).number_format = "dd/mm/yyyy"
        for col_idx in [10, 11, 12, 13, 18, 19, 20]:
            sheet.cell(row=row_idx, column=col_idx).number_format = "#,##0"
            sheet.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="right", vertical="center")
        if sheet.cell(row=row_idx, column=20).value not in ("", None):
            sheet.cell(row=row_idx, column=20).fill = summary_fill
        if sheet.cell(row=row_idx, column=17).value not in ("", None):
            sheet.cell(row=row_idx, column=17).font = Font(bold=True)

    sheet.freeze_panes = "A2"


def build_bank_garansi_output_df(facilities_df):
    facility_data_df = ensure_facility_data_columns(facilities_df)
    if facility_data_df.empty:
        return pd.DataFrame(columns=BANK_GARANSI_COLUMNS)

    facility_type = facility_data_df["Jenis Fasilitas"].fillna("").astype(str).str.strip()
    bank_garansi_df = facility_data_df[facility_type == BANK_GARANSI_FACILITY_TYPE].copy()
    if bank_garansi_df.empty:
        return pd.DataFrame(columns=BANK_GARANSI_COLUMNS)

    bank_garansi_df["Tanggal Akad Awal"] = bank_garansi_df["Awal"]
    for column in BANK_GARANSI_COLUMNS:
        if column not in bank_garansi_df.columns:
            bank_garansi_df[column] = pd.NA

    return bank_garansi_df[BANK_GARANSI_COLUMNS].copy()


def style_bank_garansi_sheet(sheet):
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")

    widths = {
        "A": 20,
        "B": 18,
        "C": 32,
        "D": 16,
        "E": 20,
        "F": 18,
        "G": 28,
        "H": 16,
        "I": 16,
        "J": 18,
        "K": 8,
        "L": 16,
        "M": 16,
        "N": 16,
        "O": 16,
        "P": 24,
        "Q": 24,
        "R": 24,
        "S": 16,
    }

    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    headers = {cell.value: cell.column for cell in sheet[1]}
    currency_columns = {"Plafond (Rp)", "Nominal (Rp)", "Setoran Jaminan (Rp)"}
    date_columns = {"Tanggal Akad Awal", "Tanggal Akad Akhir", "Tanggal Diterbitkan", "Jatuh Tempo"}

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)

    for row_idx in range(2, sheet.max_row + 1):
        for column_name in currency_columns:
            col_idx = headers.get(column_name)
            if not col_idx:
                continue
            cell = sheet.cell(row=row_idx, column=col_idx)
            numeric_value = numeric_cell_value(cell.value)
            if numeric_value is not None:
                cell.value = numeric_value
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        kol_col = headers.get("Kol.")
        if kol_col:
            kol_cell = sheet.cell(row=row_idx, column=kol_col)
            kol_value = integer_cell_value(kol_cell.value)
            if kol_value is not None:
                kol_cell.value = kol_value
            kol_cell.number_format = "0"

        for column_name in date_columns:
            col_idx = headers.get(column_name)
            if col_idx:
                sheet.cell(row=row_idx, column=col_idx).number_format = "dd-mm-yyyy"

    sheet.freeze_panes = "A2"


def write_bank_garansi_sheet(writer, facilities_df):
    bank_garansi_df = build_bank_garansi_output_df(facilities_df)
    if bank_garansi_df.empty:
        return False

    bank_garansi_df.to_excel(writer, sheet_name="Bank_Garansi", index=False)
    style_bank_garansi_sheet(writer.sheets["Bank_Garansi"])
    return True


def write_installment_sheet(writer, facilities_df):
    schedule_df, summary_df = build_installment_schedule_rows(facilities_df)

    schedule_columns = [
        "No",
        "Nama",
        "Bank",
        "Penggunaan",
        "Kol.",
        "Rate",
        "Bunga Bulanan",
        "Ke",
        "Jadwal",
        "Pokok",
        "Bunga",
        "Total",
        "Baki Debet",
    ]
    summary_columns = ["Nama", "Bank", "Tahun", "CPLTD", "LTD", "Bunga"]

    if schedule_df.empty:
        schedule_df = pd.DataFrame(columns=schedule_columns)
    if summary_df.empty:
        summary_df = pd.DataFrame(columns=["No", *summary_columns])

    schedule_df.to_excel(writer, sheet_name="Angsuran", index=False, columns=schedule_columns)
    sheet = writer.sheets["Angsuran"]

    start_col = len(schedule_columns) + 2
    for offset, header in enumerate(summary_columns, start=start_col):
        cell = sheet.cell(row=1, column=offset, value=header)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)

    summary_by_no = {facility_no: group for facility_no, group in summary_df.groupby("No", sort=False)}
    for row_idx in range(2, sheet.max_row + 1):
        facility_no = sheet.cell(row=row_idx, column=1).value
        period = sheet.cell(row=row_idx, column=8).value
        if period != 1 or facility_no not in summary_by_no:
            continue

        for summary_offset, (_, summary_row) in enumerate(summary_by_no[facility_no].iterrows()):
            target_row = row_idx + summary_offset
            for col_offset, column_name in enumerate(summary_columns, start=start_col):
                sheet.cell(row=target_row, column=col_offset, value=summary_row[column_name])

    style_angsuran_sheet(sheet)


def _report_progress(progress_callback, **payload):
    if not progress_callback:
        return

    try:
        progress_callback(**payload)
    except TypeError:
        progress_callback(payload["processed"], payload["total"], payload["file"])


def extract_slik_facilities(pdf_files, progress_callback=None, row_metadata_callback=None):
    """
    Parse multiple SLIK PDF files and return:
    - facilities_df
    - duplicate_count
    - errors_df

    This function only extracts data. It does not write Excel files.
    """
    pdf_files = [Path(pdf_file) for pdf_file in pdf_files]
    all_results = []
    errors = []

    for index, pdf_file in enumerate(pdf_files, start=1):
        try:
            df = parse_slik_pdf(pdf_file)
            if row_metadata_callback:
                row_metadata = row_metadata_callback(pdf_file) or {}
                for column, value in row_metadata.items():
                    df[column] = value
            real_rows_found = count_real_facility_rows(df)
        except Exception as exc:
            errors.append({"file": pdf_file.name, "error": str(exc)})
            _report_progress(
                progress_callback,
                processed=index,
                total=len(pdf_files),
                file=pdf_file,
                rows_found=0,
                error=str(exc),
            )
            continue

        if not df.empty:
            all_results.append(df)

        _report_progress(
            progress_callback,
            processed=index,
            total=len(pdf_files),
            file=pdf_file,
            rows_found=real_rows_found,
            error=None,
        )

    errors_df = pd.DataFrame(errors, columns=["file", "error"])
    if not all_results:
        return pd.DataFrame(columns=FACILITY_DATA_COLUMNS), 0, errors_df

    facilities_df = pd.concat(all_results, ignore_index=True)
    before_dedupe_count = len(facilities_df)
    facilities_df = deduplicate_facilities(facilities_df)
    duplicate_count = before_dedupe_count - len(facilities_df)

    return facilities_df, duplicate_count, errors_df


def export_slik_excel_dashboard(
    facilities_df,
    output_file,
    include_excel_dashboard=True,
    include_angsuran=True,
    data_quality_df=None,
    include_data_quality=False,
):
    """
    Export the final SLIK workbook while preserving the existing sheet layout.
    """
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    facility_data_df = ensure_facility_data_columns(facilities_df)

    if facility_data_df.empty:
        final_df = pd.DataFrame(columns=OUTPUT_COLUMNS)
        dashboard_facilities_df = pd.DataFrame(columns=OUTPUT_COLUMNS)
    else:
        dashboard_facilities_df = ensure_output_columns(facility_data_df)
        final_df = add_total_rows(dashboard_facilities_df)

    installment_facilities_df = ensure_output_columns(filter_non_bank_garansi_facilities(facility_data_df))

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        final_df.to_excel(writer, sheet_name="SLIK", index=False)
        style_slik_sheet(writer.sheets["SLIK"])
        write_bank_garansi_sheet(writer, facility_data_df)

        if include_excel_dashboard:
            write_dashboard_sheets(writer, dashboard_facilities_df)

        if include_angsuran:
            write_installment_sheet(writer, installment_facilities_df)

        if include_data_quality and data_quality_df is not None and not data_quality_df.empty:
            data_quality_df.to_excel(writer, sheet_name="Data_Quality", index=False)

    return output_file


def export_to_excel(df, output_file, facilities_df):
    output_file.parent.mkdir(parents=True, exist_ok=True)
    df = ensure_output_columns(df)
    facility_data_df = ensure_facility_data_columns(facilities_df)
    facilities_df = ensure_output_columns(facility_data_df)
    installment_facilities_df = ensure_output_columns(filter_non_bank_garansi_facilities(facility_data_df))
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="SLIK", index=False)
        style_slik_sheet(writer.sheets["SLIK"])
        write_bank_garansi_sheet(writer, facility_data_df)
        write_dashboard_sheets(writer, facilities_df)
        write_installment_sheet(writer, installment_facilities_df)


def process_pdf_files(pdf_files, output_file):
    def print_progress(processed, total, file, rows_found=0, error=None):
        print(f"Reading: {file}")
        if error:
            print(f"Error reading {file}: {error}")
        elif not rows_found:
            print(
                "No active, written-off/charged-off, or Kol. other than 1 "
                f"credit facilities found: {file}"
            )

    facilities_df, duplicate_count, errors_df = extract_slik_facilities(
        pdf_files,
        progress_callback=print_progress,
    )

    if facilities_df.empty:
        print(
            "No data written because no active, written-off/charged-off, "
            "or Kol. other than 1 credit facilities were found."
        )
        return None

    if not errors_df.empty:
        print(f"{len(errors_df)} file(s) failed and were skipped.")

    print("Writing Excel...")
    export_slik_excel_dashboard(facilities_df, output_file)

    return {
        "output_file": output_file,
        "facility_count": len(facilities_df),
        "duplicate_count": duplicate_count,
        "error_count": len(errors_df),
        "errors_df": errors_df,
    }


def choose_folders_with_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except ImportError:
        return None, None, None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    messagebox.showinfo(
        "SLIK PDF to Excel",
        "Choose the folder that contains your SLIK PDF files.",
        parent=root,
    )
    slik_dir = filedialog.askdirectory(
        title="Choose SLIK PDF folder",
        initialdir=str(SLIK_DIR.resolve()) if SLIK_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not slik_dir:
        root.destroy()
        return None, None, None

    messagebox.showinfo(
        "SLIK PDF to Excel",
        "Choose the folder where the Excel output should be saved.",
        parent=root,
    )
    output_dir = filedialog.askdirectory(
        title="Choose output folder",
        initialdir=str(EXCEL_DIR.resolve()) if EXCEL_DIR.exists() else str(Path.cwd()),
        parent=root,
    )
    if not output_dir:
        root.destroy()
        return None, None, None

    output_name = simpledialog.askstring(
        "SLIK PDF to Excel",
        "Enter Excel filename pattern:\nUse {folder} to name each SLIK subfolder workbook.",
        initialvalue="{folder}.xlsx",
        parent=root,
    )
    if output_name is None:
        root.destroy()
        return None, None, None

    root.destroy()
    return Path(slik_dir), Path(output_dir), ensure_xlsx_name(output_name)


def show_gui_result(message, is_error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
    except ImportError:
        return

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    if is_error:
        messagebox.showerror("SLIK PDF to Excel", message, parent=root)
    else:
        messagebox.showinfo("SLIK PDF to Excel", message, parent=root)
    root.destroy()


def parse_args():
    parser = argparse.ArgumentParser(
        description="Extract selected SLIK credit facility fields to an Excel table."
    )
    parser.add_argument(
        "pdf_name",
        nargs="?",
        help="Optional PDF filename in the slik folder. Extension can be omitted.",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Optional output Excel filename or pattern. Folder mode supports {folder}. Default: {folder}.xlsx",
    )
    parser.add_argument(
        "--slik-dir",
        default=str(SLIK_DIR),
        help="Folder containing SLIK PDF files. Subfolders are included. Default: slik",
    )
    return parser.parse_args()


def main():
    global SLIK_DIR

    args = parse_args()
    use_gui = len(sys.argv) == 1

    if use_gui:
        selected_slik_dir, output_dir, output_name_template = choose_folders_with_gui()
        if not selected_slik_dir or not output_dir:
            print("Cancelled. No folders selected.")
            return
        SLIK_DIR = selected_slik_dir
    else:
        SLIK_DIR = Path(args.slik_dir)
        output_dir = EXCEL_DIR
        output_name_template = args.output

    output_dir.mkdir(parents=True, exist_ok=True)
    written_files = []

    if args.pdf_name:
        pdf_files = find_pdf_files(args.pdf_name)
        output_file = resolve_output_file(args.output)
        result = process_pdf_files(pdf_files, output_file)
        if result:
            written_files.append(result)
    else:
        groups = discover_pdf_groups(SLIK_DIR)
        force_unique = len(groups) > 1

        for group_dir, pdf_files in sorted(groups.items(), key=lambda item: str(item[0]).lower()):
            output_file = output_file_for_group(
                group_dir,
                output_dir,
                output_name_template,
                force_unique=force_unique,
            )
            print(f"\nProcessing folder: {group_dir}")
            result = process_pdf_files(pdf_files, output_file)
            if result:
                written_files.append(result)

    if not written_files:
        message = (
            "No data written because no active, written-off/charged-off, "
            "or Kol. other than 1 credit facilities were found."
        )
        if use_gui:
            show_gui_result(message, is_error=True)
        return

    print("\nSemua proses selesai.")
    for result in written_files:
        print(f"Extracted {result['facility_count']} credit facility row(s).")
        if result["duplicate_count"]:
            print(f"Removed {result['duplicate_count']} duplicate credit facility row(s).")
        print(f"Data successfully written to {result['output_file']}")

    if use_gui:
        saved_files = "\n".join(str(result["output_file"]) for result in written_files)
        show_gui_result(
            "Done.\n\nSaved files:\n" + saved_files
        )


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}")
        if len(sys.argv) == 1:
            show_gui_result(str(exc), is_error=True)
            sys.exit(1)
        raise
