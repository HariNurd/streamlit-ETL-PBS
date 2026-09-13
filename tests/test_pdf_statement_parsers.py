import copy
import importlib
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from services.pdf_statement_adapter import statement_frames, summary_metrics
from statement_pdf import bri_recap, order_same_timestamp, read_mandiri, validate_statement, validate_recap


def sample(bank='BCA'):
    # All identities and amounts in this module are synthetic, not PDF fixtures.
    return {'bank_name':bank,'account_number':'00001','account_type':'REKENING GIRO',
            'statement_period':'JANUARI 2025','currency':'IDR',
            'summary':{'opening_balance':'1000.00','debit_amount_total':'30.00',
                       'credit_amount_total':'10.00','closing_balance':'980.00',
                       'debit_transaction_count':1,'credit_transaction_count':1},
            'transactions':[
                {'date':'01/01','description':'BIAYA ADMIN','debit':'30.00','credit':None,
                 'balance':'970.00','posting_datetime':'2025-01-01T01:00:00','source_row':1},
                {'date':'31/01','description':'JASA GIRO','debit':None,'credit':'10.00',
                 'balance':'980.00','posting_datetime':'2025-01-31T23:59:00','source_row':2}]}


class PdfStatementTests(unittest.TestCase):
    def test_bri_recap_amounts_on_next_page(self):
        text = '''Saldo Awal    Total Transaksi Debet    Total Transaksi Kredit    Saldo Akhir
Opening Balance    Total Debit Transaction    Total Credit Transaction    Closing Balance



Created By IBBIZ
07/29/2026 17:57:37

LAPORAN TRANSAKSI FINANSIAL
STATEMENT OF FINANCIAL TRANSACTION

Halaman 4 dari 4
Page 4 of 4

1,000.00    300.00    500.00    1,200.00
Terbilang / In Words
'''
        summary = bri_recap(text)
        self.assertEqual(summary, {'opening_balance':'1000.00', 'debit_amount_total':'300.00',
                                   'credit_amount_total':'500.00', 'closing_balance':'1200.00'})
        validate_recap(summary)

    def test_bri_recap_does_not_pick_amounts_beyond_terbilang(self):
        text = '''Saldo Awal    Total Transaksi Debet    Total Transaksi Kredit    Saldo Akhir
Created By IBBIZ
Terbilang / In Words
1,000.00    300.00    500.00    1,200.00
'''
        with self.assertRaisesRegex(ValueError, 'Ambiguous printed BRI recap amounts'):
            bri_recap(text)

    def test_bri_recap_rejects_multiple_amount_rows(self):
        text = '''Saldo Awal    Total Transaksi Debet    Total Transaksi Kredit    Saldo Akhir
1,000.00    300.00    500.00    1,200.00
2,000.00    300.00    500.00    2,200.00
Terbilang / In Words
'''
        with self.assertRaisesRegex(ValueError, 'Ambiguous printed BRI recap amounts'):
            bri_recap(text)

    def test_inconsistent_printed_recap_requires_review(self):
        summary=sample()['summary'];summary['closing_balance']='981.00'
        with self.assertRaisesRegex(ValueError,'Printed recap is inconsistent'):validate_recap(summary)

    def test_rejects_balanced_but_wrong_printed_side_totals(self):
        parsed=sample();parsed['summary']['debit_amount_total']='31.00'
        with self.assertRaisesRegex(ValueError,'debit total mismatch'):validate_statement(parsed)

    def test_rejects_wrong_count_even_when_amounts_balance(self):
        parsed=sample();parsed['summary']['debit_transaction_count']=2
        with self.assertRaisesRegex(ValueError,'count mismatch'):validate_statement(parsed)

    def test_rejects_intermediate_balance_mismatch(self):
        parsed=sample();parsed['transactions'][0]['balance']='971.00'
        with self.assertRaisesRegex(ValueError,'printed balance mismatch'):validate_statement(parsed)

    def test_same_day_order_preserves_amounts_and_source_rows(self):
        rows=[{'date':'01/01','posting_datetime':'2025-01-01T01:00:00','description':'',
               'debit':None,'credit':'1.00','balance':'101.00','source_row':1},
              {'date':'01/01','posting_datetime':'2025-01-01T12:00:00','description':'',
               'debit':None,'credit':'100.00','balance':'100.00','source_row':2}]
        before=copy.deepcopy(rows)
        ordered=order_same_timestamp(rows,'0')
        self.assertEqual([r['source_row'] for r in ordered],[2,1])
        self.assertEqual(rows,before)

    def test_valid_repeated_balance_sequence_keeps_source_order(self):
        rows=[]
        for i,(db,cr,bal) in enumerate([('10',None,'90'),(None,'5','95'),(None,'5','100'),('20',None,'80')]):
            rows.append({'date':'01/01','posting_datetime':f'2025-01-01T12:0{i}:00',
                         'debit':db,'credit':cr,'balance':bal})
        self.assertEqual(order_same_timestamp(rows,'100'),rows)

    def test_mandiri_wrapped_remark_and_reference(self):
        rows=read_mandiri(['''01/04/2024 12:54:
08
TRANSFER ONE
- 10.00 0.00 90.00
01/04/2024 17:23:
57
REFERENCE PREFIX
WRAPPED REMARK
REF123 0.00 20.00 110.00
'''])
        self.assertEqual(rows[0]['description'],'TRANSFER ONE')
        self.assertEqual(rows[1]['description'],'REFERENCE PREFIX WRAPPED REMARK REF123')
        self.assertEqual(rows[1]['credit'],'20.00')

    def test_mandiri_never_silently_drops_incomplete_transaction(self):
        with self.assertRaisesRegex(ValueError,'missing amounts'):
            read_mandiri(['01 Jan 2025,\n12:00:00\nTRANSFER\n'])

    def test_fee_classification_excludes_transfer_principal_and_unlabeled_rows(self):
        df=pd.DataFrame([
            {'Keterangan':'MCM InhouseTrf KE SOMEONE Transfer Fee','DB':Decimal('1000000'),'CR':pd.NA},
            {'Keterangan':'FEE_PIS_Payroll','DB':Decimal('5000'),'CR':pd.NA},
            {'Keterangan':'Biaya Admin','DB':Decimal('3000'),'CR':pd.NA},
            {'Keterangan':'Tax','DB':Decimal('200'),'CR':pd.NA},
            {'Keterangan':'Interest on Account','DB':pd.NA,'CR':Decimal('1000')},
            {'Keterangan':'','DB':Decimal('10000'),'CR':pd.NA},
        ])
        df.attrs['account_type']='Giro Special Rate'
        metrics=summary_metrics(df)
        self.assertEqual(metrics['Adm'],Decimal('8000'))
        self.assertEqual(metrics['Pajak'],Decimal('200'))
        self.assertEqual(metrics['JaGir'],Decimal('1000'))
        self.assertTrue(pd.isna(metrics['Bunga']))

    def test_all_pdf_converters_use_shared_entry_point(self):
        for bank,module in [('BCA','bca'),('BNI','bni'),('BRI','BRI'),('DKI','dki'),('Mandiri','mandiri')]:
            with self.subTest(bank=bank),patch('services.pdf_statement_adapter.extract_statement',return_value=({'gt_parse':sample(bank)},1)) as parser:
                converter=importlib.import_module('convert_mutasi_'+module)
                df,recap=converter.process_pdf(Path('sample.pdf'))
                parser.assert_called_once_with(Path('sample.pdf'),bank=bank)
                self.assertEqual(len(df),2)
                self.assertIn('Saldo Awal',recap.columns)

    def test_all_single_pdf_exports_include_summary_fee_values(self):
        metadata={'account':'00001','month_label':'Jan-25','month_order':1,'year':'2025'}
        with tempfile.TemporaryDirectory() as directory:
            for bank,module in [('BCA','bca'),('BNI','bni'),('BRI','BRI'),('DKI','dki'),('Mandiri','mandiri')]:
                converter=importlib.import_module('convert_mutasi_'+module)
                df,recap=statement_frames(sample(bank))
                path=Path(directory)/f'{bank}.xlsx'
                with self.subTest(bank=bank),patch.object(converter,'extract_pdf_metadata',return_value=metadata):
                    converter.export_to_excel(df,recap,path,Path('sample.pdf'))
                    book=load_workbook(path,data_only=True)
                    sheet=book['Summary']
                    self.assertEqual(sheet.cell(3,9).value,30)
                    self.assertEqual(sheet.cell(3,13).value,10)
                    book.close()

    def test_jakarta_multi_account_pdf_exports_separate_balances(self):
        import convert_mutasi_dki as converter
        first=sample('DKI');second=sample('DKI');second['account_number']='00002'
        label={'gt_parse':first,'statements':[first,second]}
        metadata={'account':first['account_number'],'month_label':'Jan-25','month_order':1,'year':'2025'}
        with patch('services.pdf_statement_adapter.extract_statement',return_value=(label,1)):
            df,recap=converter.process_pdf(Path('sample.pdf'))
        self.assertEqual(df['Account'].nunique(),2)
        self.assertEqual(len(recap),2)
        with tempfile.TemporaryDirectory() as directory,patch.object(converter,'extract_pdf_metadata',return_value=metadata):
            output=Path(directory)/'multi.xlsx'
            converter.export_to_excel(df,recap,output,Path('sample.pdf'))
            book=load_workbook(output,data_only=True)
            self.assertEqual(len(book.sheetnames),3)
            for row,account in [(3,first['account_number']),(4,second['account_number'])]:
                self.assertIn(account,book['Summary'].cell(row,2).value)
                self.assertEqual(book['Summary'].cell(row,7).value,980)
                self.assertEqual(book['Summary'].cell(row,9).value,30)
            book.close()


if __name__=='__main__':unittest.main()
