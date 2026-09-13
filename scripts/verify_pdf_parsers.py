"""Exercise every local PDF through detection, parsing and Excel export."""
import contextlib
import importlib
import io
import json
import sys
from collections import Counter
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
from converters.bank_detector import detect_bank_from_pdf
from services.pdf_statement_adapter import summary_metrics
import pandas as pd

MODULES={'BCA':'bca','BNI':'bni','BRI':'BRI','DKI':'dki','Mandiri':'mandiri'}


def main():
    out=ROOT/'outputs'/'pdf_parser_validation';out.mkdir(parents=True,exist_ok=True)
    results=[]
    for index,path in enumerate(sorted((ROOT/'pdf').rglob('*.pdf')),1):
        entry={'file':str(path.relative_to(ROOT))}
        try:
            bank=detect_bank_from_pdf(path);entry['bank']=bank
            module=importlib.import_module('convert_mutasi_'+MODULES[bank])
            with contextlib.redirect_stdout(io.StringIO()):
                df,summary=module.process_pdf(path)
                target=out/'workbooks'/f'{index:03d}.xlsx'
                target.parent.mkdir(exist_ok=True)
                module.export_to_excel(df,summary,target,path)
            book=load_workbook(target,data_only=False)
            assert 'Summary' in book.sheetnames
            header=[book['Summary'].cell(1,col).value for col in range(1,14)]
            assert all(key in header for key in ['Adm','Pajak','JaGir']),header
            if bank!='DKI' or df['Account'].nunique()==1:
                for key,value in summary_metrics(df).items():
                    actual=book['Summary'].cell(3,header.index(key)+1).value
                    assert actual==(None if pd.isna(value) else float(value)),(key,actual,value)
            entry.update(status='validated',rows=len(df),sheets=book.sheetnames,
                         fees={k:None if pd.isna(v) else str(v) for k,v in summary_metrics(df).items()},
                         workbook=str(target.relative_to(ROOT)))
            book.close()
        except Exception as exc:
            entry.update(status='rejected',error=str(exc))
        results.append(entry)
        print(index,path.name,entry['status'],entry.get('error',''),flush=True)
    (out/'results.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
    print(Counter((r.get('bank'),r['status']) for r in results))


if __name__=='__main__':main()
